"""Revise the user-owned SRS only; no application/database mutation."""
from pathlib import Path
from copy import deepcopy
import hashlib, json, re, sys, uuid, zipfile
from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader
import revise_srs_php as P

ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT.parent/'參考資料/Fitness_Tracking_Management_System_SRS_中文版_純文字排版版.docx'
OUT=ROOT/'output/specification/Fitness_Tracking_Management_System_SRS_中文版_教師建議修訂版.docx'
TMP=ROOT/'tmp/srs_teacher_feedback'
TMP.mkdir(parents=True,exist_ok=True)
P.TMP=TMP; P.OUTPUT=OUT
def toc_pages_from_outline(pdf_path,check_only=False):
    pdf=PdfReader(pdf_path)
    page_map={}
    def walk(items):
        for item in items:
            if isinstance(item,list):walk(item)
            elif isinstance(item,dict) and '/Title' in item:
                page_map[str(item['/Title'])]=pdf.get_destination_page_number(item)+1
    walk(pdf.outline)
    headings=json.loads((TMP/'headings.json').read_text())
    missing=[title for title,level,name in headings if level<=2 and title not in page_map]
    assert not missing,missing
    d=Document(OUT);count=0
    for p in d.paragraphs:
        if p.style.name.lower() not in ('toc 1','toc 2'):continue
        texts=p._p.xpath('.//w:t');expected=str(page_map[texts[0].text])
        if check_only:assert texts[-1].text==expected,(texts[0].text,texts[-1].text,expected)
        else:texts[-1].text=expected
        count+=1
    if not check_only:
        d.save(OUT);(TMP/'toc_pages.json').write_text(json.dumps(page_map,ensure_ascii=False,indent=2))
    print(json.dumps({'toc_entries':count,'pdf_pages':len(pdf.pages),'verified':check_only},ensure_ascii=False))
if '--toc-pdf' in sys.argv:
    toc_pages_from_outline(sys.argv[sys.argv.index('--toc-pdf')+1], '--check' in sys.argv)
    raise SystemExit
D=Document(SRC); digest=hashlib.sha256(SRC.read_bytes()).hexdigest()
original=list(D.element.body)
def pchange(i,text): P.change(Paragraph(original[i],D),text)
def x(tag,**attrs):
    e=OxmlElement('w:'+tag)
    for k,v in attrs.items():e.set(qn('w:'+k),str(v))
    return e
def embed_editable_font(docx_path,font_path,font_name):
    """Embed an editable-installable TrueType font for deterministic CJK rendering."""
    key=uuid.uuid4();font=bytearray(Path(font_path).read_bytes());guid_bytes=key.bytes_le
    for i in range(min(32,len(font))):font[i]^=guid_bytes[15-(i%16)]
    tmp=TMP/'font_embedded.docx'
    W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    PR='http://schemas.openxmlformats.org/package/2006/relationships'
    CT='http://schemas.openxmlformats.org/package/2006/content-types'
    from lxml import etree
    with zipfile.ZipFile(docx_path,'r') as zin:
        parts={n:zin.read(n) for n in zin.namelist()}
    ft=etree.fromstring(parts['word/fontTable.xml'])
    node=next((n for n in ft.findall(f'{{{W}}}font') if n.get(f'{{{W}}}name')==font_name),None)
    if node is None:
        node=etree.SubElement(ft,f'{{{W}}}font');node.set(f'{{{W}}}name',font_name)
    for old in node.findall(f'{{{W}}}embedRegular'):node.remove(old)
    emb=etree.SubElement(node,f'{{{W}}}embedRegular')
    emb.set(f'{{{R}}}id','rIdEmbeddedArialUnicode')
    emb.set(f'{{{W}}}fontKey','{'+str(key).upper()+'}')
    emb.set(f'{{{W}}}subsetted','0')
    parts['word/fontTable.xml']=etree.tostring(ft,xml_declaration=True,encoding='UTF-8',standalone=True)
    rel_name='word/_rels/fontTable.xml.rels'
    rel=etree.fromstring(parts[rel_name]) if rel_name in parts else etree.Element(f'{{{PR}}}Relationships',nsmap={None:PR})
    for old in list(rel):
        if old.get('Id')=='rIdEmbeddedArialUnicode':rel.remove(old)
    item=etree.SubElement(rel,f'{{{PR}}}Relationship')
    item.set('Id','rIdEmbeddedArialUnicode');item.set('Type',R+'/font');item.set('Target','fonts/arialunicode.odttf')
    parts[rel_name]=etree.tostring(rel,xml_declaration=True,encoding='UTF-8',standalone=True)
    cts=etree.fromstring(parts['[Content_Types].xml'])
    if not any(n.get('Extension')=='odttf' for n in cts.findall(f'{{{CT}}}Default')):
        item=etree.SubElement(cts,f'{{{CT}}}Default');item.set('Extension','odttf');item.set('ContentType','application/vnd.openxmlformats-officedocument.obfuscatedFont')
    parts['[Content_Types].xml']=etree.tostring(cts,xml_declaration=True,encoding='UTF-8',standalone=True)
    settings=etree.fromstring(parts['word/settings.xml'])
    for tag in ('embedTrueTypeFonts','embedSystemFonts'):
        n=settings.find(f'{{{W}}}{tag}')
        if n is None:n=etree.SubElement(settings,f'{{{W}}}{tag}')
        n.set(f'{{{W}}}val','1')
    subset=settings.find(f'{{{W}}}saveSubsetFonts')
    if subset is not None:settings.remove(subset)
    parts['word/settings.xml']=etree.tostring(settings,xml_declaration=True,encoding='UTF-8',standalone=True)
    parts['word/fonts/arialunicode.odttf']=bytes(font)
    with zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as zout:
        for name,data in parts.items():zout.writestr(name,data)
    tmp.replace(docx_path)
def fmt(p,size=11):
    p.paragraph_format.widow_control=True
    p._p.get_or_add_pPr().append(x('snapToGrid',val='0'))
    for r in p.runs:
        r.font.name='Arial'; r.font.size=Pt(size);r.font.color.rgb=RGBColor(0,0,0)
        r._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'PingFang TC')
def ins(a,text,style='Normal',page=False):
    p=a.insert_paragraph_before(text,style)
    fmt(p,18 if style=='Heading 1' else 14 if style=='Heading 2' else 12 if style=='Heading 3' else 11)
    p.paragraph_format.space_after=Pt(6)
    p.paragraph_format.line_spacing=1.25
    if page:p.paragraph_format.page_break_before=True
    if style.startswith('Heading'):p.paragraph_format.keep_with_next=True
    return p
def tab(a,heads,rows,widths=None,size=9.5):
    t=D.add_table(rows=1, cols=len(heads));t.autofit=False
    widths=widths or [16.7/len(heads)]*len(heads)
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for vals in [heads]+list(rows):
        row=t.rows[0] if vals is heads else t.add_row()
        row._tr.get_or_add_trPr().append(x('cantSplit'))
        for c,v,w in zip(row.cells,vals,widths):
            c.width=Cm(w);c.text=str(v)
            pr=c._tc.get_or_add_tcPr();m=x('tcMar')
            for edge in ('top','bottom'):m.append(x(edge,w=75,type='dxa'))
            for edge in ('left','right'):m.append(x(edge,w=85,type='dxa'))
            pr.append(m)
            for p in c.paragraphs:
                fmt(p,size);p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.1
                p.paragraph_format.keep_with_next=False
                for r in p.runs:r.bold=vals is heads
    t.rows[0]._tr.get_or_add_trPr().append(x('tblHeader'))
    b=x('tblBorders')
    for k in ('top','left','bottom','right','insideH','insideV'):b.append(x(k,val='single',sz=4,color='808080'))
    t._tbl.tblPr.append(b);a._p.addprevious(t._tbl);ins(a,'')
    return t
def code(a,s):
    for line in s.strip().splitlines():
        p=ins(a,line or ' ')
        p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
        for r in p.runs:r.font.name='Courier New';r.font.size=Pt(9)
def h(a,s):return ins(a,s,'Heading 2',True)
def pic(a,path,caption,width=16.5):
    p=ins(a,'');p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.keep_with_next=True;p.alignment=1
    r=p.add_run();r.add_picture(str(path),width=Cm(width))
    for e in r._r.xpath('.//wp:docPr'):e.set('descr',caption+'；完整外鍵與基數見第 4.3 節。')
    cp=ins(a,caption,'Caption');cp.alignment=1

# Minimal local changes to functional requirements; preserve user background and scope.
changes={
76:'多表寫入使用同一 PDO 連線。beginTransaction() 後、第一次資料查詢前執行 SET TRANSACTION ISOLATION LEVEL READ COMMITTED；COMMIT 成功才回應成功，任何 SQL 或提交錯誤均 rollBack()。錯誤訊息不得包含 SQL、憑證或他人資料。完整併發協定見 4.4。',
88:'3. 使用 firebase/php-jwt 簽發固定 HS256 Access JWT：sub 為 user_id 的十進位字串、jti=user_session_id，另含 iat、exp，效期 7,200 秒且不含 role。另產生 256-bit 不透明 Refresh Token；瀏覽器只以 Secure、HttpOnly、SameSite=Strict、Path=/ 且不設定 Domain 的 __Host-refresh_token Cookie 保存原值，資料庫只存 SHA-256。提交成功後才回傳 Access Token 並設定 Cookie。',
98:'2. 每次受保護請求驗證簽章、固定 HS256、exp、iat 與 jti，再 JOIN user_sessions、user_auth 確認 sub 對應本人、期限有效、未撤銷且帳號啟用。只有管理端點另查 admins；一般會員端點不查管理員表。頁面重整後以 POST /api/auth/refresh、同源 Origin 檢查及 HttpOnly Cookie 輪替 Refresh Token，取得新的 Access Token；舊 Refresh Token 立即失效。管理寫入與資格撤銷的併發順序見 4.7。',
121:'輸出／回應：GET /api/me 或 PUT /api/me/profile → 200；回傳 profile、目前體重、訓練設定、依 priority 排序的偏好部位與目前 role。role 由後端查 admins 推導，只供介面顯示，不能作為後端授權依據。',
130:'2. 在 users 更新選項；每份新課表另以 planning_profiles 與 planning_profile_body_parts 保存當次輸入值及偏好順序。已發布課表與其輸入快照不因目前會員資料改變而更新。',
148:'1. 以參數化查詢取得 exercises 各自 MAX(version_no) 的 exercise_versions，JOIN body_parts；keyword 比對該版本名稱與部位名稱。預設只顯示 exercises.is_active=TRUE，依名稱及 exercise_id 排序。此為目前目錄，不用於歷史顯示。',
155:'輸入資料：exercise_name、body_part_id、difficulty_level、equipment、description、image_url、is_active。識別與啟用狀態存 exercises，描述性內容存 exercise_versions，限制見 4.2。',
157:'1. 在管理寫入交易內確認 admins 資格；圖片只接受本站已核准之不可覆寫資源 URL，不下載任意用戶提供的 URL。資源由部署程序驗證授權、保存實體檔案並計算 SHA-256。',
158:'2. 新增動作同交易建立 exercises 及版本 1。管理員修改名稱、器材、說明、圖片或難度時，先鎖 exercises，再建立下一版 exercise_versions；舊版本不 UPDATE、不 DELETE。DELETE /api/admin/exercises/{id} 定義為停用 is_active=FALSE，非物理刪除。',
159:'3. 同一 exercise_id 代表同一動作身份；本版主部位固定，修正主部位或動作身份須新增 exercise_id 並停用舊身份。名稱不同不代表不同動作，名稱不設跨版本 UNIQUE。版本號由鎖內 MAX(version_no)+1 配置，UNIQUE(exercise_id,version_no) 防重複。',
161:'例外處理：非管理員回 403；同一身份改主部位回 409；未核准圖片、無效欄位回 422。舊課表保留原版本，新產生課表使用最新版本。',
178:'2. 依 3.11 固定規則，查詢啟用動作的最新 exercise_versions；只選難度不高於使用者等級的版本，配置每日順序、組次與休息。',
179:'3. 依 4.4 鎖定並重新驗證輸入與候選版本；同一交易新增 workout_plans、planning_profiles、planning_profile_body_parts、plan_days、plan_exercises，驗證後設定 published_at。plan_exercises 僅保存 exercise_version_id，不重複保存 exercise_id、名稱或圖片。任何步驟失敗全部回滾。',
183:'FR-11 查看複製修訂與封存課表',
184:'觸發條件：進入「我的課表」、複製修訂或確認封存。介面不將封存標示為永久刪除。',
185:'輸入資料：來源 workout_plan_id；POST /api/plans/{id}/revisions 接受 plan_name 與 exercises 陣列，各項含來源 plan_exercise_id 及欲替換的 exercise_version_id、target_sets、target_reps、rest_seconds。不得混入其他課表項目。',
187:'1. 讀取時 JOIN plan_days、plan_exercises 與其固定 exercise_versions，依 day_number、exercise_order 排序；不改連最新版本。',
188:'2. 修訂時完整複製來源週課表與其規劃輸入，再套用修改、檢查難度、主要部位覆蓋與時間；新 workout_plan_id 的 copied_from_plan_id 指向來源。原始課表不變。本次是人工修訂，algorithm_version 記錄 manual-revision-v1。',
189:'3. 已發布課表的名稱、規劃輸入、訓練日及動作不可原地修改。封存只填 archived_at，保留來源關聯；封存後禁止新開始，既有進行中場次仍可繼續。進行中若發現目標重量或組數不合適，使用者仍可照實記錄實際 weight_kg 與 reps，但不得改寫來源處方；要更改後續處方時，須先取消本場再複製修訂，或完成本場後修訂下一次課表。介面必須提示此差異。舊課表與動作版本不得物理刪除。',
190:'輸出／回應：GET /api/plans、GET /api/plans/{id} → 200；POST /api/plans/{id}/revisions → 201 與新 ID；POST /api/plans/{id}/archive → 204，可重送。原 PATCH/DELETE 課表端點不再提供，回 405。',
191:'例外處理：非本人或不存在回 404；無效動作版本、同日相同動作身份、超時回 422。封存不刪歷史；改用停用動作的新修訂不得發布。',
192:'複製修訂不增加／刪除訓練日與動作列，也不改順序；調整整體日數須依目前會員設定重新產課。新課表發布前允許交易內編輯，COMMIT 時不得留下未發布或不完整課表。',
199:'1. READ COMMITTED 交易先鎖 users；鎖成功後以另一個 SELECT 查目前 in_progress。若已有同一 plan_day_id，回原場次與已存組数；另一日回 409。只有沒有進行中場次才鎖來源課表並驗證所有權、published_at、archived_at 及動作啟用狀態。',
200:'2. 同一交易 INSERT workout_sessions（user_id、started_at、status=in_progress）及 workout_session_plans（workout_session_id、plan_day_id）。完整歷史處方由來源日連回不可變週課表及固定動作版本取得，不再逐組複製名稱。',
201:'3. uq_active_workout_per_user 防止同一會員提交兩筆 in_progress。遇此索引的 23505 必須先 ROLLBACK，再用新交易查既有場次；同日回 200、不同日回 409。其他唯一性錯誤不得誤判為續練。來源與完成組數的 DB 約束見 4.5。',
203:'例外處理：來源非本人或不存在回 404；另一日進行中回 409 WORKOUT_ALREADY_ACTIVE；新開始的來源已封存或動作停用回 409 PLAN_REQUIRES_UPDATE。既有同日場次續練不受事後封存影響。',
207:'輸入資料：workout_session_id 取自網址；JSON 傳 plan_exercise_id、set_number 1–100、weight_kg 0–999.99、reps 1–999。不得另傳 user_id、exercise_id、名稱或圖片。首版不接受課表外臨時動作；若日後支援，須新增獨立的場次動作來源模型，不得以放寬 FK 取代。',
209:'1. 依固定順序鎖 users 與本人 workout_sessions，確認 in_progress；驗證 plan_exercise_id 所屬 plan_day_id 等於場次來源日。後端先檢查，資料庫觸發器再驗證，防止跨課表日組數。',
210:'2. 以 (workout_session_id, plan_exercise_id, set_number) 唯一鍵保存；同鍵同值重送回既有結果，不多建一組。同鍵不同值回 409，需使用 FR-14 修改。',
219:'1. 經 workout_sessions 驗證本人並先鎖父場次；允許修正 in_progress 或 completed 的 weight_kg、reps，不改所屬場次、plan_exercise_id 或 set_number。',
220:'2. cancelled 不可改組數。completed 刪除最後一組回 409；資料庫在提交前也檢查非空，不能只靠前端與 API。需刪全部時使用刪除整場；本版不保存歷次數值修訂的稽核軌跡。',
229:'1. 完成時先鎖本人場次並確認至少一組；將 in_progress 改為 completed 並填 ended_at。延遲約束觸發器於 COMMIT 前確認仍有一組，否則回滾。重送完成回既有結果，不改結束時間。',
240:'1. Workout History 預設列出本人 completed 與 cancelled。workout_session_details 以 INNER JOIN 取得非空 source_plan_day_id，從 started_at 換算台北 session_date，再依日期與 ID 倒序。',
241:'2. 明細 JOIN workout_sets → plan_exercises → exercise_versions 顯示當時名稱、部位、器材、步驟、圖片、目標組次及實際組數；另由來源日回查整份週課表及規劃輸入。Exercise History 以 exercise_versions.exercise_id 篩選，跨版本視為同一動作。',
251:'3. 部位排名依固定版本的 body_part_id；動作排名依穩定 exercise_id 彙總所有版本。兩者以完成組數 COUNT(*) 排名、ID 升序破同分。排行榜名稱取期間內最近一次完成紀錄所用版本，並標註「最近訓練時名稱」，不 JOIN 目前名稱。AVG(weight_kg) 用於體重平均。',
269:'訓練頁按不可变來源課表順序顯示當時動作、部位、器材、版本圖片與說明、目標組次、休息與輸入欄位。今天指使用者選擇的課表日，不依星期猜測；計時僅為操作提示。歷史完整性範圍與不保存項目見 4.6。',
360:'NFR-05：密碼以 PHP bcrypt cost=10 儲存。Access JWT 固定 HS256、效期 2 小時且不含 role，只存於前端記憶體；Refresh Token 為 256-bit 不透明值，效期 7 天，只放在 Secure、HttpOnly、SameSite=Strict Cookie，資料庫只存雜湊並在每次刷新時輪替。每次受保護請求查登入工作階段，管理端點另查 admins。撤銷生效與重用偵測見 4.7、TC-05、TC-18。',
364:'NFR-08：登入限流以 PostgreSQL 的 login_rate_limits 作為整個部署共用的原子計數器，不存於單一 PHP Worker。來源 IP 先以環境變數 RATE_LIMIT_SECRET 做 HMAC-SHA-256；15 分鐘內允許 20 次失敗，第 21 次回 429。計數以 INSERT ... ON CONFLICT 的原子 UPSERT 更新，成功登入刪除該列；DB 不可用時登入採 fail closed 並回 503。依 4.9 與 TC-27 驗收。',
371:'NFR-13：Phinx migration 必須重建 19 張應用資料表、23 條 FK、索引、觸發器、唯讀 view 及 seed；另有工具表 phinxlog。composer.json 與 composer.lock 固定 PHP 套件版本。',
372:'NFR-14：每日備份，保留最近 7 份；同時備份不可覆寫圖片資源及 SHA-256 清單，單備份 URL 不算完成。至少一次還原到獨立測試庫，核對表数、筆數、FK、版本及圖片雜湊；RPO≤24 小時、RTO 目標≤60 分鐘。',
}
for i,t in changes.items():pchange(i,t.replace('组','組').replace('变','變').replace('数','數'))
# Correct existing minimum height to the user's current FR; don't silently restore old edits.
for t in D.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                if '15 張核心' in p.text:P.change(p,p.text.replace('15 張核心','19 張應用'))

# Load old dictionaries as preservation baseline before replacing chapter 4.
old={}
for hi,ti in [(287,288),(290,291),(293,294),(296,297),(299,300),(302,303),(305,306),(308,309),(311,312),(314,315),(318,319),(321,322),(324,325),(327,328),(330,331)]:
    name=Paragraph(original[hi],D).text.split()[-1]
    old[name]=[[c.text for c in row.cells] for row in Table(original[ti],D).rows][1:]

# Strip old chapter 4 as an internally inconsistent unit, retaining preceding/following chapters.
for el in original[272:353]:D.element.body.remove(el)
a=Paragraph(original[353],D)
ins(a,'4. 資料模型與結構設計','Heading 1',True)
ins(a,'本版以 18 張領域資料表保存目前狀態、不可變處方與實際訓練，另以 login_rate_limits 保存跨 PHP Worker 共用的登入限流狀態，共 19 張應用資料表。exercise_versions 保存動作版本；planning_profiles 與 planning_profile_body_parts 保存產課時輸入。歷史經固定外鍵取得，不使用目前會員設定或最新動作版本代替當時資料。1 個唯讀 view workout_session_details 不另存數值。')
ins(a,'本章定義待實作的資料庫契約。隔離、觸發器及效能均須依第 6 章測試後才可宣稱通過。資料表增加與欄位替換須用 migration，並非僅修改 Word 就已更新網站。')
ins(a,'4.1 實體關係圖','Heading 2')
ins(a,'ERD 為鍵層級模型，PK／FK／UQ 分別表示主鍵、外鍵與唯一鍵。端點數字代表可參與筆數；圖示以基本鍵約束的基數為準，發布完整性及每場恰一來源由 4.5 的提交前約束補強。各圖重複節點是同一資料表，不是資料複製。R01 至 R23 與 4.3 節逐列 1:1 對應；login_rate_limits 沒有 FK，因此不編列 R 編號。')

# Four small deterministic ERDs; each edge has a matching FK in the dictionary.
F=lambda n:ImageFont.truetype('/Library/Fonts/Arial Unicode.ttf',n)
def diagram(name,nodes,edges):
    im=Image.new('RGB',(1900,1750),'white');dr=ImageDraw.Draw(im)
    def edge_label(xy,text,size=27):
        xx,yy=xy
        box=dr.textbbox((xx,yy),text,font=F(size))
        dr.rectangle((box[0]-5,box[1]-3,box[2]+5,box[3]+3),fill='white')
        dr.text((xx,yy),text,font=F(size),fill='black')
    for src,dst,ca,cb,route,rid in edges:
        dr.line(route,fill='black',width=3)
        for pt,other,label in [(route[0],route[1],ca),(route[-1],route[-2],cb)]:
            xx,yy=pt;ox,oy=other
            tx=xx+12 if ox>=xx else xx-dr.textlength(label,font=F(27))-12
            ty=yy+12 if oy>yy else yy-43
            edge_label((tx,ty),label)
        segments=list(zip(route,route[1:]))
        (p1,p2)=max(segments,key=lambda s:abs(s[1][0]-s[0][0])+abs(s[1][1]-s[0][1]))
        mx=(p1[0]+p2[0])//2;my=(p1[1]+p2[1])//2
        edge_label((mx-25,my-38 if p1[1]==p2[1] else my-14),rid,25)
    for n,(xx,yy,fields) in nodes.items():
        w=650;hh=90+len(fields)*49
        dr.rectangle((xx,yy,xx+w,yy+hh),outline='black',width=3)
        dr.line((xx,yy+70,xx+w,yy+70),fill='black',width=2)
        dr.text((xx+16,yy+15),n,font=F(32),fill='black')
        for j,t in enumerate(fields):dr.text((xx+16,yy+87+j*49),t,font=F(29),fill='black')
    dr.text((45,1635),'PK 主鍵　FK 外鍵　UQ 唯一鍵　1 一筆　0..1 至多一筆　0..N 零至多筆',font=F(28),fill='black')
    dr.text((45,1690),'教師建議修訂版　2026-10-06　完整欄位與跨表規則見 4.2–4.5',font=F(28),fill='black')
    path=TMP/(name+'.png');im.save(path)
    return path
# Keep the source-backed account diagram and add the standalone rate-limit table.
acct=D.part.related_parts[Paragraph(original[276],D)._p.xpath('.//a:blip')[0].get(qn('r:embed'))].blob
(TMP/'accounts.png').write_bytes(acct)
acct_im=Image.open(TMP/'accounts.png').convert('RGB');acct_dr=ImageDraw.Draw(acct_im)
acct_dr.rectangle((35,35,555,305),fill='white',outline='#66717A',width=3)
acct_dr.rectangle((35,35,555,105),fill='#E9ECEF',outline='#66717A',width=2)
acct_dr.text((50,49),'login_rate_limits',font=F(34),fill='#111111')
for j,t in enumerate(['PK ip_hash','window_started_at / failed_count','blocked_until / updated_at']):
    acct_dr.text((51,126+j*50),t,font=F(30),fill='#111111')
acct_dr.text((36,320),'獨立安全維運表，沒有 FK；故不編列 R 關係。',font=F(24),fill='#333333')
acct_path=TMP/'accounts_with_rate_limit.png';acct_im.save(acct_path)
pic(a,acct_path,'圖 2　帳號、目前偏好與登入限流 ERD')
def figure(nodes,edges,name,caption):
    ins(a,'',page=True);pic(a,diagram(name,nodes,edges),caption)
figure({
'users':(50,40,['PK user_id']), 'workout_plans':(1150,40,['PK workout_plan_id','FK user_id','FK copied_from_plan_id (NULL)']),
'training_goals':(50,550,['PK training_goal_id']), 'planning_profiles':(1150,550,['PK FK workout_plan_id','FK training_goal_id','experience / days / duration']),
'body_parts':(50,1130,['PK body_part_id']), 'planning_profile_body_parts':(1150,1130,['PK FK workout_plan_id','PK FK body_part_id','UQ (workout_plan_id, priority)']),
},[('users','workout_plans','1','0..N',[(700,150),(1150,150)],'R08'),('training_goals','planning_profiles','1','0..N',[(700,660),(1150,660)],'R09'),('workout_plans','planning_profiles','1','0..1',[(1475,277),(1475,550)],'R10'),('planning_profiles','planning_profile_body_parts','1','0..N',[(1475,787),(1475,1130)],'R11'),('body_parts','planning_profile_body_parts','1','0..N',[(700,1240),(1150,1240)],'R12'),('workout_plans','workout_plans','0..1','0..N',[(1800,170),(1850,170),(1850,400),(1700,400),(1700,277)],'R23')], 'inputs','圖 3　課表與產課輸入 ERD')
figure({
'workout_plans':(50,40,['PK workout_plan_id','published_at / archived_at']), 'exercises':(1150,40,['PK exercise_id','is_active']),
'plan_days':(50,560,['PK plan_day_id','FK workout_plan_id','day_number / day_title']), 'exercise_versions':(1150,560,['PK exercise_version_id','FK exercise_id / FK body_part_id','UQ (exercise_id, version_no)','name / description / image']),
'plan_exercises':(50,1150,['PK plan_exercise_id','FK plan_day_id','FK exercise_version_id','order / sets / reps / rest']), 'body_parts':(1150,1150,['PK body_part_id']),
},[('workout_plans','plan_days','1','0..N',[(375,228),(375,560)],'R13'),('plan_days','plan_exercises','1','0..N',[(375,797),(375,1150)],'R16'),('exercises','exercise_versions','1','0..N',[(1475,228),(1475,560)],'R15'),('exercise_versions','plan_exercises','1','0..N',[(1150,760),(950,760),(950,1300),(700,1300)],'R17'),('body_parts','exercise_versions','1','0..N',[(1475,1150),(1475,846)],'R14')],'prescriptions','圖 4　不可變處方與動作版本 ERD')
figure({
'users':(50,40,['PK user_id']), 'plan_days':(1150,40,['PK plan_day_id','FK workout_plan_id']),
'workout_sessions':(50,560,['PK workout_session_id','FK user_id','status / started_at / ended_at']), 'workout_session_plans':(1150,560,['PK FK workout_session_id','FK plan_day_id']),
'workout_sets':(50,1150,['PK workout_set_id','FK workout_session_id','FK plan_exercise_id','set_number / weight_kg / reps']), 'plan_exercises':(1150,1150,['PK plan_exercise_id','FK plan_day_id','FK exercise_version_id']),
},[('users','workout_sessions','1','0..N',[(375,179),(375,560)],'R18'),('workout_sessions','workout_session_plans','1','0..1',[(700,685),(1150,685)],'R19'),('plan_days','workout_session_plans','1','0..N',[(1475,228),(1475,560)],'R20'),('workout_sessions','workout_sets','1','0..N',[(375,797),(375,1150)],'R21'),('plan_exercises','workout_sets','1','0..N',[(1150,1310),(700,1310)],'R22')],'tracking','圖 5　訓練來源與實際組數 ERD')
ins(a,'圖 3 的課表自關聯 copied_from_plan_id 只記修訂來源；原課表不會被覆寫。圖 5 的 workout_sets 不再保存 exercise_id 或名稱，版本由 plan_exercise_id 唯一決定。場次來源日與組數所屬課表日必須一致，由 DB 觸發器驗證。')

# Dictionary rows: compact 3-column layout, no stale columns from old design.
h(a,'4.2 資料表詳細規格')
ins(a,'除標示可空者外一律 NOT NULL。主鍵 INTEGER 使用 GENERATED BY DEFAULT AS IDENTITY；时间戳為 TIMESTAMPTZ，created_at 預設 CURRENT_TIMESTAMP。未註明 DEFAULT 的欄位須由後端提供。FK 刪除政策見 4.3；CHAR(64) 雜湊須 CHECK 為 64 個小寫十六進位字元。')
schemas={}
for name,rows in old.items():
    schemas[name]=[(r[0],r[1].replace('\n',' ')+'；'+r[2]+('；可空' if r[3]=='是' else '；必填'),r[5]+(' DEFAULT '+r[4].replace('\n','') if r[4] not in ('無','IDENTITY') else '')) for r in rows]
schemas['users']=[(c,t,s.replace('100 AND 250','90 AND 250')) for c,t,s in schemas['users']]
schemas['users'].insert(-1,('profile_completed_at','TIMESTAMPTZ；可空','完成基本資料且初始體重同交易建立後寫入；有值時 body_records 必至少一筆。'))
schemas['user_sessions']=[
('user_session_id','UUID；PK','作為 Access JWT 的 jti；DEFAULT gen_random_uuid()。'),
('auth_id','INTEGER；FK','→ user_auth；CASCADE。'),
('refresh_token_hash','CHAR(64)','只存 256-bit 不透明 Refresh Token 的 SHA-256，不保存 Cookie 原值。'),
('created_at','TIMESTAMPTZ','DEFAULT CURRENT_TIMESTAMP。'),
('expires_at','TIMESTAMPTZ','Access JWT 到期時間；CHECK expires_at > created_at，登入或刷新後為目前時間 + 2 小時。'),
('refresh_expires_at','TIMESTAMPTZ','Refresh Token 絕對到期時間；CHECK refresh_expires_at > expires_at，登入後為目前時間 + 7 天。'),
('revoked_at','TIMESTAMPTZ；可空','登出、改密碼、停用帳號或 Refresh Token 重用時寫入；非 NULL 即失效。'),
('last_refreshed_at','TIMESTAMPTZ；可空','成功輪替 Refresh Token 的時間；DEFAULT NULL。')]
schemas['body_records']=[(c,('NUMERIC(6,2)；—；必填' if c=='weight_kg' else t),s) for c,t,s in schemas['body_records']]
schemas['exercises']=[('exercise_id','INTEGER；PK','動作穩定身份；IDENTITY。'),('is_active','BOOLEAN','DEFAULT TRUE；FALSE 排除新課表／新開始，不抹除既有歷史。'),('created_at','TIMESTAMPTZ','DEFAULT CURRENT_TIMESTAMP。')]
schemas['exercise_versions']=[('exercise_version_id','INTEGER；PK','IDENTITY；完整描述版本的唯一識別。'),('exercise_id','INTEGER；FK','→ exercises；RESTRICT。'),('version_no','INTEGER','CHECK >=1；UNIQUE(exercise_id,version_no)。'),('exercise_name','VARCHAR(120)','去空白 1–120 字；不同版本可同名，不以名稱當身份。'),('body_part_id','INTEGER；FK','→ body_parts；RESTRICT；同一身份各版本必須相同。'),('difficulty_level','VARCHAR(20)','CHECK IN Beginner / Intermediate / Advanced。'),('equipment','VARCHAR(100)；可空','徒手填 Bodyweight；缺資料才 NULL。'),('description','TEXT','長度 1–4,000，包含步驟與注意事項；不可覆寫。'),('image_url','TEXT；可空','本站不可覆寫資源 URL，最長 2,048 字。'),('image_sha256','CHAR(64)；可空','與 image_url 同時為 NULL 或同時有值；識別圖片內容。'),('created_at','TIMESTAMPTZ','DEFAULT CURRENT_TIMESTAMP；版本建立時間。')]
schemas['workout_plans']=[('workout_plan_id','INTEGER；PK','IDENTITY；每次產課／修訂取得新 ID。'),('user_id','INTEGER；FK','→ users；RESTRICT。'),('copied_from_plan_id','INTEGER；FK；可空','→ workout_plans；RESTRICT；原始產課 NULL，同會員複製來源。'),('plan_name','VARCHAR(120)','去空白 1–120 字；發布後不可改。'),('created_at','TIMESTAMPTZ','DEFAULT CURRENT_TIMESTAMP。'),('published_at','TIMESTAMPTZ；可空','僅組裝交易中允許 NULL；提交前必填且 >=created_at。'),('archived_at','TIMESTAMPTZ；可空','DEFAULT NULL；有值須 >=published_at；只能 NULL→時間，不可解封存。')]
schemas['planning_profiles']=[('workout_plan_id','INTEGER；PK/FK','→ workout_plans；RESTRICT；每課表一份產課時輸入。'),('training_goal_id','INTEGER；FK','→ training_goals；RESTRICT；目標名稱／代碼首版固定。'),('age_at_planning','SMALLINT','CHECK BETWEEN 13 AND 100。'),('height_cm_at_planning','NUMERIC(5,2)','CHECK BETWEEN 90 AND 250。'),('weight_kg_at_planning','NUMERIC(6,2)','CHECK BETWEEN 20 AND 500；輸入時體重值，非目前體重快取。'),('weight_record_date','DATE','該體重來源日期；由交易檢查不得晚於 captured_at 的台北日期。'),('experience_at_planning','VARCHAR(20)','CHECK IN Beginner / Intermediate / Advanced。'),('requested_days','SMALLINT','CHECK BETWEEN 2 AND 6；產課請求天數，發布時須等於日表筆數。'),('duration_minutes','SMALLINT','CHECK IN (30,60,90)；本次處方時間上限。'),('algorithm_version','VARCHAR(60)','rule-v1 或 manual-revision-v1；長度 1–60；規則代碼保留於 Git。'),('captured_at','TIMESTAMPTZ','取得規劃輸入的時間；複製修訂沿用原輸入時間。')]
schemas['planning_profile_body_parts']=[('workout_plan_id','INTEGER；PK/FK','→ planning_profiles；RESTRICT。'),('body_part_id','INTEGER；PK/FK','→ body_parts；RESTRICT；複合 PK(workout_plan_id,body_part_id)。'),('priority','SMALLINT','CHECK BETWEEN 1 AND 8；UNIQUE(workout_plan_id,priority)；發布驗證連續 1…N。')]
schemas['plan_days']=[(c,t,s.replace('CASCADE','RESTRICT')) for c,t,s in schemas['plan_days']]
schemas['plan_exercises']=[(c,t,s.replace('CASCADE','RESTRICT')) for c,t,s in schemas['plan_exercises'] if c!='exercise_id']
schemas['plan_exercises'].insert(2,('exercise_version_id','INTEGER；FK','→ exercise_versions；RESTRICT；固定版本，不連最新值。'))
schemas['workout_sessions']=[(c,t,s.replace('CASCADE','RESTRICT')) for c,t,s in schemas['workout_sessions']]
schemas['workout_sets']=[(c,t,s) for c,t,s in schemas['workout_sets'] if c!='exercise_id']
schemas['workout_sets'].insert(2,('plan_exercise_id','INTEGER；FK','→ plan_exercises；RESTRICT；須屬於本場來源日。'))
schemas['workout_session_plans']=[('workout_session_id','INTEGER；PK/FK','→ workout_sessions；CASCADE；每場恰一來源由延遲觸發器保證。'),('plan_day_id','INTEGER；FK','→ plan_days；RESTRICT；不可 UPDATE／單獨 DELETE。')]
schemas['login_rate_limits']=[
('ip_hash','CHAR(64)；PK','正規化來源 IP 經 HMAC-SHA-256 後的雜湊；金鑰由 RATE_LIMIT_SECRET 環境變數提供，不保存原始 IP。'),
('window_started_at','TIMESTAMPTZ','目前 15 分鐘計數視窗的起點。'),
('failed_count','SMALLINT','DEFAULT 0；CHECK BETWEEN 0 AND 32767。'),
('blocked_until','TIMESTAMPTZ；可空','第 21 次失敗前設為視窗結束時間；NULL 表示目前未封鎖。'),
('updated_at','TIMESTAMPTZ','DEFAULT CURRENT_TIMESTAMP；每次失敗、重設或解除時更新。')]
order=['users','user_auth','user_sessions','admins','login_rate_limits','training_goals','body_parts','user_body_parts','body_records','exercises','exercise_versions','workout_plans','planning_profiles','planning_profile_body_parts','plan_days','plan_exercises','workout_sessions','workout_session_plans','workout_sets']
for n,name in enumerate(order,1):
    ins(a,f'表 4.2.{n}　{name}','Heading 3')
    tab(a,['欄位','型態與鍵','語意、預設值及約束'],schemas[name],[3.6,3.7,9.4],9.0)
ins(a,'workout_session_details 為 INNER JOIN workout_sessions 與 workout_session_plans 的唯讀 view，提供來源日與台北 session_date。沒有來源的損毀資料不可被默默解讀為正常歷史；一致性巡檢須另外以 LEFT JOIN 找出並報錯。')

h(a,'4.3 外鍵關係與刪除政策')
rels=[
('R01','users','user_auth','user_id','1 : 0..1','CASCADE'),('R02','user_auth','user_sessions','auth_id','1 : 0..N','CASCADE'),('R03','users','admins','user_id','1 : 0..1','CASCADE'),('R04','training_goals','users','training_goal_id','0..1 : 0..N','RESTRICT'),('R05','users','user_body_parts','user_id','1 : 0..N','CASCADE'),('R06','body_parts','user_body_parts','body_part_id','1 : 0..N','RESTRICT'),('R07','users','body_records','user_id','1 : 0..N','CASCADE'),('R08','users','workout_plans','user_id','1 : 0..N','RESTRICT'),('R09','training_goals','planning_profiles','training_goal_id','1 : 0..N','RESTRICT'),('R10','workout_plans','planning_profiles','workout_plan_id','1 : 0..1','RESTRICT'),('R11','planning_profiles','planning_profile_body_parts','workout_plan_id','1 : 0..N','RESTRICT'),('R12','body_parts','planning_profile_body_parts','body_part_id','1 : 0..N','RESTRICT'),('R13','workout_plans','plan_days','workout_plan_id','1 : 0..N','RESTRICT'),('R14','body_parts','exercises','body_part_id','1 : 0..N','RESTRICT'),('R15','exercises','exercise_versions','exercise_id','1 : 0..N','RESTRICT'),('R16','plan_days','plan_exercises','plan_day_id','1 : 0..N','RESTRICT'),('R17','exercise_versions','plan_exercises','exercise_version_id','1 : 0..N','RESTRICT'),('R18','users','workout_sessions','user_id','1 : 0..N','RESTRICT'),('R19','workout_sessions','workout_session_plans','workout_session_id','1 : 0..1','CASCADE'),('R20','plan_days','workout_session_plans','plan_day_id','1 : 0..N','RESTRICT'),('R21','workout_sessions','workout_sets','workout_session_id','1 : 0..N','CASCADE'),('R22','plan_exercises','workout_sets','plan_exercise_id','1 : 0..N','RESTRICT'),('R23','workout_plans','workout_plans','copied_from_plan_id','0..1 : 0..N','RESTRICT')]
tab(a,['編號','父表 → 子表','子表外鍵','父端 : 子端／刪除'],[(rid,u+' → '+v,k,c+'；'+d) for rid,u,v,k,c,d in rels],[1.3,5.5,4.2,5.7],9.0)
ins(a,'共 23 條 FK。發布交易須讓每份課表有且僅有一份 planning_profiles、1–8 個偏好、2–6 個訓練日及每日 3–6 動作。每場訓練在 COMMIT 時必有一筆來源；completed 必至少一組。圖中的 0 下限不是允許正常業務提交空內容，而是區分單純 FK 與跨表約束能力。')
ins(a,'課表與版本不物理刪除：課表封存、動作停用，相關 RESTRICT 避免歷史斷鏈。帳號停用保留資料。使用者明確刪除 completed/cancelled 整場時，場次及來源關聯、實際組數一併刪除；該操作不刪課表。歷史保存承諾不包含已由使用者確認刪除的場次。正式資料保留／個資清除屬獨立維運流程，不以停用或 CASCADE 冒充。')
tab(a,['唯一性或檢查','規則'],[
('帳號','UNIQUE(user_auth.user_id)、UNIQUE(email)；CHECK email=lower(btrim(email))。'),
('目前偏好','PK(user_id,body_part_id)、UNIQUE(user_id,priority)；順位連續由更新交易驗證。'),
('輸入偏好','PK(workout_plan_id,body_part_id)、UNIQUE(workout_plan_id,priority)。'),
('版本與處方','UNIQUE(exercise_id,version_no)；UNIQUE(workout_plan_id,day_number)；UNIQUE(plan_day_id,exercise_order)。'),
('同日動作身份','同日不可使用相同 exercise_id 的兩個版本；跨表判斷由發布／子表觸發器保證，不以版本 ID 唯一冒充身份唯一。'),
('每組與體重','UNIQUE(workout_session_id,plan_exercise_id,set_number)；UNIQUE(user_id,record_date)。'),
('訓練狀態','CHECK status IN (in_progress,completed,cancelled)；in_progress 的 ended_at 為 NULL；終态的 ended_at 非 NULL 且 >=started_at。'),
],[5,11.7])

h(a,'4.4 隔離等級與併發交易協定')
ins(a,'本版所有業務寫入明確採 READ COMMITTED。PHP 在 beginTransaction() 後先執行 SET TRANSACTION ISOLATION LEVEL READ COMMITTED，驗收以 SHOW transaction_isolation 確認，不依賴伺服器預設。每一 SQL 指令取新的已提交快照；SELECT FOR UPDATE 僅鎖既有資料列，不鎖「查不到的空集合」，也不等同 SERIALIZABLE。[1]')
ins(a,'目前資料、偏好及體重寫入先鎖 users；課表建立、封存、開始及訓練修改採 users → workout_plans（需要時）→ workout_sessions（需要時）→ exercises（需要時，ID 升序）→ 子表。DB 觸發器補做所需父列鎖並重新查詢。所有流程不得反向取得資源鎖；管理員修改動作不再取得會員資源鎖。')
ins(a,'只在短交易內做 DB 讀寫；圖片下載、HTTP 呼叫、等待使用者操作與昂貴雜湊均不放在持鎖區間。SET LOCAL lock_timeout=\'3s\'、statement_timeout=\'5s\' 作為本版上限，部署驗收確認與效能門檻相容。')
ins(a,'4.4.1 開始訓練的交易步驟','Heading 3')
code(a,"""BEGIN;
SET TRANSACTION ISOLATION LEVEL READ COMMITTED;
SELECT user_id FROM users
WHERE user_id = :authenticated_user_id FOR UPDATE;
-- 必須是鎖成功後的另一條 SQL，不能合併成同一個舊快照查詢。
SELECT workout_session_id FROM workout_sessions
WHERE user_id = :authenticated_user_id
  AND status = 'in_progress';
-- 已有：相同來源日回既有 ID；不同日回 409。
-- 沒有：鎖來源課表，再驗證本人、已發布、未封存與動作啟用。
-- 同一交易建立場次與來源；COMMIT 前執行延遲完整性檢查。
COMMIT;""")
ins(a,'SQL 中的 :authenticated_user_id 由已驗證的登入狀態取得，不採用客戶端任意 user_id。BEGIN 例示資料庫交易；PHP 實作以 PDO beginTransaction() 啟動，不重複 BEGIN。')
code(a,"""CREATE UNIQUE INDEX uq_active_workout_per_user
ON workout_sessions (user_id)
WHERE status = 'in_progress';""")
tab(a,['時間','Request A','Request B'],[
('T1','取得 users(U) 的 FOR UPDATE。','嘗試鎖同一 users(U)，等待。'),
('T2','另查 active，為空；新增場次 S 與來源。','尚未取得鎖，不得依先前空結果 INSERT。'),
('T3','COMMIT，釋放鎖。','取得鎖，另查看到已提交的 S。'),
('T4','回 201 與 S。','同來源日回 200 與 S；不同日回 409，筆數不增加。'),
],[1.5,7.6,7.6])
ins(a,'若 A 回滾，B 取得鎖後可建立自己的場次。即使錯誤寫入路徑省略 users 鎖，部分唯一索引仍阻止兩筆 active 同時提交：衝突 INSERT 可能等待對方結束，對方提交後失敗。[2] 只把 constraint/index 名稱為 uq_active_workout_per_user 的 SQLSTATE 23505 映射為續練衝突。')
ins(a,'4.4.2 錯誤分類及重試界線','Heading 3')
tab(a,['情況','處理要求'],[
('23505 active 唯一性','ROLLBACK 後新交易查 active；同日 200、不同日 409。若查無（例如場次已完成），回 409 請重新載入，不盲目再建。'),
('40P01 死結／40001 序列化失敗','整筆回滾；本版不在控制器內自動無限重試，回 409 TRANSACTION_RETRY_REQUIRED。若日後自動重試，必須重跑整筆且有次數上限。'),
('55P03 鎖逾時／57014 指令逾時','整筆回滾，回 503 RETRY_LATER；前端保留輸入。'),
('提交結果不明','不宣稱失敗亦不立即重建；先查 current 場次或該組唯一鍵。active 索引不是所有時間的 idempotency key，完成後的新開始是新場次。'),
('其他 DB 錯誤','回滾，记录不含機密的 request_id／SQLSTATE；回一般錯誤，不冒充已有場次。'),
],[5,11.7])

h(a,'4.5 資料庫完整性與觸發器')
ins(a,'「completed 至少一組」不是資料庫做不到，而是一般單列 CHECK 與 FK 無法表達父列必須有子列。CHECK 不得以查詢其他表的函式假裝跨表約束；FK 只保證子列的父列存在。[2] 本版採後端預檢加資料庫約束觸發器，正常應用／維運 DML 均受保護。')
ins(a,'4.5.1 完成訓練與最後一組','Heading 3')
ins(a,'workout_sets 的 BEFORE INSERT/UPDATE/DELETE 觸發器鎖定父 workout_sessions；取得鎖後重新讀狀態。INSERT 只允許 in_progress；UPDATE 只許重量與次數，禁止換場次、換 plan_exercise_id、換組號；cancelled 不允許逐組改刪。completed 的 DELETE 最終是否違規由提交前檢查決定。')
ins(a,'父場次 INSERT／狀態 UPDATE，以及組數 INSERT／UPDATE／DELETE，各自登記 AFTER ROW CONSTRAINT TRIGGER，設定 DEFERRABLE INITIALLY DEFERRED。檢查於交易結束時讀取最終狀態：若父場次仍存在且 status=completed，必須 EXISTS 至少一筆 workout_sets；不符拋出命名約束錯誤，COMMIT 失敗並回滾。[3]')
code(a,"""-- 約束觸發器的檢查邏輯；需分別掛在父表與子表事件。
affected_session := NEW/OLD.workout_session_id;
SELECT status FROM workout_sessions
WHERE workout_session_id = affected_session FOR UPDATE;
IF parent_not_found THEN
    RETURN; -- 僅允許已合法刪整場所導致的 CASCADE
END IF;
IF status = 'completed' AND NOT EXISTS (
    SELECT 1 FROM workout_sets
    WHERE workout_session_id = affected_session
) THEN
    RAISE check_violation 'completed_session_requires_set';
END IF;""")
ins(a,'上述為規格偽碼，不是可直接執行的 migration。實作須區分 TG_OP 的 NEW/OLD，並測試 INSERT completed、改為 completed、刪最後一組及刪整場。BEFORE 觸發器序列化同父場次的改動；延遲觸發器檢查提交後應成立的不變量，兩者不能互相取代。')
ins(a,'兩個交易各刪一組時，即使原有兩組，也必先競爭同一父列鎖。第一筆刪除提交後剩一組，第二筆取得鎖後刪除將在提交前被拒絕。完成與存組同時發生時，先存組者提交後完成可成功；先完成者若無組則失敗，若已有組則後續新增因終態而被拒。')
ins(a,'整場刪除限定 completed/cancelled。父列刪除後 CASCADE 清除 sets 與來源，子表觸發器發現父列在本交易已不存在時跳過非空檢查；禁止直接刪來源關聯。角色不得取得 TRUNCATE、停用 trigger 或改 session_replication_role 的權限。擁有者／超級使用者能繞過部分保護，屬受管制維運邊界，不宣稱可防所有特權操作。')
ins(a,'4.5.2 來源與不可變版本','Heading 3')
tab(a,['觸發時機','保證內容'],[
('workout_sessions／workout_session_plans 提交前','每個存在的場次恰一來源；來源課表 user_id 等於場次 user_id。開始時必已發布未封存；事後封存不使既有場次失效。'),
('workout_sets BEFORE 寫入','plan_exercise_id 的 plan_day_id 等於本場來源日；禁止引用另一日或其他會員處方。'),
('課表及子表 BEFORE 變動','先鎖 workout_plans。published_at 非 NULL 後禁止名稱、copied_from、輸入、日與動作的 INSERT/UPDATE/DELETE；只允许首次 archive。'),
('workout_plans 提交前','published_at 非 NULL；一份輸入、1–8 偏好順位連續、日數=requested_days、每日 3–6 動作，順序連續、同日 exercise_id 不重複。'),
('exercise_versions BEFORE 變動','只准 INSERT，不准 UPDATE/DELETE；INSERT 先鎖 exercises，版本號連續且同身份主部位不變。不在版本列複製可變 is_active。'),
('workout_session_plans BEFORE 變動','只可開始時 INSERT；禁止 UPDATE、父場次仍存在時 DELETE。'),
],[5.5,11.2])
ins(a,'資料庫觸發器保證結構及時間不變量；產課的選擇策略與預估時間由具版本的後端規則驗證，不能把有 FK 等同於授權或健身合理性。所有測試須確認失敗交易的欄位與筆數均不變。')

ins(a,'4.5.3 基本資料完成與最低一筆體重','Heading 3')
ins(a,'首次完成基本資料時，後端在同一交易鎖定 users、UPSERT 當日 body_records，最後設定 users.profile_completed_at。只要 profile_completed_at 非 NULL，該會員就必須保有至少一筆體重紀錄。這項規則避免刪除全部體重後，目前體重變成 NULL，造成產課與介面欄位失去必要輸入。')
ins(a,'body_records 的 BEFORE DELETE／user_id UPDATE 觸發器先以 FOR UPDATE 鎖定父 users，使同一會員的並行刪除依序執行。users 的 profile_completed_at 變動，以及 body_records 的 DELETE 或 user_id UPDATE，另登記 DEFERRABLE INITIALLY DEFERRED 約束觸發器。提交前若會員仍存在、profile_completed_at 非 NULL 且查無 body_records，交易以 23514 拒絕。刪除完整帳號時，父 users 已不存在，CASCADE 清理不受此限制。')
code(a,"""-- 提交前不變量偽碼
IF EXISTS (
    SELECT 1 FROM users AS u
    WHERE u.user_id = affected_user_id
      AND u.profile_completed_at IS NOT NULL
) AND NOT EXISTS (
    SELECT 1 FROM body_records AS br
    WHERE br.user_id = affected_user_id
) THEN
    RAISE check_violation 'completed_profile_requires_weight_record';
END IF;""")
ins(a,'刪除最後一筆體重時 API 回 409 LAST_BODY_RECORD_REQUIRED，介面改為提供「修改體重」。產課交易仍會鎖 users 並讀取最新體重；資料庫觸發器負責最低一筆的結構保障，後端負責把可理解的錯誤訊息回傳給使用者。')

h(a,'4.6 正規化與完整歷史保存')
ins(a,'4.6.1 保存範圍與時間語意','Heading 3')
ins(a,'本版的「完整歷史」指能還原當次採用的整週課表、來源訓練日、動作版本、目標處方、產課輸入及現存實際組數，不表示備份整個網站或追蹤所有操作。開始訓練只綁定已發布的不可變課表；不同場次可共享同一份不可變處方，不必每次再複製整份。')
tab(a,['歷史項目','唯一保存位置','修改後的歷史行为'],[
('課表名稱／整週安排','workout_plans、plan_days','發布後不可變；修訂產生新 ID，來源仍指舊 ID。'),
('目標、程度、天數、時間、年齡、身高、體重與來源日期','planning_profiles','為產課當時值；不隨 users 或 body_records 更新。姓名、Email、性別不是 rule-v1 輸入，不複製進快照。'),
('偏好部位及優先順序','planning_profile_body_parts','每個部位一列，非 VARCHAR 清單；不跟目前偏好同步。'),
('動作名稱、器材、难度、步驟、圖片','exercise_versions','固定版本不可覆寫；images 實體資源依 SHA-256 保存。'),
('動作顺序、目標組次、休息','plan_exercises','引用單一固定版本；不改連最新版本。'),
('場次開始／結束、狀態、實際重量／次數','workout_sessions、workout_sets','實際量與目標量是不同事實。允許本人修正錯記的數值，歷史顯示修正後值。'),
],[4.4,5,7.3])
ins(a,'不保存：未提交的輸入、已確認刪除的場次、每次修改實際重量前的舊值、整個候選目錄在當時的所有版本集合、瀏覽器像素畫面及登入憑證。因此本版可完整查閱處方與現存訓練事實，但不是全面稽核日誌，也不宣稱僅憑 algorithm_version 就能重新執行得到完全相同的候選搜尋結果。若課程要求數值修改稽核，須另增 append-only audit 表，不能稱目前已具備。')
ins(a,'圖片 URL 並非圖片快照。發布版本前，部署流程須把合法素材存於不可覆寫路徑（例如 /assets/exercises/<sha256>.png），並在版本保存 URL 與雜湊；備份包括圖片實體。外部網址被覆寫或失效不能影響舊版本。沒有圖片時明確記 NULL，顯示非此動作示範的佔位圖，不宣稱還原不存在的圖片。')
ins(a,'三年前名稱為「槓鈴臥推」的版本 V1，改名後新增 V2「平板槓鈴臥推」。舊 plan_exercises 仍指向 V1，因此三年前歷史顯示 V1 的名稱、器材、說明與圖片；新課表才指向 V2。exercise_id 不變，可彙總同一動作跨版本進步。')
ins(a,'4.6.2 刻意重複與第三正規化','Heading 3')
ins(a,'正規化的目標是排除錯誤的函數相依與更新異常，不是禁止不同時間事實出現相同數值。本版不在每組重複保存名稱、圖片與部位，也不保存可計算的 volume、總組數、session_date 或目前體重。名稱與圖片在 exercise_version_id 下只保存一次。')
tab(a,['關係與候選鍵','主要函數相依','判定與維護規則'],[
('exercise_versions；PK version_id，AK(exercise_id,version_no)','版本键 → 名稱、部位、難度、器材、說明、圖片','同一身份跨版本主部位固定，故 exercise_id → body_part_id 形成非鍵相依；此欄刻意保留於版本以自足顯示，屬受控反正規化，不宣稱嚴格 3NF。DB 保證同身份部位一致。'),
('planning_profiles；PK workout_plan_id','課表 ID → 當次輸入與捕捉時間','獨立的時間事實；不是 user_id → 目前設定的重複快取，與目前值不同是預期。'),
('planning_profile_body_parts；PK(plan_id,body_part_id)，AK(plan_id,priority)','完整複合鍵 → priority','消除多值欄位；非鍵不依賴部分 PK。'),
('plan_exercises；PK plan_exercise_id，AK(day_id,order)','處方動作 ID → 版本、目標組次與休息','不另存可由版本推導的 exercise_id／名稱。'),
('workout_sets；PK set_id，AK(session_id,plan_exercise_id,set_number)','組的鍵 → weight_kg、reps、created_at','實際數值只存每組；volume 查詢計算。'),
('workout_sessions 與來源關聯','session_id → 執行者、來源日（在各自關係）','user_id 可沿自己的來源課表推導，屬跨關係的受控資訊重複；保留以建立單會員 active 唯一索引並直接授權，DB 驗證兩側相同。'),
],[5,5.2,6.5])
ins(a,'因此報告不能籠統宣稱「所有表均嚴格 3NF 且零重複」。上述版本部位欄位及場次執行者重複必須列入例外。本版改採更嚴格方案：將固定 body_part_id 留在 exercises，不存於 exercise_versions；歷史透過不可改的 exercise 身份讀取固定部位。下列資料字典與 ERD 以此消除非鍵相依的方案為準，場次執行者仍為有完整性約束的跨表重複。')
# Apply strict 3NF refinement to model in a postpass below; no duplicate body part in versions.
ins(a,'課表複製修訂會重用既有 exercise_versions，但複製該次的日與處方列；這是以不同 workout_plan_id 識別的兩個不可變處方事實。重複值增加儲存量，交換獨立可重建、無回寫同步的歷史。此取捨與列內函數相依分開評估。')
ins(a,'所有 immutable 表的 UPDATE/DELETE 權限應以最小權限及觸發器雙重限制；對可變表只給必要欄位更新權。任何版本修訂都先建立新列再引用，禁止背景程序「同步修正」舊版本。版本與圖片保留期間須涵蓋引用它們的所有課表與場次。')

h(a,'4.7 JWT 身分與資料庫授權')
ins(a,'JWT 只攜帶 sub、jti、iat、exp，不保存 role。Token 簽章證明內容未被改動，不能保證管理員資格至今未變。本版本來就須查 user_sessions 以支援登出撤銷，因此採有伺服器狀態的 JWT 設計，不宣稱完全無狀態。')
ins(a,'每次受保護請求驗證 Token 與登入狀態；一般資源再驗所有權，管理端點再查 admins。GET /api/me 的 role 只是即時查詢結果供顯示。即使前端仍顯示管理按鈕，後端也以當次授權為準；DB 不可用或查詢失敗時不放行。[4]')
ins(a,'管理寫入先以 FOR SHARE 鎖 user_auth、對應 user_sessions 及 admins 資格列，再修改資源並提交。撤銷管理員資格的維運程序依相同順序鎖 user_auth 後刪 admins；兩者互斥。若撤銷先提交，之後管理寫入回 403；若管理寫入先取得資格鎖，該筆可先提交，撤銷等待其結束。本版不承諾撤銷能取消已通過授權並持鎖的交易。讀取端點在授權查詢時作判斷。')

h(a,'4.8 SQL 查詢與索引')
ins(a,'歷史查詢一定走處方固定版本，不使用目前目錄的 MAX(version_no)。下例保留 JOIN、GROUP BY 與 SUM；:user_id 來自伺服器認證。')
code(a,"""SELECT
    ws.workout_session_id,
    ev.exercise_version_id,
    ev.exercise_name,
    SUM(s.weight_kg * s.reps) AS total_volume
FROM workout_sessions AS ws
JOIN workout_sets AS s
    ON s.workout_session_id = ws.workout_session_id
JOIN plan_exercises AS pe
    ON pe.plan_exercise_id = s.plan_exercise_id
JOIN exercise_versions AS ev
    ON ev.exercise_version_id = pe.exercise_version_id
WHERE ws.user_id = :user_id
  AND ws.status = 'completed'
GROUP BY ws.workout_session_id,
         ev.exercise_version_id, ev.exercise_name
ORDER BY ws.workout_session_id DESC;""")
ins(a,'每週次數由場次表直接 COUNT(*)，或組表 JOIN 後 COUNT(DISTINCT workout_session_id)，避免一場多組被重複計數。每日／每週日期均從 started_at 換算 Asia/Taipei；日期篩選以 UTC 半開區間 [起日,迄日+1日) 套用到 started_at。')
tab(a,['索引','用途與界線'],[
('uq_active_workout_per_user','UNIQUE(user_id) WHERE status=\'in_progress\'；DB 最終防重複。'),
('idx_sessions_user_started','workout_sessions(user_id,started_at DESC)；會員歷史日期篩選。'),
('每組複合 UNIQUE','(workout_session_id,plan_exercise_id,set_number)；重送去重並涵蓋父場次 FK 前綴。'),
('idx_sets_plan_exercise','workout_sets(plan_exercise_id)；處方引用與動作歷史 JOIN。'),
('版本複合 UNIQUE','exercise_versions(exercise_id,version_no)；每個動作取最新版及版本唯一。'),
('idx_exercises_body_active','exercises(body_part_id) WHERE is_active=TRUE；部位候選。完整 body_part_id FK 索引另行評估。'),
('idx_plan_exercises_version','plan_exercises(exercise_version_id)；版本引用。'),
('idx_session_plans_day','workout_session_plans(plan_day_id)；來源查詢，刪除為 RESTRICT。'),
('idx_sessions_cleanup','user_sessions(refresh_expires_at)；另對 revoked_at IS NOT NULL 建部分索引，支援排程分批清理。'),
('idx_login_rate_updated','login_rate_limits(updated_at)；支援過期限流列清理。'),
('其餘 FK','對未被既有 PK／UNIQUE 左前綴涵蓋的 FK 建 B-tree，尤其 copied_from_plan_id、planning_profiles.training_goal_id、兩個偏好表的 body_part_id。'),
],[5.7,11])
ins(a,'既有 PK／UNIQUE 已有索引，不重複建立相同前綴。以 EXPLAIN (ANALYZE, BUFFERS) 驗證讀取計畫及資料量，不能由「已建索引」推論必然變快。migration 與備份帳號和應用帳號分離；Phinx 版本與種子需可重跑。')

h(a,'4.9 登入限流與過期資料清理')
ins(a,'登入失敗計數放在 PostgreSQL 的 login_rate_limits，使多個 PHP Worker 或 Container 共用同一份狀態。PHP 以伺服器端 RATE_LIMIT_SECRET 對正規化 IP 做 HMAC-SHA-256，不保存原始 IP。單一交易使用 INSERT ... ON CONFLICT DO UPDATE 原子更新視窗與 failed_count；成功登入刪除對應列。第 20 次失敗提交後設定 blocked_until，第 21 次請求在驗證密碼前回 429；封鎖到期後重設新視窗。')
code(a,"""INSERT INTO login_rate_limits (
    ip_hash, window_started_at, failed_count, blocked_until, updated_at
) VALUES (:ip_hash, CURRENT_TIMESTAMP, 1, NULL, CURRENT_TIMESTAMP)
ON CONFLICT (ip_hash) DO UPDATE
SET failed_count = CASE
        WHEN login_rate_limits.window_started_at <= CURRENT_TIMESTAMP - INTERVAL '15 minutes'
        THEN 1 ELSE login_rate_limits.failed_count + 1 END,
    window_started_at = CASE
        WHEN login_rate_limits.window_started_at <= CURRENT_TIMESTAMP - INTERVAL '15 minutes'
        THEN CURRENT_TIMESTAMP ELSE login_rate_limits.window_started_at END,
    blocked_until = CASE
        WHEN login_rate_limits.window_started_at > CURRENT_TIMESTAMP - INTERVAL '15 minutes'
         AND login_rate_limits.failed_count + 1 >= 20
        THEN login_rate_limits.window_started_at + INTERVAL '15 minutes'
        ELSE NULL END,
    updated_at = CURRENT_TIMESTAMP
RETURNING failed_count, blocked_until;""")
ins(a,'Render Cron Job 每日執行 php bin/cleanup-expired.php。腳本以 1,000 筆為一批，使用短交易清除 refresh_expires_at 或 revoked_at 已超過 7 天的 user_sessions，並清除 updated_at 超過 1 天且 blocked_until 已到期的 login_rate_limits。Access Token 的 expires_at 已到期但 Refresh Token 尚有效時不可刪除該登入工作階段。每批提交後再處理下一批，避免長交易與大量鎖定。')
code(a,"""WITH doomed AS (
    SELECT user_session_id
    FROM user_sessions
    WHERE refresh_expires_at < CURRENT_TIMESTAMP - INTERVAL '7 days'
       OR revoked_at < CURRENT_TIMESTAMP - INTERVAL '7 days'
    ORDER BY COALESCE(revoked_at, refresh_expires_at)
    LIMIT 1000
    FOR UPDATE SKIP LOCKED
)
DELETE FROM user_sessions AS us
USING doomed
WHERE us.user_session_id = doomed.user_session_id;""")
ins(a,'排程記錄開始時間、刪除筆數、執行秒數與 request_id，但不記錄 Cookie、Token、雜湊或原始 IP。排程失敗不得影響使用者既有資料；監控連續失敗並由維運人員重跑。TC-27 驗證跨 Worker 限流，TC-31 驗證分批清理不會刪除仍可刷新的工作階段。')

# Add the missing referenced algorithm at the end of chapter 3, not a dangling appendix.
anchor4=next(p for p in D.paragraphs if p.text=='4. 資料模型與結構設計')
h(anchor4,'3.11 規則式排課與版本記錄')
ins(anchor4,'rule-v1 將 Beginner／Intermediate／Advanced 排為 1／2／3，只選不高於會員程度的啟用動作最新版本。先安排八個部位週覆蓋，再優先增加使用者所選部位頻率；每日不重複同一 exercise_id。候選依難度升序、exercise_id 升序決定，沒有隨機抽樣。')
ins(anchor4,'2／3／4／5／6 天的部位模板分別採：① 胸背肩腿／二頭三頭臀核心；② 胸肩三頭／背二頭核心／腿臀；③ 胸肩三頭／背二頭／腿臀核心／優先部位補強；④ 在前三天外加上肢補強、下肢核心；⑤ 推／拉／腿核心各重複兩次。每日先保證 3 動作，依可用時間最多 6 動作；補強日按偏好順位循環，且仍受每日時間限制。')
ins(anchor4,'增肌 3×10、減脂 3×12、力量 4×5、一般健身 2×12；休息分別 60／45／120／45 秒。預估秒數 = 300 秒暖身 + Σ[sets×reps×3 + (sets−1)×rest] + (動作數−1)×60。只能在不破壞覆蓋與最低 3 動作的前提下選 3–6 動作；無可行配置回 409，不能默默超過時間。')
ins(anchor4,'年齡、身高、體重保存為產課上下文，本版不以其推定安全負重；實際重量由使用者輸入。候選版本、設定及偏好在持鎖交易內再次確認；保存輸入快照與完整輸出处方，published_at 之前驗證全部條件。人工修訂沿用原規劃上下文，改天數或改目標需重新產課。')

# Replace obsolete acceptance table with concrete tests, covering all FR plus new rules.
ta=Table(original[379],D);ta._tbl.getparent().remove(ta._tbl)
at=Paragraph(original[380],D)
tests=[
('TC-01','FR-01 註冊','新 Email','註冊，查 users/auth/sessions 與回應標頭。','201；三表各 +1；Access JWT 無 role、exp−iat=7200；設定 Secure、HttpOnly、SameSite=Strict Refresh Cookie；無 admins。'),
('TC-02','FR-01 重複帳號','已有小寫 Email','以不同大小寫再註冊。','409 EMAIL_EXISTS；各表筆數不變。'),
('TC-03','FR-05／06 設定','會員已登入','存 4 天 60 分偏好胸背；再送 7 天或重複部位。','合法 200；非法 422；舊設定保留。'),
('TC-04','FR-10 產課','初階增肌4天60分，各部位足夠候選','產課後 JOIN 輸入、日、動作版本。','201；4天、每日3–6動作、全部位覆蓋且不超時；完整輸入與published_at。'),
('TC-05','FR-02 登入刷新登出','有效帳號','登入→模擬 F5 清除記憶體→refresh→以舊 Refresh 重用→登出→以舊 Access 呼叫。','200；刷新後取得新 Access 且 Refresh 雜湊輪替；重用舊 Refresh 撤銷該工作階段並回 401；登出清 Cookie；舊 Access 401。'),
('TC-06','FR-03 密碼與停用','兩個登入階段','改密碼驗舊 Token；重新登入後停用。','所有舊 Token 401；停用後不能登入，訓練仍在 DB。'),
('TC-07','FR-04 基本資料','已有會員','存身高90及250；送89.99；同日改體重；檢查 profile_completed_at。','邊界可存，非法422；同日體重僅1筆；完成時間與初始體重同交易提交，users 無體重副本。'),
('TC-08','FR-07 動作搜尋','同動作有V1,V2','目錄搜尋與V1舊課表查詢。','目錄顯V2；舊課表顯V1；空搜尋items=[]。'),
('TC-09','FR-08 版本修改','已引用V1','改名、說明、圖片建立V2；直接SQL UPDATE V1。','V2新增；舊值不變；UPDATE被DB拒絕；不改exercise_id。'),
('TC-10','FR-09 參考資料','管理員／一般會員','管理員改description；會員嘗試；嘗試改分類名。','200／403／422；固定身份與名稱不變。'),
('TC-11','FR-11 修訂封存','已發布P1及進行中場次','嘗試改 P1 目標；實際以不同重量完成；取消後複製修訂 P2；封存 P1。','來源處方不可改並顯示取消／完成後修訂提示；實際值可照實保存；P2 為新 ID 且 P1 不變；封存後不得新開始。'),
('TC-12','FR-12 同日併發','會員無active；两連線與同步屏障','同時POST同日，確認第二連線在users鎖等待。','僅1場與1來源；一個201、一個200；ID相同。'),
('TC-13','FR-12 異日併發','同上，兩個不同來源日','同時POST異日；另以測試SQL繞過users鎖競爭active索引。','API一201一409；直接SQL至多1筆提交，另一23505。'),
('TC-14','FR-13 組數重送','in_progress','同鍵同值送2次，再同鍵不同值；傳另一日或課表外 plan_exercise_id。','201／200／409；只1組。另一日與課表外動作在 API 及 DB 皆被拒絕，v1 不建立臨時動作。'),
('TC-15','FR-14 組數修正','completed恰1組','改重量→刪該組→SQL換所屬場次。','修改200且量改變；刪組409；換場次被DB拒絕。'),
('TC-16','FR-15 DB非空約束','空場次；非superuser測試角色','直接SQL把空場次設completed並COMMIT；再以INSERT建空completed。','兩筆交易均在提交前失敗；無空completed留存，不能只看API回應。'),
('TC-17','FR-15 並行刪最後組','completed有2組','兩連線各刪不同一組，第一筆提交後第二筆再提交。','至多一筆刪除提交；另一約束失敗；最終≥1組。'),
('TC-18','管理員撤銷','有效JWT及admins資格','先撤銷提交再寫管理API；反向先持授權鎖再撤銷。','前者403；後者撤銷等待既有寫入結束，後續403；一般會員API不需admins。'),
('TC-19','FR-16 完整歷史','P1使用V1，有已完成場次','改會員設定、刪原體重、發布V2、修訂P2、封存P1。','仍查得P1整週、當初輸入與偏好順序、V1名稱說明圖片、當初目標組次；實際值不變。'),
('TC-20','FR-17 統計','完成組60×12、65×10、65×8及徒手0×15；另有取消組','查外加重量量、徒手次數、週次數與跨V1/V2同身份排名。','外加重量訓練量1890 kg·reps；bodyweight_reps=15、completed_sets=4；1場；取消排除；同exercise_id合併。'),
('TC-21','FR-18 體重 CRUD','兩日各1筆且 Profile 已完成','同日新增、改值、刪最新，再刪最後一筆。','重複409；改值成功；刪最新後目前體重回退；刪最後一筆回409 LAST_BODY_RECORD_REQUIRED；已發布輸入快照不變。'),
('TC-22','隔離／死結／逾時','兩測試連線','SHOW隔離；故意互鎖造死結，再持鎖超3秒。','read committed；回滾後409或503；不得部分寫入或假成功。'),
('TC-23','原子性／來源','無active','開始時來源INSERT前注入錯誤；直接刪既有來源後COMMIT。','前者全回滾；後者DB拒絕，原來源仍在。'),
('TC-24','封存／開始競爭','已發布未封存P','兩連線同時封存P及開始P。','按鎖順序：開始先成功則既有場次保留；封存先提交則開始409；無斷鏈。'),
('TC-25','完成／存組競爭','場次已有1組','並行完成及新增；重送完成或取消。','存組先則完成可見；完成先則新增409；同終態重送200時間不變，換終態409。'),
('TC-26','整場刪除','completed或cancelled','DELETE場次，查組、來源、課表、版本。','204；場次/組/來源刪除，課表/版本保留；in_progress刪除409。'),
('TC-27','安全與限流','兩PHP worker及一般會員','同一 ip_hash 15分鐘21次失敗；成功登入後重查；模擬DB不可用。','login_rate_limits 原子計數跨 Worker 一致；第21次429；成功後刪列；DB不可用時登入503且不繞過限流。'),
('TC-28','圖片與備份','已引用V1圖片','還原DB和圖片到隔離環境；核對sha256；模擬目前圖更新。','V1圖片雜湊一致；舊URL內容不被覆寫；缺圖有文字，不假稱圖片已保存。'),
('TC-29','效能／響應式','第5章資料量及環境','200人5分鐘；另測3種視窗及指定瀏覽器。','依NFR實測記錄P95、錯誤率及截圖；未達不得標示通過。'),
('TC-30','發布／不可變約束','已發布P；可寫入測試角色','直接加日、換版本、改profile、改published_at；提交空draft。','全部DB拒絕；合法產課原子發布；同日同exercise不同版本也拒絕。'),
('TC-31','Session 與限流清理','準備有效Refresh、過期與撤銷超過7日Session及過期限流列','執行 cleanup-expired.php；檢查每批上限、刪除範圍與日誌。','仍可刷新的Session保留；過期／撤銷超過7日者刪除；過期限流列刪除；每批≤1000且無Token、雜湊或原始IP進入日誌。'),
]
tab(at,['案例','功能','前置條件','操作步驟','預期結果'],tests,[1.6,3.5,4.1,7.4,8.8],9.5)
# The section ending after tests is landscape; other user content stays portrait.
sec=original[380].xpath('.//w:sectPr')[0];sz=sec.find(qn('w:pgSz'));sz.set(qn('w:w'),'16838');sz.set(qn('w:h'),'11906');sz.set(qn('w:orient'),'landscape')

# Refresh functional trace references, eliminating old dangling TC numbers.
mapping={1:'01、02',2:'05',3:'06',4:'07',5:'03',6:'03',7:'08',8:'09',9:'10',10:'04、23、30',11:'11、24',12:'12、13、24',13:'14',14:'15、17',15:'16、17、25、26',16:'19',17:'20',18:'21'}
current=None
for p in D.paragraphs:
    m=re.match(r'FR-(\d+) ',p.text)
    if m:current=int(m[1])
    if p.text.startswith('資料庫對應：') and current:
        P.change(p,'驗收對應：'+ '、'.join('TC-'+n for n in mapping[current].split('、'))+'；相關資料表與 FK 見第 4 章。')

# References are explanatory sources, not proof of implementation.
ref=next(p for p in D.paragraphs if p.text=='7. 參考資料')
for s in [
'[1] PostgreSQL 16 Transaction Isolation：https://www.postgresql.org/docs/16/transaction-iso.html',
'[2] PostgreSQL 16 Constraints：https://www.postgresql.org/docs/16/ddl-constraints.html',
'[3] PostgreSQL 16 CREATE TRIGGER：https://www.postgresql.org/docs/16/sql-createtrigger.html',
'[4] OWASP Authorization Cheat Sheet：https://cheatsheetseries.owasp.org/cheatsheets/Authorization_Cheat_Sheet.html',
]:D.add_paragraph(s)

# Apply teacher-feedback refinements to the functional, security, and UX contracts.
for p in D.paragraphs:
    t=p.text
    if t.startswith('前端網頁與後端 API 使用相同的網站來源。'):
        P.change(p,'前端網頁與後端 API 使用相同的網站來源。受保護 API 以 Authorization: Bearer <Access JWT> 驗證身分；只有 /api/auth/refresh 與 /api/auth/logout 會讀取同源的 HttpOnly Refresh Cookie。系統不開放任意外部網站跨來源讀取 API，Refresh 請求另驗證 Origin。所有新增、修改及刪除操作均由後端重新檢查資料格式、身分與權限，不依賴前端驗證結果。')
    elif t.startswith('輸出／回應：POST /api/auth/register'):
        P.change(p,'輸出／回應：POST /api/auth/register → 201；JSON 回傳 user_id、access_token、token_type=Bearer、expires_in=7200，回應同時設定 __Host-refresh_token Cookie。Refresh Token 不出現在 JSON、URL、前端儲存區或日誌。')
    elif t.startswith('1. PHP 依正規化 Email 查詢 user_auth'):
        P.change(p,'1. PHP 依正規化 Email 查詢 user_auth，以 password_verify() 比對密碼並檢查 is_active。成功登入時建立 user_sessions：Access JWT 效期 2 小時，Refresh Token 絕對效期 7 天。資料庫只存 Refresh Token 雜湊。')
    elif t.startswith('3. 登出時將對應工作階段'):
        P.change(p,'3. POST /api/auth/refresh 以 Cookie 原值雜湊查找工作階段，驗證 Origin、未撤銷與 refresh_expires_at 後，在同一交易輪替 Refresh Token 雜湊、更新 expires_at 與 last_refreshed_at，再發新 Access JWT。舊 Refresh Token 重用時撤銷整個工作階段。登出時寫入 revoked_at 並以過期 Set-Cookie 清除 Cookie。')
    elif t.startswith('輸出／回應：POST /api/auth/login'):
        P.change(p,'輸出／回應：POST /api/auth/login → 200 與 Access JWT，並設定 Refresh Cookie；POST /api/auth/refresh → 200 與新 Access JWT、輪替 Cookie；POST /api/auth/logout → 204 並清除 Cookie。')
    elif t.startswith('例外處理：錯誤帳密統一回 401'):
        P.change(p,'例外處理：錯誤帳密統一回 401 INVALID_CREDENTIALS；停用帳號、過期、撤銷或重用 Refresh Token 回 401。登入限流依 NFR-08 回 429；限流儲存不可用時回 503，不得繞過。')
    elif t.startswith('2. PUT /api/me/profile 驗證所有必填值'):
        P.change(p,'2. PUT /api/me/profile 驗證所有必填值，鎖定 users，在同一交易更新 users、UPSERT 今日 body_records，最後設定 profile_completed_at。只有 COMMIT 成功才視為基本資料完成。')
    elif t.startswith('3. 目前體重由 body_records'):
        P.change(p,'3. 目前體重由 body_records 依 record_date DESC、body_record_id DESC 取得最新一筆；users 不另存 weight_kg。profile_completed_at 非 NULL 的會員在資料庫層必須至少保有一筆體重紀錄，因此產課與基本資料頁不會取得 NULL。')
    elif t.startswith('例外處理：缺少必填資料、範圍或等級錯誤'):
        P.change(p,'例外處理：缺少必填資料、範圍或等級錯誤回 422；不得修改其他使用者。初始體重寫入失敗時整筆回滾，profile_completed_at 不得單獨留下。')
    elif t.startswith('1. 只計入本人 status=completed 的場次'):
        P.change(p,'1. 只計入本人 status=completed 的場次。weighted_volume（介面顯示「外加重量訓練量」）= SUM(weight_kg×reps)；weight_kg=0 的徒手組不估算身體負荷，但仍納入 completed_sets、total_reps 與 bodyweight_reps。')
    elif t.startswith('2. 總量使用 SUM(weight_kg*reps)'):
        P.change(p,'2. 外加重量訓練量使用 SUM(weight_kg*reps)，每週次數使用 COUNT(DISTINCT workout_session_id)，另以 COUNT(*)、SUM(reps) 及 FILTER (WHERE weight_kg=0) 計算完成組數、總次數與徒手次數；日期以台北日期及週一分組。')
    elif t.startswith('3. 最常部位及動作以完成組數'):
        P.change(p,'3. 最常部位及動作以完成組數 COUNT(*) 排名，平手依對應 ID 升序；體重歷史讀取 body_records，平均值用 AVG(weight_kg)。未來若要估算徒手有效負荷，須建立動作別係數、版本及驗證依據；首版不以任意比例混入 kg·reps。')
    elif t.startswith('4. 無訓練資料回總量 0'):
        P.change(p,'4. 無訓練資料回 weighted_volume=0、completed_sets=0、total_reps=0、bodyweight_reps=0、每週次數 0 與空排行榜。完成基本資料後體重至少一筆；尚未完成 Profile 時體重趨勢為空並顯示引導，不讓前端直接對 NULL 做計算。')
    elif t.startswith('輸出／回應：GET /api/dashboard'):
        P.change(p,'輸出／回應：GET /api/dashboard → 200；回傳 weighted_volume、completed_sets、total_reps、bodyweight_reps、每週次數、部位與動作排行、體重趨勢。')
    elif t.startswith('2. 查詢依 record_date 排序；修改與刪除均驗證本人。'):
        P.change(p,'2. 查詢依 record_date 排序；修改與刪除均驗證本人。刪除最新紀錄後，目前體重回退到上一筆；若該筆已是 Profile 完成後的最後一筆，資料庫與 API 均拒絕刪除，使用者只能修改該筆或先新增另一日紀錄。')
    elif t.startswith('例外處理：重複日期回 409；未來日期'):
        P.change(p,'例外處理：重複日期回 409；刪除最後一筆回 409 LAST_BODY_RECORD_REQUIRED；未來日期或不合法體重回 422；非本人回 404。')
    elif t.startswith('介面必須區分「尚未儲存」'):
        P.change(p,'介面必須區分「尚未儲存」「儲存中」「已儲存」「儲存失敗」。只有成功回應才標示已儲存；失敗保留輸入並允許重試。F5 後 Access Token 雖從記憶體消失，前端先呼叫 /api/auth/refresh，成功即恢復登入，不要求重新輸入帳密；Refresh 失敗才顯示登入頁。已提交的每組資料不受影響，尚未送出的欄位仍只存在目前頁面，離開前須提示。')
    elif t.startswith('除標示可空者外一律 NOT NULL。主鍵 INTEGER'):
        P.change(p,'除標示可空者外一律 NOT NULL。主鍵 INTEGER 使用 GENERATED BY DEFAULT AS IDENTITY；時間戳為 TIMESTAMPTZ，created_at 預設 CURRENT_TIMESTAMP。身高欄位使用 NUMERIC(5,2) 並限制 90–250 cm；所有 kg 欄位（體重快照、體重紀錄及訓練重量）統一使用 NUMERIC(6,2)，再以各欄 CHECK 限制合理範圍。未註明 DEFAULT 的欄位須由後端提供。FK 刪除政策見 4.3；CHAR(64) 雜湊須 CHECK 為 64 個小寫十六進位字元。')
    elif t.startswith('JWT 只攜帶 sub、jti、iat、exp'):
        P.change(p,'Access JWT 只攜帶 sub、jti、iat、exp，不保存 role。Refresh Token 是隨機不透明值，只由 HttpOnly Cookie 傳送，資料庫只存雜湊；兩者共用 user_session_id。JWT 簽章只能證明內容未被改動，不能保證管理員資格至今未變。本版查 user_sessions 以支援登出、密碼變更、Refresh 輪替與重用撤銷，因此不宣稱完全無狀態。')

# Clarify the database-side source guard and the v1 ad-hoc exercise decision.
for t in D.tables:
    for row in t.rows:
        if row.cells and row.cells[0].text=='workout_sets BEFORE 寫入':
            row.cells[1].text='plan_exercise_id 的 plan_day_id 必須等於本場來源日；禁止引用另一日、其他會員處方或任何課表外動作。首版不支援臨時加做；後端先檢查，DB 觸發器再鎖父場次並驗證。'

# Freeze-part strict 3NF final model: body part belongs to stable exercise identity.
for p in D.paragraphs:
    if p.text.startswith('因此報告不能籠統'):
        P.change(p,'本版的資料表以 3NF 為設計目標：exercise_id → body_part_id，因此固定主部位只存 exercises，不在每個版本重複；exercise_versions 只保存會隨版本改變的描述。歷史分類經固定 exercise_id 取得，不受目前名稱更動影響。場次 user_id 的跨表重複仍須明列用途與 DB 一致性檢查，不能宣稱完全沒有重複。')
    if '固定版本的 body_part_id' in p.text:P.change(p,p.text.replace('固定版本的 body_part_id','固定動作身份的 body_part_id'))
    if p.text.startswith('1. 以參數化查詢取得 exercises'):
        P.change(p,p.text.replace('JOIN body_parts','依 exercises.body_part_id JOIN body_parts'))
    if p.text.startswith('歷史經固定'):pass
for t in D.tables:
    for row in t.rows:
        vals=[c.text for c in row.cells]
        if vals and vals[0].startswith('exercise_versions；PK'):
            row.cells[2].text='exercise_id 不能決定跨版本描述；固定部位放 exercises，版本表不複製。候選鍵以外欄位依版本鍵，沒有當前名稱副本。'
            row.cells[1].text='版本鍵 → 名稱、難度、器材、說明、圖片'
        if vals and vals[0]=='body_parts → exercise_versions':row.cells[0].text='body_parts → exercises'
# Rewrite exercise/version tables in place.
for p in D.paragraphs:
    if p.text in ('表 4.2.9　exercises','表 4.2.10　exercise_versions'):
        t=Table(p._p.getnext(),D)
        if p.text.endswith('exercise_versions'):
            for row in list(t.rows)[1:]:
                if row.cells[0].text in ('body_part_id','image_url'):t._tbl.remove(row._tr)
                elif row.cells[0].text=='image_sha256':row.cells[2].text='圖片內容的 SHA-256；NULL 表示無圖。URL 由固定資源前綴＋雜湊＋.png 推導，不另存 image_url。'
        else:
            rr=t.add_row();rr.cells[0].text='body_part_id';rr.cells[1].text='INTEGER；FK';rr.cells[2].text='→ body_parts；RESTRICT；同一動作身份固定，禁止 UPDATE。'
# Eliminate asset URL/hash functional-dependency duplication and keep API vs DB distinct.
for p in D.paragraphs:
    if '版本保存 URL 與雜湊' in p.text:
        P.change(p,p.text.replace('版本保存 URL 與雜湊','版本只保存內容雜湊；URL 由固定資源前綴、雜湊與 .png 副檔名推導'))
    if p.text.startswith('1. 在管理寫入交易內確認'):
        P.change(p,'1. 在管理寫入交易內確認 admins 資格；image_url 只作 API 輸入，必須指向本站已核准的不可覆寫 PNG 路徑。部署程序驗證授權、保存圖片實體並計算 SHA-256，後端只存 image_sha256，不再存可由雜湊推導的 URL；不得下载任意用戶 URL。')
    if p.text.startswith('目前資料、偏好及體重寫入先鎖'):
        P.change(p,'授權鎖先於業務鎖：需要持鎖驗證的寫入先取得 user_auth → user_sessions → admins（僅管理操作），再進入 users → workout_plans（需要時）→ workout_sessions（需要時）→ exercises（需要時，ID 升序）→ 子表。目前資料、偏好及體重寫入先鎖 users，使產課快照讀取一致。DB 觸發器補做所需父列鎖並重新查詢；管理員修改動作不反向取得會員資源鎖。死結仍須處理，固定鎖序不是「永不死結」的保證。')
    if p.text.startswith('資料庫觸發器保證結構'):
        P.change(p,p.text+' copied_from_plan_id 只允許指向已發布且同會員的既有課表，來源在新課表之前建立，因此不得形成自循環或修訂循環。')
# Remove user-source margin comments; final delivery remains a clean editable file.
for e in list(D.element.xpath('.//w:commentRangeStart | .//w:commentRangeEnd | .//w:commentReference')):e.getparent().remove(e)
for rid,rel in list(D.part.rels.items()):
    if 'comment' in rel.reltype.lower():del D.part.rels[rid]
for e in D.settings.element.xpath('./w:revisionView'):D.settings.element.remove(e)
D.settings.element.append(x('revisionView',markup='0',comments='0',insDel='0',formatting='0'))

# Plain readable formatting; preserve existing user title, names and text.
for t in D.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
for p in D.paragraphs:
    if p.style.name=='Heading 1':p.paragraph_format.page_break_before=True
for e in list(D.element.xpath('.//w:shd')):e.getparent().remove(e)
# Patch references inside all text after strict schema finalization.
for e in D.element.xpath('.//w:t'):
    if e.text:e.text=e.text.replace('时间','時間').replace('允许','允許').replace('行为','行為').replace('难度','難度').replace('顺序','順序').replace('终态','終態').replace('记录','記錄').replace('两連線','兩連線')

# Architecture and diagram finalization via generated assets; never edit source images.
import revise_srs_rubric as R
R.ASSETS=TMP/'architecture_assets';R.ASSETS.mkdir(exist_ok=True)
# Recreate the architecture rather than retaining its old 15-table count.
im=Image.new('RGB',(1900,1100),'white');dr=ImageDraw.Draw(im)
for y,title,body in [(25,'使用者瀏覽器','PHP 範本 / HTML / CSS / JavaScript / Fetch API'),(295,'PHP 網頁與 API','Access JWT / HttpOnly Refresh Cookie / admins 授權 / 規則式排課'),(565,'PDO 資料存取層與排程','READ COMMITTED / 固定列鎖 / 參數化 SQL / 每日清理 Job'),(835,'雲端 PostgreSQL','19 張應用表 / 不可變版本 / FK / UNIQUE / Trigger')]:
    dr.rectangle((80,y,1820,y+175),outline='black',width=3)
    dr.text((115,y+20),title,font=F(42),fill='black');dr.text((115,y+95),body,font=F(35),fill='black')
    if y<835:dr.line((950,y+175,950,y+265),fill='black',width=4);dr.polygon([(950,y+269),(938,y+246),(962,y+246)],fill='black')
arch=TMP/'architecture.png';im.save(arch)
ap=Paragraph(original[63],D);ap.clear();ap.alignment=1;ap.add_run().add_picture(str(arch),width=Cm(16.5))
# Modify source diagram model to reflect fixed part on exercises, not on versions.
# Generate replacement prescription diagram with corrected relation, same source-backed keys.
fixed=diagram('prescriptions_final',{
'workout_plans':(50,40,['PK workout_plan_id','published_at / archived_at']), 'exercises':(1150,40,['PK exercise_id','FK body_part_id (immutable)','is_active']),
'plan_days':(50,560,['PK plan_day_id','FK workout_plan_id','day_number / day_title']), 'exercise_versions':(1150,560,['PK exercise_version_id','FK exercise_id','UQ (exercise_id, version_no)','name / description / image']),
'plan_exercises':(50,1150,['PK plan_exercise_id','FK plan_day_id','FK exercise_version_id','order / sets / reps / rest']), 'body_parts':(1150,1150,['PK body_part_id']),
},[('workout_plans','plan_days','1','0..N',[(375,228),(375,560)],'R13'),('plan_days','plan_exercises','1','0..N',[(375,797),(375,1150)],'R16'),('exercises','exercise_versions','1','0..N',[(1475,277),(1475,560)],'R15'),('exercise_versions','plan_exercises','1','0..N',[(1150,760),(950,760),(950,1300),(700,1300)],'R17'),('body_parts','exercises','1','0..N',[(1800,1240),(1860,1240),(1860,160),(1800,160)],'R14')])
cp=next(p for p in D.paragraphs if p.text=='圖 4　不可變處方與動作版本 ERD');dp=Paragraph(cp._p.getprevious(),D);dp.clear();dp.alignment=1;dp.add_run().add_picture(str(fixed),width=Cm(16.5))
# Give every durable figure a human-readable alternative description.
figure_alts=[
'系統架構：瀏覽器經 PHP 網頁與 API、PDO 資料存取層連線雲端 PostgreSQL，另有每日清理排程。',
'帳號、目前偏好與登入限流 ERD：顯示 users、認證、登入工作階段、管理員、目標、部位、體重及獨立限流表。',
'課表與產課輸入 ERD：顯示 R08 至 R12 與 R23。',
'不可變處方與動作版本 ERD：顯示 R13 至 R17。',
'訓練來源與實際組數 ERD：顯示 R18 至 R22，實際組數引用來源課表動作。']
assert len(D.inline_shapes)==len(figure_alts)
for shape,alt in zip(D.inline_shapes,figure_alts):shape._inline.docPr.set('descr',alt)
# Replace the complete old TOC region, including its cached hyperlink paragraphs.
for e in original[13:45]:
    if e.getparent() is not None:e.getparent().remove(e)
tp=Paragraph(original[46],D).insert_paragraph_before()
rr=tp.add_run()._r
rr.append(x('fldChar',fldCharType='begin'))
it=x('instrText');it.text=' TOC \\o "1-2" \\h ';rr.append(it)
rr.append(x('fldChar',fldCharType='end'))
P.regenerate_toc(D)
for i,section in enumerate(D.sections):
    section.footer.is_linked_to_previous=False
    section.different_first_page_header_footer=(i==0)
    fp=section.footer.paragraphs[0];fp.clear();fp.alignment=2
    fp.add_run('第 ')
    fr=fp.add_run()._r;fr.append(x('fldChar',fldCharType='begin'))
    fi=x('instrText');fi.text=' PAGE ';fr.append(fi)
    fr.append(x('fldChar',fldCharType='end'));fp.add_run(' 頁')
    fmt(fp,9)
for e in D.element.xpath('.//w:t'):
    if e.text:
        e.text=e.text.replace('下载','下載').replace('表数','表數')
        e.text=re.sub(r'(?i)\bpostgresql\b','PostgreSQL',e.text)
        e.text=re.sub(r'(?i)\bjwt\b','JWT',e.text)
        e.text=re.sub(r'(?i)\bjson\b','JSON',e.text)
        e.text=e.text.replace('application/JSON','application/json')
# Normalize Traditional Chinese glyph fallback across source text, TOC, tables,
# captions, code comments, headers and footers while preserving Latin/code fonts.
for style in D.styles:
    rpr=style.element.get_or_add_rPr();fonts=rpr.get_or_add_rFonts()
    for key in ('ascii','hAnsi','eastAsia','cs'):fonts.set(qn('w:'+key),'Arial Unicode MS')
for run in D.element.xpath('.//w:r'):
    rpr=run.get_or_add_rPr();fonts=rpr.get_or_add_rFonts()
    ascii_font=fonts.get(qn('w:ascii'))
    if ascii_font not in ('Courier New','Menlo','Consolas'):
        for key in ('ascii','hAnsi','eastAsia','cs'):fonts.set(qn('w:'+key),'Arial Unicode MS')
    else:
        fonts.set(qn('w:eastAsia'),'Arial Unicode MS')
OUT.parent.mkdir(parents=True,exist_ok=True);D.save(OUT)
embed_editable_font(OUT,'/System/Library/Fonts/Supplemental/Arial Unicode.ttf','Arial Unicode MS')
assert hashlib.sha256(SRC.read_bytes()).hexdigest()==digest
assert len(order)==19 and len(rels)==23
print(OUT)
print('source unchanged; 19 entities; 23 FKs; 31 acceptance cases')
