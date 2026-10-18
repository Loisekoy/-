"""Revise the current editable SRS for a PHP/PDO implementation, not the app.

The user-edited source document remains untouched. Render the output, then call
this script with --toc-pdf to cache the actual pages in its native Word TOC.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import io
import json
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader

import revise_srs_for_teacher as layout

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'output/specification/Fitness_Tracking_Management_System_SRS_中文版.docx'
OUTPUT = SOURCE.with_name('Fitness_Tracking_Management_System_SRS_中文版_PHP版.docx')
TMP = ROOT / 'tmp/srs_php'


def change(p, text):
    """Keep paragraph formatting and first-run styling while replacing text."""
    style = copy.deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear()
    r = p.add_run(text)
    if style is not None:
        r._r.insert(0, style)
    return p


def before(anchor, text, style='Normal'):
    p = anchor.insert_paragraph_before(text, style)
    p.paragraph_format.widow_control = True
    for run in p.runs:
        layout.fonts(run, mono=style == 'Code SRS')
    return p


def code_before(anchor, text):
    lines = text.strip('\n').splitlines()
    for i, line in enumerate(lines):
        p = before(anchor, line or ' ', 'Code SRS')
        p.paragraph_format.keep_with_next = i < len(lines) - 1
        p._p.get_or_add_pPr().append(layout.x('shd', fill='F3F4F5'))
        for r in p.runs:
            layout.fonts(r, 9, mono=True)
    before(anchor, '')


def architecture():
    """Deterministic layered architecture; ERD assets are not altered."""
    im = Image.new('RGB', (1900, 1180), 'white')
    draw = ImageDraw.Draw(im)
    font_path = '/Library/Fonts/Arial Unicode.ttf'
    def text(x, y, value, size):
        draw.multiline_text((x, y), value, font=ImageFont.truetype(font_path, size), fill='#111111', spacing=11)
    nodes = [
        ('browser', 20, '使用者瀏覽器', 'HTML / CSS / JavaScript\n課表與表單互動　Fetch API　Chart.js 統計圖'),
        ('php', 295, 'PHP 網頁與 API 層', 'PHP 範本產生 HTML；/api 路由回傳 JSON\nJWT 驗證　權限與欄位檢查　規則式排課'),
        ('pdo', 575, 'PHP 資料存取層', 'PDO_PGSQL　參數化 SQL　交易與列鎖\nPhinx 管理 migration 與 seed'),
        ('db', 855, '雲端 PostgreSQL', '14 張核心業務資料表　PK / FK / CHECK\n帳號、偏好、課表、訓練組數與體重紀錄'),
    ]
    for ident, top, title, body in nodes:
        draw.rectangle((255, top, 1645, top + 190), outline='#7C848B', width=3, fill='#F8F9FA')
        text(290, top + 15, title, 43)
        text(290, top + 77, body, 37)
    for start, end, label in [(210,295,'HTTPS：HTML / JSON / Bearer JWT'), (485,575,'完成驗證後呼叫資料存取'), (765,855,'SQL / TLS　提交或回滾交易')]:
        draw.line((910,start,910,end-5),fill='#65717A',width=4)
        draw.polygon([(910,end),(899,end-18),(921,end-18)],fill='#65717A')
        text(955, start+21, label, 27)
    text(190,1100,'同源部署；首版無寄信、推播或 LLM 服務依賴。資料庫不直接開放瀏覽器連線。',29)
    stream = io.BytesIO()
    im.save(stream, format='PNG')
    (TMP/'architecture_php.png').write_bytes(stream.getvalue())
    return stream.getvalue()


def regenerate_toc(d):
    begin = next(p for p in d.paragraphs if any('TOC ' in (n.text or '') for n in p._p.xpath('.//w:instrText')))
    anchor = begin.insert_paragraph_before('[[TOC]]')
    current = begin._p
    while current is not None:
        following = current.getnext()
        end = bool(current.xpath('.//w:fldChar[@w:fldCharType="end"]'))
        current.getparent().remove(current)
        if end:
            break
        current = following
    headings = []
    ident = 1000
    for p in d.paragraphs:
        if not p.style.name.startswith('Heading '):
            continue
        level = int(p.style.name[-1])
        bookmarks = p._p.xpath('./w:bookmarkStart')
        if bookmarks:
            name = bookmarks[0].get(qn('w:name'))
        else:
            ident += 1
            name = f'PHP_SRS_{ident}'
            p._p.insert(1 if p._p.pPr is not None else 0, layout.x('bookmarkStart',id=ident,name=name))
            p._p.append(layout.x('bookmarkEnd',id=ident))
        headings.append((p.text,level,name))
    layout.DOC = d
    layout.HEADINGS = headings
    layout.build_toc(anchor)
    for r in d.element.xpath('.//w:r'):
        r.attrib.pop('data-page-bookmark',None)
    (TMP/'headings.json').write_text(json.dumps(headings,ensure_ascii=False,indent=2))


def revise():
    TMP.mkdir(parents=True,exist_ok=True)
    digest = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    d = Document(SOURCE)
    # Change specific paragraphs, not the authored background or database model.
    revisions = {
        8: '版本 1.2 PHP 版　｜　2026 年 10 月 2 日',
        71: '本系統採 PHP 網頁範本與同源 REST API。HTML、CSS 與 JavaScript 負責介面，PHP 負責產生頁面、登入授權、欄位驗證與排課，並透過 PDO 讀寫 PostgreSQL。資料庫憑證只留在伺服器，瀏覽器以 Fetch API 呼叫 /api 取得 JSON。',
        76: '本版以原生 PHP 分層開發，不要求大型框架或 ORM。public/index.php 為網頁與 API 入口，依路徑及 HTTP 方法分派至控制器；商業邏輯與 PDO SQL 分開存放。登入後以 JavaScript 切換內容，Token 僅存記憶體，重新整理需再次登入。',
        82: 'PHP 統一驗證 JSON 欄位的型別、長度、範圍及日期；不合法回 422，資源重複或狀態衝突回 409；成功查詢／更新回 200，新增 201，刪除 204。這些狀態碼由 PHP 控制器明確設定。',
        84: '多表新增與關聯替換由同一 PDO 連線使用 beginTransaction()、commit() 與 rollBack()。交易提交成功才回覆成功；失敗回滾並回傳不含 SQL 或機密的錯誤訊息，介面保留輸入供重試。',
        92: "2. PHP 以 password_hash($password, PASSWORD_BCRYPT, ['cost' => 10]) 雜湊，salt 由函式產生；同一交易建立 users、user_auth、user_sessions。註冊不建立 admins，角色為 USER。",
        93: '3. 使用 firebase/php-jwt 簽發 HS256 JWT，包含 sub=user_id、jti=user_session_id、role、iat、exp；效期 7,200 秒。驗證演算法固定為 HS256，不採用用戶傳入的演算法；提交成功後才回傳 Token。',
        102: '1. PHP 依正規化 Email 查詢 user_auth，以 password_verify() 比對密碼並檢查 is_active。成功登入時建立 user_sessions，過期时间為建立時間加 2 小時。'.replace('时间','時間'),
        112: '1. PHP 以 password_verify() 驗證目前密碼，使用 password_hash() 產生新雜湊；更新 password_hash 與撤銷所有登入工作階段必須在同一交易完成。',
        317: '使用 Phinx 管理 schema 與 seed，部署帳號執行 DDL，應用帳號僅有必要 DML 權限。seed 重跑不得新增重複資料；Phinx 的 phinxlog 為工具版本紀錄，不列入 14 張核心業務資料表。',
        326: 'NFR-05：PHP 以 password_hash() 的 bcrypt cost=10 保存密碼，以 password_verify() 驗證；JWT 使用 SECRET_KEY，效期 2 小時。登出或改密碼後舊工作階段立即被拒絕。',
        327: 'NFR-06：PHP PDO 使用 prepare() 與參數綁定，關閉模擬預備語句；排序欄位採允許清單，個人資源檢查所有權。範本輸出以 htmlspecialchars() 編碼，JavaScript 以 textContent 顯示使用者文字。',
        336: 'NFR-13：schema 變更具 Phinx migration；新空白資料庫可重建 14 張核心業務表、約束與 seed。工具另有 phinxlog 版本紀錄表；交付 composer.json 與 composer.lock 固定 PHP 套件版本。',
        338: 'NFR-15：README 包含 PHP、Composer 與 PDO_PGSQL 安裝、環境變數、啟動、migration、seed、PHPUnit 測試與部署。公開目錄限 public/，伺服器重啟不影響雲端資料。',
        357: '以下 SQL 由 PHP PDO::prepare() 與 execute() 執行；:user_id 取自已驗證 JWT。每個 SQL 敘述分別 prepare，同一交易共用 PDO 連線；DBeaver 以參數視窗綁定，不把使用者輸入拼入 SQL。',
        492: 'APP_ENV=development',
        494: '# Optional future integration; unused in version 1.2',
        497: '設定由 PHP 載入；雲端平台注入環境變數，本機以 vlucas/phpdotenv 載入 .env。組態讀取支援 $_ENV 與平台 getenv()；缺少必要值即拒絕啟動，禁止把機密印入回應或日誌。LLM_API_KEY 留空不影響首版排課。',
        500: '本系統以 PHP 與關聯式資料庫支援帳號、訓練設定、課表、每組紀錄及統計。使用者可由當日課表開始訓練，完成組數逐筆保存，再由 PHP PDO 執行 SQL 查詢結果。實作先完成資料表、約束、交易與 API，再整合介面及雲端驗收。',
        501: '本文件版本 1.2 為 PHP 實作規格，保留資料庫課程的六大章節、正規化及 SQL 需求。內容定義待實作與驗收的目標，完成情形以實際測試報告為準。',
    }
    original_paragraphs = list(d.paragraphs)
    for idx, value in revisions.items():
        change(original_paragraphs[idx],value)
    # Preserve table shapes, widths, repeated headers and every core dictionary.
    tech = [
        ('前端','HTML / CSS / JavaScript','PHP 範本產生頁面，JavaScript 處理表單、切換課表與逐組即時儲存，CSS 提供響應式排版。'),
        ('後端','PHP 8.4 / Composer','PHP 控制器處理路由、驗證、授權及排課；Composer 管理依賴，實際修補版本於部署時固定。'),
        ('資料庫','PostgreSQL','保留關聯鍵、CHECK、交易、JOIN、GROUP BY 及彙總查詢，資料模型不因 PHP 而更換。'),
        ('資料存取','PDO / PDO_PGSQL','使用手寫參數化 SQL、交易與列鎖，方便展示資料庫課程內容；不採 ORM。'),
        ('版本管理','Phinx / Git / GitHub','以 Phinx migration 與 seed 管理 schema；Git 管理程式及 SQL 版本。'),
        ('圖表','Chart.js','依 PHP API 的統計資料繪製訓練量與體重趨勢；固定套件版本並隨站提供資源。'),
        ('身份與密碼','firebase/php-jwt / PHP 密碼函式','JWT 效期 2 小時；password_hash() 明確指定 bcrypt cost=10，password_verify() 驗證。'),
        ('部署與機密','Docker / PHP / Apache / 環境變數','Render 以 Docker 或由學校 PHP 主機部署；公開 HTTPS，連線至雲端 PostgreSQL。'),
        ('資料管理與測試','DBeaver / PHPUnit / HTTP 測試工具','驗證資料表、ERD、SQL、PHP 規則、API 與交易；測試僅連線隔離的測試資料庫。'),
    ]
    for row, values in zip(d.tables[2].rows[1:],tech):
        for cell,value in zip(row.cells,values):
            change(cell.paragraphs[0],value)
    changes = {
        (0,3,1): '完成 14 張核心業務資料表及約束；以雲端 PostgreSQL 保存資料，18 項 FR 均可追溯到验收案例；統計須得到指定數值。'.replace('验','驗'),
        (4,4,5): 'PHP password_hash() 產生 bcrypt hash，cost=10；禁止明文密碼。',
        (20,1,4): '201；users/user_auth/user_sessions 各 +1；admins 不新增；password_verify() 驗證成功；JWT exp−iat=7200。',
        (21,2,2): '以 PHPUnit／HTTP 工具檢查 HTTPS/TLS、SQL 注入、跨用戶存取與文字跳脫；掃描機密；第 21 次失敗登入回 429，撤銷 Token 回 401。',
        (21,4,2): '執行 Phinx migration/seed，核對 14 張業務表與 phinxlog；備份還原後比對筆數、FK；確認 PHP 套件及 README 可重現。',
        (26,1,1): '沿用 repository；提交 PHP、SQL、migration、seed、測試、composer.json、composer.lock、README 與 .env.example；排除 .env、vendor/、備份與機密。',
        (26,2,1): 'Render 使用 Docker 部署 PHP／Apache，或使用支援 PHP 的學校雲端主機；網頁與 /api 同源，首頁公開，個人資料需登入，無需 GitHub 帳號。',
        (26,3,1): '雲端 PostgreSQL；由 PHP 解析 DATABASE_URL，建立 PDO 的 pgsql: DSN，帳密另行傳入。啟用 TLS，正式資料不得僅存在 localhost。',
        (26,4,1): '建立雲端 DB → 建置 PHP 映像與 Composer 依賴 → 注入環境變數 → migration／seed → 啟動 Web 服務 → HTTPS 健康檢查 → 核心驗收。',
        (26,8,1): '包含專案說明、資料庫設計、ERD、PHP／Composer 安裝、環境變數、本機開發、migration／seed、測試、部署與 Live Demo。',
        (27,7,1): 'PHP PDO、Phinx migration、seed、最小 DB 權限、索引及備份還原；TC-29。',
    }
    for (t,r,c),value in changes.items():
        change(d.tables[t].rows[r].cells[c].paragraphs[0],value)
    # Native PDO needs a distinct named placeholder at each occurrence.
    for p in d.paragraphs:
        if '本版包含 14 張資料表。' in p.text:
            change(p,p.text.replace('本版包含 14 張資料表。','本版包含 14 張核心業務資料表。'))
        elif 'e.exercise_name ILIKE :pattern' in p.text:
            change(p,p.text.replace(':pattern',':exercise_pattern'))
        elif 'bp.body_part_name ILIKE :pattern' in p.text:
            change(p,p.text.replace(':pattern',':body_part_pattern'))
    before(original_paragraphs[78], '同源介面僅以 Authorization 標頭攜帶 JWT，不開放任意來源 CORS。後端不信任前端驗證結果，所有寫入重新驗證。', 'Normal')
    # The system diagram is replaced; original two ER diagrams remain byte-identical.
    shape = d.inline_shapes[0]
    rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
    d.part.related_parts[rid]._blob = architecture()
    shape._inline.docPr.set('descr','瀏覽器以 HTTPS 呼叫 PHP 網頁及 API；PHP 經 PDO 執行參數化 SQL 與交易，資料保存於雲端 PostgreSQL。')
    # Keep its original physical ratio so the surrounding source layout is stable.
    # Add a practical PHP example without removing any original SQL.
    anchor = next(p for p in d.paragraphs if p.text == '附錄 C　課表推薦演算法')
    heading = before(anchor,'B.7 PHP PDO 執行範例','Heading 2')
    heading.paragraph_format.page_break_before = True
    before(anchor,'以下示範安全讀取本人訓練紀錄；$pdo 已由設定建立，$userId 來自伺服器驗證後的 JWT，而非表單輸入。對外回應以 JSON 傳送，錯誤訊息不得包含資料庫連線內容。')
    code_before(anchor,"""<?php
$pdo->setAttribute(PDO::ATTR_ERRMODE, PDO::ERRMODE_EXCEPTION);
$pdo->setAttribute(PDO::ATTR_EMULATE_PREPARES, false);

$sql = <<<'SQL'
SELECT workout_session_id, session_date, status
FROM workout_sessions
WHERE user_id = :user_id
ORDER BY session_date DESC, workout_session_id DESC
LIMIT :page_size OFFSET :offset
SQL;

$stmt = $pdo->prepare($sql);
$stmt->bindValue(':user_id', $userId, PDO::PARAM_INT);
$stmt->bindValue(':page_size', $pageSize, PDO::PARAM_INT);
$stmt->bindValue(':offset', $offset, PDO::PARAM_INT);
$stmt->execute();
$rows = $stmt->fetchAll(PDO::FETCH_ASSOC);
""")
    before(anchor,'分頁參數先依 3.1 驗證，LIMIT／OFFSET 以整數型別綁定。附錄 B.5 的 :exercise_pattern 與 :body_part_pattern 分別綁定相同搜尋字串，例如 %press%；不可在關閉模擬模式時重用同名佔位符。')
    before(anchor,'交易寫入使用同一 $pdo，先 beginTransaction()，依規格鎖定場次、驗證所有權及狀態，再 INSERT／UPDATE，成功 commit()。捕捉例外且 inTransaction() 為真時 rollBack()；不要把多個 SQL 指令合成一個 prepare() 呼叫。')
    before(anchor,'資料庫完整性錯誤依規格映射為 409；格式與範圍問題由 PHP 預先驗證並回 422。非預期資料庫錯誤回一般化 500，詳細資訊仅寫入遮蔽機密的伺服器日誌。'.replace('仅','僅'))
    # Extend deployment in place, preserving the existing six teacher chapters.
    env_anchor = original_paragraphs[487]
    before(env_anchor,'D.1 環境變數與資料庫連線','Heading 2')
    p = original_paragraphs[493]
    for line in ['APP_URL=', 'DB_SSLMODE=verify-full', 'DB_SSLROOTCERT=']:
        new = before(p,line,'Code SRS')
        new._p.get_or_add_pPr().append(layout.x('shd',fill='F3F4F5'))
        for r in new.runs: layout.fonts(r,9,mono=True)
    anchor = next(p for p in d.paragraphs if p.text == '附錄 E　資料庫課程重點與結論')
    before(anchor,'DATABASE_URL 須拆解 host、port、database、user 與 password；URL 編碼的帳密須解碼，轉成 PDO 可接受的 pgsql: DSN，不能直接把 postgresql:// URL 當作 PDO DSN。正式環境以 verify-full 驗證憑證與主機名，DB_SSLROOTCERT 指向供應商 CA；本機隔離環境可依需求調整 TLS。')
    before(anchor,'D.2 專案結構與本機開發','Heading 2').paragraph_format.page_break_before = True
    before(anchor,'採輕量分層結構，控制器負責請求與回應，服務層處理排課及交易規則，資料存取層保留可閱讀的 SQL。網站根目錄設為 public/，避免 .env、SQL 與套件目錄被下載。')
    code_before(anchor,"""public/                 # index.php, router.php, CSS and JavaScript
src/Controllers/        # HTTP validation and JSON responses
src/Services/           # workout planning and transaction rules
src/Repositories/       # PDO and parameterized SQL
src/Middleware/         # JWT validation and permission checks
views/                  # PHP page templates
config/                 # environment and database configuration
db/migrations/          # Phinx schema migrations
db/seeds/               # reference data and exercise catalog
tests/                  # PHPUnit and API tests
phinx.php               # environment-based database settings
composer.json           # required PHP extensions and packages
composer.lock           # locked dependency versions
""")
    before(anchor,'開發環境安裝 PHP 8.4、Composer、PDO_PGSQL 及 mbstring，套件包含 firebase/php-jwt、vlucas/phpdotenv 與 robmorgan/phinx；PHPUnit 為開發測試依賴。下列命令為後續實作完成後的操作契約，不代表現有程式已轉為 PHP。')
    code_before(anchor,"""composer install
# Create a local .env from .env.example and fill in local values.
php vendor/bin/phinx migrate -e development
php vendor/bin/phinx seed:run -e development
php -S 127.0.0.1:8000 -t public public/router.php
# In another terminal, using the isolated test database:
php vendor/bin/phpunit
""")
    before(anchor,'router.php 為本專案需實作的開發路由檔：靜態檔直接交由開發伺服器處理，其餘導向 public/index.php。前端資源隨網站提供，沒有獨立前端編譯步驟。')
    before(anchor,'D.3 正式部署與技術參考','Heading 2')
    before(anchor,'正式環境使用 PHP／Apache Docker 容器或學校支援 PHP 的主機，禁止以 php -S 作為公開服務。Render 選擇 Docker，容器 Web server 監聽平台指定連接埠；Apache 將 /api 等非靜態路徑導向入口檔，網站啟用 HTTPS。')
    before(anchor,'部署安裝 Composer 正式依賴，Phinx 必須存在於執行 migration 的部署環境；以專用帳號完成 migration／seed 後啟動應用程式。正式應用帳號不持有 DDL 權限；套件版本由 composer.lock 固定，切勿將 .env 打包進映像或 Git。')
    before(anchor,'驗收需確認重複部署不新增重複 seed、逐組存檔可重讀、PHP 程式原始碼及 .env 無法從公開網址下載，並執行第 6 章案例。14 張業務表及既有 PK／FK 保留；不得因更換語言而重置正式使用者資料。')
    d.add_paragraph('技術文件參考','Heading 3')
    for text in [
        'PHP PDO 與參數化查詢：https://www.php.net/manual/en/pdo.prepare.php',
        'PHP 密碼處理：https://www.php.net/manual/en/function.password-hash.php',
        'JWT 套件：https://github.com/googleapis/php-jwt',
        'Phinx migration：https://book.cakephp.org/phinx/0/en/index.html',
        'Chart.js 圖表：https://www.chartjs.org/docs/latest/',
        'Render Docker 部署：https://render.com/docs/docker',
    ]:
        p = d.add_paragraph(text)
        p.paragraph_format.space_after = Pt(3)
        for r in p.runs: layout.fonts(r,9)
    # Comments mark the significant changes without adding visible revision prose.
    for paragraph, note in [
        (original_paragraphs[71], '改為 PHP 網頁與同源 API 架構；原有 PostgreSQL 資料模型維持。'),
        (original_paragraphs[93], '保留老師要求的 JWT 兩小時效期，改由 PHP 套件執行。'),
        (original_paragraphs[357], 'SQL 改由 PHP PDO 執行，補上原生預備語句、整數綁定與交易規則。'),
    ]:
        d.add_comment(paragraph.runs,text=note,author='規格修訂',initials='SRS')
    # Keep technology selection together on its own page instead of a tail row page.
    original_paragraphs[74].paragraph_format.page_break_before = True
    regenerate_toc(d)
    d.core_properties.title = '健身紀錄管理系統 PHP 版系統規格書'
    d.core_properties.version = '1.2'
    d.save(OUTPUT)
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == digest
    result={'output':str(OUTPUT),'source_sha256':digest,'tables':len(d.tables),'images':len(d.inline_shapes),'comments':len(d.comments)}
    (TMP/'revision.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
    print(json.dumps(result,ensure_ascii=False))


def toc_pages(pdf_path, check_only=False):
    headings=json.loads((TMP/'headings.json').read_text())
    pdf=PdfReader(pdf_path)
    normalize=lambda t: re.sub(r'\s+','',t)
    page_map={}
    for n,page in enumerate(pdf.pages,1):
        lines=(page.extract_text() or '').splitlines()
        for title,level,name in headings:
            if level>2: continue
            if any(normalize(''.join(lines[i:i+count]))==normalize(title)
                   for i in range(len(lines)) for count in (1,2,3)):
                page_map[title]=n
    missing=[t for t,l,n in headings if l<=2 and t not in page_map]
    assert not missing,missing
    d=Document(OUTPUT)
    count=0
    for p in d.paragraphs:
        if p.style.name.lower() not in ('toc 1','toc 2'): continue
        texts=p._p.xpath('.//w:t')
        expected=str(page_map[texts[0].text])
        if check_only: assert texts[-1].text==expected,(texts[0].text,texts[-1].text,expected)
        else: texts[-1].text=expected
        count+=1
    if not check_only:
        d.save(OUTPUT)
        (TMP/'toc_pages.json').write_text(json.dumps(page_map,ensure_ascii=False,indent=2))
    print(json.dumps({'toc_entries':count,'pdf_pages':len(pdf.pages),'verified':check_only}))


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--toc-pdf')
    parser.add_argument('--check',action='store_true')
    args=parser.parse_args()
    if args.toc_pdf: toc_pages(args.toc_pdf,args.check)
    else: revise()
