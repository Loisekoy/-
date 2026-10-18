"""Revise the user's current PHP SRS against the supplied course rubric.

Document artifacts only. No application, production database or Git mutations.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
import re
import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader

import revise_srs_for_teacher as L
import revise_srs_php as P

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT/'output/specification/Fitness_Tracking_Management_System_SRS_中文版_PHP版.docx'
OUT = SOURCE.with_name('Fitness_Tracking_Management_System_SRS_中文版_PHP評分規範版.docx')
TMP = ROOT/'tmp/srs_rubric'
ASSETS = TMP/'assets'
D = None


def find(text):
    return next(p for p in D.paragraphs if p.text == text)


def contains(text):
    return next(p for p in D.paragraphs if text in p.text)


def replace(text, new):
    return P.change(contains(text), new)


def para(anchor, text='', style='Normal', page=False):
    p = P.before(anchor, text, style)
    if page:
        p.paragraph_format.page_break_before = True
    return p


def table(anchor, headers, rows, widths, size=10):
    L.DOC = D
    t = L.table(headers, rows, widths, size)
    spacer = D.paragraphs[-1]
    anchor._p.addprevious(t._tbl)
    anchor._p.addprevious(spacer._p)
    for ri, row in enumerate(t.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = ri == 0
    return t


def setcell(t, row, col, text):
    P.change(t.rows[row].cells[col].paragraphs[0], text)


def addrow(t, values):
    row = copy.deepcopy(t.rows[-1]._tr)
    t._tbl.append(row)
    for c, value in zip(t.rows[-1].cells, values):
        P.change(c.paragraphs[0], value)
    return t.rows[-1]


def code(anchor, content):
    P.code_before(anchor, content)


RELATIONS = [
    ('R01','users','user_auth','1','0..1','user_id','CASCADE','註冊交易另保證每個有效會員有 1 筆認證。'),
    ('R02','user_auth','user_sessions','1','0..N','auth_id','CASCADE','一次成功登入新增一個可撤銷工作階段。'),
    ('R03','users','admins','1','0..1','user_id','CASCADE','子表 PK 同時為 FK，表示管理員資格。'),
    ('R04','training_goals','users','0..1','0..N','training_goal_id','RESTRICT','會員尚未完成設定時，FK 可為 NULL。'),
    ('R05','users','user_body_parts','1','0..N','user_id','CASCADE','與 R06 共同實現使用者與部位 M:N。'),
    ('R06','body_parts','user_body_parts','1','0..N','body_part_id','RESTRICT','(user_id, body_part_id) 為複合 PK。'),
    ('R07','users','body_records','1','0..N','user_id','CASCADE','UNIQUE(user_id, record_date)，一天一筆。'),
    ('R08','users','workout_plans','1','0..N','user_id','CASCADE','每份課表只有一名擁有者。'),
    ('R09','training_goals','workout_plans','1','0..N','training_goal_id','RESTRICT','保存產生課表時的目標。'),
    ('R10','workout_plans','plan_days','1','0..N','workout_plan_id','CASCADE','DB 允許 0；成功產課交易必須有 2–6 天。'),
    ('R11','plan_days','plan_exercises','1','0..N','plan_day_id','CASCADE','DB 允許 0；成功產課交易每天 3–6 動作。'),
    ('R12','exercises','plan_exercises','1','0..N','exercise_id','RESTRICT','與 R11 實現訓練日與動作 M:N。'),
    ('R13','body_parts','exercises','1','0..N','body_part_id','RESTRICT','每個動作記錄一個主要訓練部位。'),
    ('R14','users','workout_sessions','1','0..N','user_id','CASCADE','保存實際訓練者，不從來源課表推導。'),
    ('R15','workout_sessions','workout_session_plans','1','0..1','workout_session_id','CASCADE','每場至多一個來源；來源刪除不刪場次。'),
    ('R16','plan_days','workout_session_plans','1','0..N','plan_day_id','CASCADE','刪課表日只連帶刪除來源關聯列。'),
    ('R17','workout_sessions','workout_sets','1','0..N','workout_session_id','CASCADE','completed 場次至少一組，由交易驗證。'),
    ('R18','exercises','workout_sets','1','0..N','exercise_id','RESTRICT','保留歷史實際動作，不依附課表動作列。'),
]


def draw_diagrams():
    ASSETS.mkdir(parents=True, exist_ok=True)
    font='/Library/Fonts/Arial Unicode.ttf'
    def F(n): return ImageFont.truetype(font,n)
    # All boxes use key-level compartments; full types and all attributes are in 4.2.
    keys={
      'users':['PK user_id','FK training_goal_id (NULL)','name / training_experience'],
      'user_auth':['PK auth_id','FK UQ user_id','UQ email / password_hash'],
      'user_sessions':['PK user_session_id','FK auth_id','expires_at / revoked_at'],
      'admins':['PK FK user_id','created_at'],
      'training_goals':['PK training_goal_id','UQ goal_name'],
      'user_body_parts':['PK FK user_id','PK FK body_part_id','priority'],
      'body_parts':['PK body_part_id','UQ body_part_name'],
      'body_records':['PK body_record_id','FK user_id','record_date / weight_kg'],
      'workout_plans':['PK workout_plan_id','FK user_id','FK training_goal_id','experience_snapshot','duration_minutes'],
      'plan_days':['PK plan_day_id','FK workout_plan_id','day_number / day_title'],
      'plan_exercises':['PK plan_exercise_id','FK plan_day_id','FK exercise_id','target_sets / target_reps'],
      'exercises':['PK exercise_id','FK body_part_id','UQ exercise_name'],
      'workout_sessions':['PK workout_session_id','FK user_id','status / started_at / ended_at'],
      'workout_session_plans':['PK FK workout_session_id','FK plan_day_id'],
      'workout_sets':['PK workout_set_id','FK workout_session_id','FK exercise_id','set_number / weight_kg / reps'],
    }
    views=[
      ('erd_accounts',1800,{
       'training_goals':(640,40,520,210),'users':(640,450,520,250),
       'user_auth':(0,450,520,250),'admins':(1280,450,520,210),
       'user_sessions':(0,900,520,250),'user_body_parts':(640,900,520,250),
       'body_parts':(1280,900,520,210),'body_records':(640,1370,520,250)},[
       ('R01',[(640,540),(520,540)]),('R02',[(260,700),(260,900)]),
       ('R03',[(1160,540),(1280,540)]),('R04',[(900,250),(900,450)]),
       ('R05',[(900,700),(900,900)]),('R06',[(1280,990),(1160,990)]),
       ('R07',[(640,660),(580,660),(580,1490),(640,1490)])]),
      ('erd_plans',2080,{
       'users':(0,40,520,300),'training_goals':(1280,40,520,260),
       'workout_plans':(640,490,520,400),'plan_days':(640,1080,520,300),
       'plan_exercises':(640,1570,520,350),'exercises':(1280,1570,520,300),
       'body_parts':(1280,1080,520,260)},[
       ('R08',[(260,340),(260,650),(640,650)]),
       ('R09',[(1540,300),(1540,650),(1160,650)]),
       ('R10',[(900,890),(900,1080)]),('R11',[(900,1380),(900,1570)]),
       ('R12',[(1280,1700),(1160,1700)]),('R13',[(1540,1340),(1540,1570)])]),
      ('erd_tracking',1650,{
       'users':(0,40,570,300),'plan_days':(1230,40,570,300),
       'workout_sessions':(0,570,570,300),'workout_session_plans':(1230,570,570,260),
       'workout_sets':(0,1150,570,350),'exercises':(1230,1150,570,300)},[
       ('R14',[(285,340),(285,570)]),('R15',[(570,700),(1230,700)]),
       ('R16',[(1515,340),(1515,570)]),('R17',[(285,870),(285,1150)]),
       ('R18',[(1230,1320),(570,1320)])]),
    ]
    all_edges=[]
    for name,height,nodes,edges in views:
        im=Image.new('RGB',(1880,height),'white'); dr=ImageDraw.Draw(im)
        def label(x,y,text,size=31):
            box=dr.textbbox((x,y),text,font=F(size)); dr.rectangle((box[0]-3,box[1]-3,box[2]+3,box[3]+3),fill='white')
            dr.text((x,y),text,font=F(size),fill='#222222')
        for ident,points in edges:
            r=next(r for r in RELATIONS if r[0]==ident)
            points=[(x+40,y) for x,y in points]
            dr.line(points,fill='#59636A',width=3)
            for pt,other,value in [(points[0],points[1],r[3]),(points[-1],points[-2],r[4])]:
                x,y=pt; ox,oy=other
                if x==ox: label(x+10,y+(16 if oy>y else -50),value,29)
                else:
                    tw=dr.textlength(value,font=F(29))
                    label(x+8 if ox>x else x-tw-8,y-44,value,29)
            if len(points)==2:
                (x,y),(xx,yy)=points
                label((x+xx)/2+12 if x==xx else (x+xx)/2-32,(y+yy)/2-8 if x==xx else y+16,ident,29)
            else:
                label(points[1][0]-66 if ident=='R07' else points[1][0]+12,points[1][1]+40,ident,29)
            all_edges.append(ident)
        for node,(x,y,w,h) in nodes.items():
            x+=40
            dr.rectangle((x,y,x+w,y+h),fill='#FFFFFF',outline='#66717A',width=3)
            dr.rectangle((x,y,x+w,y+74),fill='#E9ECEF',outline='#66717A',width=2)
            title_size=36 if len(node)>20 else 39
            dr.text((x+15,y+17),node,font=F(title_size),fill='#111111')
            for i,k in enumerate(keys[node]):
                dr.text((x+16,y+96+i*51),k,font=F(36),fill='#111111')
        dr.text((40,height-95),'PK 主鍵　FK 外鍵　UQ 唯一鍵　0..1 至多一筆　0..N 零至多筆',font=F(31),fill='#333333')
        dr.text((40,height-48),'基數標於各實體端；R 編號對照 4.3。跨圖同名方塊為同一張表。',font=F(29),fill='#333333')
        im.save(ASSETS/f'{name}.png')
    assert set(all_edges)=={r[0] for r in RELATIONS}
    assert set().union(*(set(v[2]) for v in views))==set(keys)
    (ASSETS/'erd_model.json').write_text(json.dumps({'entities':keys,'relationships':RELATIONS,'views':views},ensure_ascii=False,indent=2))
    lines=['erDiagram']
    for name,fields in keys.items():
        lines.append(f'    {name} {{')
        for field in fields:
            if field.startswith(('PK','FK')):
                words=field.split(); col=next(w for w in words if w not in ('PK','FK','UQ'))
                tags=','.join(w for w in words if w in ('PK','FK','UQ')).replace('UQ','UK')
                typ='uuid' if col=='user_session_id' else 'int'
                lines.append(f'        {typ} {col} {tags}')
        lines.append('    }')
    for rid,a,b,ca,cb,fk,rule,note in RELATIONS:
        left='|o' if ca=='0..1' else '||'
        right='o|' if cb=='0..1' else 'o{'
        lines.append(f'    {a} {left}--{right} {b} : "{rid} {fk}"')
    (ASSETS/'Fitness_ERD.mmd').write_text('\n'.join(lines)+'\n')
    # Refresh source-backed architecture label only by regenerating its schematic.
    P.TMP=TMP
    original=P.architecture()
    # Native architecture generated here to reflect the now normalized business model.
    im=Image.new('RGB',(1900,1180),'white'); dr=ImageDraw.Draw(im)
    for top,title,body in [
      (20,'使用者瀏覽器','HTML / CSS / JavaScript\n課表與表單　Fetch API　Chart.js 統計'),
      (295,'PHP 網頁與 API 層','PHP 範本與 JSON API\nJWT 驗證　所有權檢查　規則式排課'),
      (575,'PHP 資料存取層','PDO_PGSQL　參數化 SQL　交易與列鎖\nPhinx 管理 migration 與 seed'),
      (855,'雲端 PostgreSQL','15 張核心資料表　1 個唯讀檢視表\nPK / FK / UNIQUE / CHECK　索引與備份')]:
        dr.rectangle((255,top,1645,top+190),fill='#F7F8F9',outline='#79838B',width=3)
        dr.text((290,top+14),title,font=F(43),fill='#111111')
        dr.multiline_text((290,top+78),body,font=F(36),fill='#222222',spacing=10)
    for y,end,label in [(210,295,'HTTPS　HTML / JSON / Bearer JWT'),(485,575,'通過驗證後執行'),(765,855,'SQL / TLS　COMMIT 或 ROLLBACK')]:
        dr.line((910,y,910,end),fill='#65717A',width=4)
        dr.polygon([(910,end),(899,end-18),(921,end-18)],fill='#65717A')
        dr.text((960,y+22),label,font=F(27),fill='#222222')
    dr.text((210,1100),'瀏覽器不直接連線資料庫；首版不依賴 LLM、寄信或推播服務。',font=F(31),fill='#333333')
    im.save(ASSETS/'architecture.png')


def insert_figure(anchor,path,caption,alt,page=False,width=16.4):
    p=para(anchor,'',page=page); p.alignment=1
    p.paragraph_format.keep_with_next=True; p.paragraph_format.line_spacing=1
    shape=p.add_run().add_picture(str(path),width=Cm(width))
    shape._inline.docPr.set('descr',alt)
    c=para(anchor,caption,'Caption'); c.alignment=1
    return p


def build():
    global D
    TMP.mkdir(parents=True,exist_ok=True)
    digest=hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    D=Document(SOURCE); L.DOC=D
    ts=list(D.tables)
    # Keep the user's cover attribution and recently rewritten introductory wording.
    para(find('資料庫系統課程專題'),'版本 1.3　PHP 版　2026 年 10 月 2 日').alignment=1
    for p in D.paragraphs:
        if p.style.name=='Normal (Web)': p.style=D.styles['Normal']
        if p.style.name.startswith('Heading') and not p.text.strip(): p.style=D.styles['Normal']
    for p in D.paragraphs:
        if '14 張' in p.text: P.change(p,p.text.replace('14 張','15 張'))
        if '版本 1.2' in p.text: P.change(p,p.text.replace('版本 1.2','版本 1.3'))
        if 'version 1.2' in p.text: P.change(p,p.text.replace('version 1.2','version 1.3'))
    for t in ts:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    if '14 張' in p.text: P.change(p,p.text.replace('14 張','15 張'))
    replace('輸入資料：name；gender','輸入資料：name 1–100 字；gender 選填 Female／Male／Other；age 13–100；height_cm 100–250；weight_kg 20–500；training_experience 為 Beginner／Intermediate／Advanced。')
    replace('3. 使用 firebase/php-jwt 簽發','3. 使用 firebase/php-jwt 簽發固定 HS256 JWT：sub 為 user_id 的十進位字串、jti=user_session_id，另含 role、iat、exp；效期 7,200 秒。不採用請求指定的演算法，提交成功後才回傳 Token。')
    replace('2. 每次受保護請求均驗證 JWT','2. 每次受保護請求均驗證簽章、exp、jti，JOIN user_sessions 與 user_auth 確認 sub 對應本人、expires_at 未到期、revoked_at 為 NULL 且帳號啟用。管理員資格以 admins 即時結果為準，不只信任 Token 中的 role。')
    replace('3. 該課表存在進行中訓練時','3. 該課表有進行中場次時禁止修改與刪除。其餘時段刪除課表會 CASCADE 刪除 days、plan_exercises 與 workout_session_plans 來源關聯；保留 workout_sessions 與 sets，歷史回傳 source_plan_day_id=NULL。')
    replace('2. 更新時驗證本人所有權','2. 更新時驗證本人所有權、啟用狀態，以及課表保存的 experience_snapshot 與 duration_minutes；同日動作不重複，修改後仍須符合附錄 C 的主要部位覆蓋與時間限制。')
    replace('2. 無進行中場次時 INSERT','2. 無進行中場次時，在同一交易 INSERT workout_sessions（user_id、started_at、status=in_progress）與 workout_session_plans（workout_session_id、plan_day_id）。顯示日期由開始時間換算，不另存 session_date。')
    replace('3. 相同課表日已有進行中場次時','3. 以 workout_session_plans 查來源日。相同日已有進行中場次則回傳該筆與已存 sets；不同日回 409。部分唯一索引防止並行建立兩筆，來源關聯寫入失敗則整筆回滾。')
    replace('1. 鎖定本人 workout_sessions，確認','1. 鎖定本人 workout_sessions 並確認 in_progress；JOIN workout_session_plans、plan_exercises 檢查 exercise_id 屬於来源日，再驗證數值。進行中場次缺少來源關聯視為完整性錯誤，不接受寫入。'.replace('来源','來源'))
    replace('1. Workout History 預設','1. Workout History 預設列出 completed 與 cancelled；只查本人。以 workout_session_details 取得台北 session_date 與可空的 source_plan_day_id，依 session_date、workout_session_id 倒序。')
    replace('輸出／回應：GET /api/workout-sessions 或','輸出／回應：GET /api/workout-sessions、GET /api/workout-sessions/{id} 或 GET /api/exercises/{id}/history → 200；清單含 items 與分頁，明細含場次及 sets，無清單資料回空陣列。')
    replace('時間戳以 UTC','时间戳以 UTC 的 TIMESTAMPTZ 儲存；API 日期採 YYYY-MM-DD、時間戳採含 Z 的 ISO 8601。訓練日期由 started_at AT TIME ZONE \'Asia/Taipei\' 取 DATE，週一為一週起點。'.replace('时间','時間'))
    replace('3. 後端驗證天數、候選 ID','3. 驗證天數、候選 ID、難度、覆蓋與時間後，單一交易新增 workout_plans、plan_days、plan_exercises。課表保存當次 experience_snapshot 與 duration_minutes；日數由 plan_days 計算。新課表不覆蓋舊課表。')
    replace('2. 在 users 更新選項','2. 在 users 更新選項；已產生課表保留當時目標、程度與時間設定，重新產生才套用新值。既有課表顯示「建立時設定」，不與現在的個人設定混淆。')
    replace('輸入資料：record_date≤台北今日','輸入資料：record_date 為 1900-01-01 至台北今日；weight_kg 20–500；notes≤1,000 字；body_record_id 用於修改／刪除。PATCH 可改日期、體重與 notes，新日期仍須唯一。')
    replace('輸入資料：workout_session_id、exercise_id','輸入資料：workout_session_id 取自網址；JSON 傳 exercise_id、set_number 1–100、weight_kg 0–999.99、reps 1–999。不得另傳不同的場次 ID 或 user_id。')
    replace('輸入資料：start_date、end_date、exercise_id','輸入資料：start_date、end_date、exercise_id、status 選填；分頁參數。省略日期時採含今日最近 28 天；只提供一端日期回 422。')
    replace('圖片載入失敗時顯示本機預設示意','2. 回傳主要部位、難度、器材、動作說明與 image_url；圖片載入失敗時改用網站隨附且標示「非此動作示範」的佔位圖，文字步驟仍可閱讀。')
    replace('15 張核心業務資料表。users','本版包含 15 張核心業務資料表與 1 個唯讀檢視表。users 保存會員資料，user_auth 保存認證，user_sessions 保存登入狀態；workout_sessions 則保存運動場次。workout_session_plans 專門記錄場次來源；workout_session_details 以 JOIN 與日期換算提供查詢，不重複儲存資料。')
    replace('設計修正：目前體重','本版的資料字典為實作依據。登入資料與運動紀錄分開；目前體重與訓練量透過查詢推導；場次來源拆為獨立關聯表，台北訓練日期由開始時間換算。workout_plans 保存建立時的程度與時間，避免修改個人設定後改變舊課表的判定標準。')
    replace('IDENTITY 為 PostgreSQL','IDENTITY 為 PostgreSQL 自動流水號；「無」表示沒有 DB DEFAULT，由 API 提供。NULL 僅供選填或未完成設定。固定列舉與單列數值範圍以 CHECK 實作，必填值用 NOT NULL；跨表、權限與依今日變動的規則由交易／API 驗證，不能籠統當作單列 CHECK。')
    replace('users 與 body_parts 透過','三張 ERD 共描述 15 張表與 18 條 FK 關係；同名實體跨圖重複僅為閱讀導航。users↔body_parts、plan_days↔exercises 皆透過關聯表實現 M:N。圖中標示 DB 可保證的最小／最大基數；註冊後 1:1、產課後每日至少 3 動作等業務最低數量由同一交易保證，詳見 4.3。')
    # Normalize the session model: dates are derived; source association is separate.
    for row in list(ts[14].rows):
        if row.cells[0].text in ('plan_day_id','session_date'):
            row._tr.getparent().remove(row._tr)
    for row in ts[11].rows:
        if row.cells[0].text=='created_at':
            setcell(ts[11],len(ts[11].rows)-1,5,'建立時間；既有課表不隨個人設定變更而改寫。')
    addrow(ts[11],('experience_snapshot','VARCHAR(20)','—','否','無','建立時程度：Beginner／Intermediate／Advanced；不隨 users 更新。'))
    addrow(ts[11],('duration_minutes','SMALLINT','—','否','無','建立時單次時間；CHECK IN (30,60,90)；課表編輯採此上限。'))
    addrow(ts[17],('偏好順位唯一性','UNIQUE (user_id, priority)；priority 依偏好陣列連續設為 1…N。'))
    setcell(ts[17],7,1,'刪除課表連帶刪除 days、plan_exercises 與來源關聯 workout_session_plans；不沿來源關聯反向刪除 workout_sessions 或 sets。')
    setcell(ts[17],9,1,'API 透過 workout_session_plans JOIN 回課表擁有者；FK 不等於授權。進行中場次必須有來源關聯；開始／產課交易驗證最低子項數。')
    setcell(ts[18],2,1,'workout_sessions(user_id, started_at DESC)；日期篩選轉為台北日起訖的 UTC 半開時間範圍。')
    setcell(ts[18],2,0,'idx_sessions_user_started')
    for values in [
      ('idx_plans_user','workout_plans(user_id)；列出本人課表。'),
      ('idx_plan_exercises_exercise','plan_exercises(exercise_id)；檢查動作引用，補足複合唯一鍵未覆蓋的 FK。'),
      ('idx_session_plans_day','workout_session_plans(plan_day_id)；檢查進行中課表引用與來源 CASCADE。'),
      ('其他 FK 索引','users(training_goal_id)、workout_plans(training_goal_id)、user_body_parts(body_part_id)、exercises(body_part_id)；檢查分類引用。既有部分索引不涵蓋停用動作。')]: addrow(ts[18],values)
    setcell(ts[20],16,4,'PATCH 200；DELETE 204；P/day/plan_exercises/來源關聯消失；S、sets 保留，檢視表 source_plan_day_id=NULL。')
    setcell(ts[20],17,4,'只一筆 in_progress 與一筆 workout_session_plans；回同 ID；另一日 409；current 載入已存 sets。')
    setcell(ts[20],24,2,'2026-10-05 至 10-11；S1 Chest 三組 60×12、65×10、65×8；S2 Back 一組 20×10；另有 cancelled 100×10；體重 50、51。')
    setcell(ts[20],24,4,'總量 2090.00、週次數 2；Chest 3 組第一、Back 1 組第二；平均體重 50.50；取消不計；空區間為 0／空陣列、平均體重 NULL。')
    for row in ts[22].rows[1:]:
        if row.cells[0].text in ('FR-11','FR-12','FR-13','FR-15','FR-16'):
            P.change(row.cells[1].paragraphs[0],row.cells[1].text+' / workout_session_plans')
    addrow(ts[23],('workout_session_plans','開始運動同交易建立','查來源課表日','來源不可改','刪課表／場次時 CASCADE'))
    setcell(ts[23],1,4,'無前端物理刪除；停用屬 Update')
    setcell(ts[24],1,1,'D1：Chest、Back、Legs；D2：Shoulders、Glutes、Core。僅當時間與欄位足夠時補 Biceps、Triceps；2×30 分鐘不保證手臂獨立動作。')
    # All read-only session queries consistently use the derived view.
    for p in D.paragraphs:
        if p.style.name=='Code SRS' and p.text=='FROM workout_sessions AS ws': P.change(p,'FROM workout_session_details AS ws')
        elif p.style.name=='Code SRS' and p.text=='JOIN workout_sessions AS ws': P.change(p,'JOIN workout_session_details AS ws')
        elif p.style.name=='Code SRS' and p.text=='FROM workout_sessions': P.change(p,'FROM workout_session_details')
    replace('workout_sessions.user_id 表示實際執行者','workout_sessions 只保存實際執行者及開始／結束時間；來源 plan_day_id 移至 workout_session_plans，避免在同一表中形成 plan_day_id→user_id 的非鍵相依。session_date 亦不儲存，而由檢視表換算，因此不存在 started_at→session_date 的重複儲存。users 與 user_auth 分表本身不是 3NF 的證明，仍須檢查候選鍵與函數相依。')
    replace('目前體重以 body_records 最新值','目前體重由 body_records 最新值推導，訓練量以 weight_kg×reps 計算，不保存總量。workout_plans 的目標、experience_snapshot 與 duration_minutes 是建立時設定，不是目前 users 設定的同步副本；角色由 admins 是否存在判定，不另存 role。')
    # New scope and rubric map: concise, concrete and not a self-awarded grade.
    a=find('2. 系統架構與技術棧')
    para(a,'1.4 系統範圍與作業責任','Heading 2')
    para(a,'本專題以健身者為一般會員，以受信任的維護者為管理員。老師可開啟公開網址自行註冊測試；管理功能使用另行安全提供的測試帳號，不將帳密寫在規格書或 GitHub。所有測試使用測試資料，不宣稱課表具有醫療處方效果。')
    table(a,['模組','完成範圍','不在本版範圍'],[
      ('會員與個資','註冊、登入／登出、改密碼、本人停用、基本資料','寄信驗證、忘記密碼、社群登入、他人帳號管理'),
      ('課表與動作','目標、偏好、規則排課、圖片與文字、修改及刪除','LLM 排課、器材可用性最佳化、影片辨識'),
      ('訓練與統計','開始／續練、每組存檔、終止、歷史、體重與 Dashboard','穿戴裝置同步、離線自動同步、營養與醫療建議'),
      ('資料管理','15 張表、1 個 view、migration、seed、索引、備份還原','讓一般使用者直接執行 SQL 或接觸 DB 管理密碼')],[3.1,7,6.6])
    para(a,'訪客只進入首頁與認證流程；登入但設定未完成的會員可修改資料，不能產課。管理員管理共用參考資料，不取得他人私人紀錄權限。資料庫維運帳號不等於網站管理員，由部署人員保管，僅在隔離環境展示 SQL 與備份還原。')
    para(a,'1.5 評分項目與文件對照','Heading 2')
    table(a,['評分項目','權重','本文件的具體對應'],[
      ('需求分析與完整度','25%','1.1–1.4 目標、角色、範圍；3.2–3.10 完整流程與例外。'),
      ('功能邏輯與規格嚴謹性','25%','18 項 FR 均有輸入／處理／輸出；3.1 與 3.10 定義格式、状态與重試。'.replace('状态','狀態')),
      ('資料庫與架構設計','20%','圖 1 架構、圖 2–4 ERD；4.2 資料字典、4.3 外鍵；附錄 A 3NF、B SQL。'),
      ('非功能需求與測試案例','15%','15 項 NFR；第 6 章案例、量測流程、資料庫驗證與可重現數值。'),
      ('圖表規範與文件排版','15%','可編輯 Word、PDF、自動目錄、分區 ERD、重複表頭與程式碼縮排。')],[5.1,1.6,10])
    # Tighten transport rules, deterministic boundaries and examples.
    a=find('3.2 會員與權限管理')
    para(a,'API 採 UTF-8 JSON；整數欄位拒絕小數與布林值。重量／身高最多兩位小數，超過精度回 422，不默默四捨五入；回傳 NUMERIC 採兩位小數字串，避免前端浮點誤差。文字長度以 Unicode 字元計算，密碼例外按 UTF-8 bytes 計算且不得含 NUL。')
    para(a,'清單回傳 {items, page, page_size, total}，page 必須為正整數；超過末頁回 items=[]。日期範圍含起訖兩日，最多 366 日；只傳一端、格式錯誤或起日大於迄日回 422。未知 JSON 欄位一律回 422；PATCH 未提供的欄位維持原值。')
    para(a,'所有成功寫入均在 COMMIT 後回應。連線失敗時不可推斷寫入失敗：先重新查詢，再重送同一組號；只有 FR-12、FR-13 及完成／取消具有明確重試規則。產課每次成功都建立新版本，前端按鈕送出後禁用至請求結束，逾時先查看課表清單。')
    a=find('4. 資料模型與結構設計')
    para(a,'3.10 訓練狀態與介面回應','Heading 2',page=True)
    table(a,['目前狀態','操作與條件','結果與資料庫變化'],[
      ('無進行中場次','開始本人課表日','201；同交易新增場次及來源關聯。'),
      ('in_progress','再次開始同日／另一日','同日 200 回原 ID；另一日 409，不建立新場次。'),
      ('in_progress','儲存一組／完全相同重送','201 新增一組；相同重送 200，筆數不變。'),
      ('in_progress','至少一組後完成','200；改 completed 並填 ended_at。空場次 409。'),
      ('in_progress','取消','200；改 cancelled 並填 ended_at，保留組數、排除統計。'),
      ('completed／cancelled','重送同種終止操作','200 回原場次；ended_at 不改。改另一終態回 409。'),
      ('completed','修改／刪除既有組','允許修正；刪最後一組 409。新增組一律 409。'),
      ('cancelled','修改、刪除組或新增組','409；只能刪除整場。'),
      ('completed／cancelled','刪除整場','204；場次、來源關聯、sets 同交易刪除。')],[3.5,6,7.2])
    para(a,'訓練頁按課表順序顯示動作、主要部位、器材、圖片、文字步驟、目標組次、休息秒數與每組輸入欄位。今天指使用者選擇的課表日，不依星期自動猜測。計時與預估時間只作操作提示，不影響實際存檔。')
    para(a,'介面必須區分「尚未儲存」「儲存中」「已儲存」「儲存失敗」。只有成功回應才標示已儲存；失敗保留輸入並允許重試。JWT 到期顯示登入提示，已存資料不受影響；未存輸入只暫存目前頁面記憶體，關閉或重新整理可能遺失，離開前須提示。')
    para(a,'逐組儲存範例','Heading 3')
    code(a,'POST /api/workout-sessions/101/sets\nAuthorization: Bearer <JWT>\nContent-Type: application/json\n\n{"exercise_id": 7, "set_number": 1, "weight_kg": 60, "reps": 12}\n\n201 Created\n{"workout_set_id": 501, "set_volume": "720.00"}\n\n409 Conflict\n{"code": "SET_VALUE_CONFLICT",\n "message": "這組已存在，請使用修改功能。"}')
    # Figures: replace old pair listings with proper connected ER diagrams.
    draw_diagrams()
    shapes=list(D.inline_shapes)
    shapes[0]._inline.graphic.graphicData.pic.blipFill.blip
    rid=shapes[0]._inline.graphic.graphicData.pic.blipFill.blip.embed
    D.part.related_parts[rid]._blob=(ASSETS/'architecture.png').read_bytes()
    for shape in shapes[1:]:
        pe=shape._inline.getparent().getparent(); pe.getparent().remove(pe)
    for p in list(D.paragraphs):
        if p.text.startswith(('圖 2　','圖 3　')): p._p.getparent().remove(p._p)
    erd_h=find('4.1 實體關係圖')._p
    next_el=erd_h.getnext()
    while next_el is not None and next_el.tag==qn('w:p') and not next_el.xpath('.//w:t|.//w:sectPr'):
        after=next_el.getnext(); next_el.getparent().remove(next_el); next_el=after
    a=contains('三張 ERD 共描述')
    para(a,'以下為鍵層級 ERD，列出 PK、FK 與辨識欄位；完整型態、NULL、DEFAULT 及約束見 4.2。連線兩端的數字表示最小與最大參與數，連線不是資料流方向。圖面依本版 schema 繪製，日期為 2026-10-02。')
    insert_figure(a,ASSETS/'erd_accounts.png','圖 2　帳號與偏好 ERD','R01–R07：會員認證、登入工作階段、管理員、目標、偏好部位與體重歷史。')
    insert_figure(a,ASSETS/'erd_plans.png','圖 3　課表規劃 ERD','R08–R13：課表擁有者與目標、訓練日、課表動作和主要部位。',page=True)
    insert_figure(a,ASSETS/'erd_tracking.png','圖 4　實際訓練 ERD','R14–R18：執行者、場次、來源關聯與每組實際紀錄；來源刪除仍保留歷史。',page=True)
    # New dictionary in the original landscape dictionary section.
    a=find('4.3 關聯唯一性與刪除規則')
    sectionbreak=a._p.getprevious()
    while sectionbreak is not None and not sectionbreak.xpath('.//w:sectPr'):
        sectionbreak=sectionbreak.getprevious()
    from docx.text.paragraph import Paragraph
    anchor=Paragraph(sectionbreak,D._body)
    para(anchor,'表 4.2.15　訓練來源關聯 workout_session_plans','Heading 3')
    table(anchor,['欄位名稱','資料型態','PK/FK','允許\nNULL','預設值','說明與約束'],[
      ('workout_session_id','INTEGER','PK/FK','否','無','→ workout_sessions.workout_session_id；ON DELETE CASCADE；每場至多一個來源。'),
      ('plan_day_id','INTEGER','FK','否','無','→ plan_days.plan_day_id；ON DELETE CASCADE；刪課表日只刪關聯，不刪場次。')],[4.7,3.6,1.5,1.5,3,11.1],10)
    para(anchor,'唯讀檢視表 workout_session_details','Heading 3')
    para(anchor,'不是第 16 張資料表，也不保存副本。輸出 workout_sessions 全部欄位、source_plan_day_id（可 NULL）與 session_date（台北開始日期）。由場次 LEFT JOIN 來源關聯取得；僅供查詢，寫入操作仍指定基礎表。定義見附錄 B.8。')
    # Complete relationship catalogue, with physical and business minima separated.
    a=find('4.4 交易索引與資料管理')
    para(a,'外鍵關係逐項對照','Heading 3')
    table(a,['編號','父表 → 子表','子表 FK／刪除規則'],[
      (rid,f'{parent}\n→ {child}',f'{fk}；{rule}\n{note}') for rid,parent,child,ca,cb,fk,rule,note in RELATIONS
    ],[1.2,7.4,8.1],9.5)
    a=find('5. 非功能性需求')
    para(a,'4.5 資料完整性與交易邊界','Heading 2',page=True)
    table(a,['規則','由何處保證','驗證方式'],[
      ('唯一帳號與資料範圍','DB UNIQUE、NOT NULL、CHECK；Email 額外 CHECK(email=lower(btrim(email)))。','繞過前端送非法 SQL，必須被 DB 拒絕。'),
      ('同一順位與同一組不重複','user_body_parts(user_id, priority) 唯一；sets 複合唯一鍵。','同鍵並行寫入，最終僅一筆。'),
      ('來源擁有者與狀態','API 授權 JOIN；以相同鎖順序在交易內再驗證。','兩名會員互換資源 ID，回 404 且筆數不變。'),
      ('課表日數與部位覆蓋','產課／編輯交易；跨列規則不可只靠單列 CHECK。','驗證日數、每日至少三動作與覆蓋，不符則全回滾。'),
      ('日期與最新體重','API 禁止未來體重日期；view 推導訓練日期；最新體重以 MAX 日期取值。','台北跨午夜與刪最新紀錄測試。'),
      ('歷史保存','來源關聯與實際場次分表；動作 FK RESTRICT。','刪課表後實際 sets、開始時間、執行者不變。')],[4.1,7.5,5.1],9.5)
    para(a,'同一使用者的訓練及課表寫入採固定鎖順序：users → workout_plans（需要時）→ workout_sessions → exercises（依 ID 升序，需要時）→ 子表。取得鎖後重新查狀態與所有權，避免「開始運動」與「刪課表」同時通過舊檢查。管理員動作修改／刪除先鎖該 exercise，再檢查引用，不反向鎖使用者。')
    para(a,'設定偏好也先鎖 users，避免兩次替換互相交錯。修改動作主部位時，若被 plan_exercises 或 workout_sets 任一引用即拒絕，避免改變既有課表覆蓋與歷史統計。跨交易衝突先 ROLLBACK，再回 409 TRANSACTION_CONFLICT；不宣稱尚未提交的結果已儲存。')
    para(a,'DB 備份與 migration 使用獨立維運帳號；網站應用帳號只取得必要 SELECT／INSERT／UPDATE／DELETE 與 sequence 權限，不授予 CREATE／DROP。上述索引是待驗證的設計，需以 EXPLAIN (ANALYZE, BUFFERS) 的實際計畫記錄使用情況，不以「有索引」直接宣稱更快。')
    replace('3. 若已有 workout_sets 引用','3. 若已有 plan_exercises 或 workout_sets 引用，不得修改 body_part_id；新增動作版本並停用舊動作。已儲存的來源課表因此不會被重新分類；修改與引用檢查遵循 4.5 的鎖順序。')
    # Statistical and algorithm definitions are executable specifications, not health advice.
    a=find('C.2 組數次數與休息規則')
    para(a,'候選選取的「本週使用次數」僅計算本次產生中的課表，不讀個人历史，以保證相同輸入可重現。未選偏好部位的 priority 視為 9。先按 day_number 升序放必備部位，再按相同日順序補未覆蓋部位，最後填偏好額外欄位；每新增一動作即更新頻率與時間。'.replace('历史','歷史'))
    para(a,'必要部位沒有任何合格候選即回 409 並列出缺少部位。額外部位沒有候選或加入後超時則跳過，繼續評估同日其餘候選；無候選可加入才結束該日。未覆蓋偏好須明列限制，不保證容量不足時仍增加偏好頻率。休息與執行秒數是軟體估計，不代表實際訓練耗時。')
    # Insert source view definition and make schema checks demonstrable.
    a=find('附錄 C　課表推薦演算法')
    para(a,'B.8 訓練來源與日期檢視表','Heading 2',page=True)
    code(a,"""CREATE TABLE workout_session_plans (
    workout_session_id INTEGER PRIMARY KEY
        REFERENCES workout_sessions(workout_session_id)
        ON DELETE CASCADE,
    plan_day_id INTEGER NOT NULL
        REFERENCES plan_days(plan_day_id)
        ON DELETE CASCADE
);

CREATE INDEX idx_session_plans_day
    ON workout_session_plans(plan_day_id);

CREATE VIEW workout_session_details AS
SELECT
    ws.*,
    wsp.plan_day_id AS source_plan_day_id,
    (ws.started_at AT TIME ZONE 'Asia/Taipei')::date
        AS session_date
FROM workout_sessions AS ws
LEFT JOIN workout_session_plans AS wsp
    ON wsp.workout_session_id = ws.workout_session_id;""")
    para(a,'本段為新版本目標 schema，實作時須寫 migration，不可對正式資料直接 DROP 重建。搬移舊 plan_day_id 前先核對擁有者、日期與來源；在同一受控遷移流程建立關聯、驗證筆數後才移除舊欄位。不得因重新設計而清除使用者紀錄。')
    para(a,'日期索引查詢','Heading 3')
    code(a,"""-- :start_utc / :end_utc are Taipei calendar boundaries.
-- The end boundary is the day after the requested end_date.
SELECT workout_session_id, started_at, status
FROM workout_sessions
WHERE user_id = :user_id
    AND started_at >= :start_utc
    AND started_at < :end_utc
ORDER BY started_at DESC, workout_session_id DESC;""")
    # Normalization evidence beyond prose assertions.
    a=find('附錄 B　SQL 查詢需求')
    para(a,'A.1 主要函數相依與異常消除','Heading 2',page=True)
    table(a,['資料表或關聯','候選鍵與相依','設計判斷'],[
      ('users','user_id → 個資與目前訓練設定','目標名稱留 training_goals；不存目前體重。'),
      ('user_auth','auth_id／user_id／email → 認證屬性','三者各為候選鍵；role 不重複存。'),
      ('user_sessions／admins','各 PK → 時間與資格屬性','登入事件及角色資格不混入 users。'),
      ('training_goals／body_parts','ID 或 UNIQUE 名稱 → description','名稱只在主表保存。'),
      ('user_body_parts','(user_id, body_part_id) → priority；(user_id, priority) → body_part_id','排序依整個鍵，不依單一 user_id 或部位。'),
      ('exercises','exercise_id 或 exercise_name → 動作屬性','只存主要部位 ID，不存部位名稱。'),
      ('workout_plans','workout_plan_id → 建立時目標、程度、時間','快照屬於課表事件，不要求與目前 users 值相等。'),
      ('plan_days','plan_day_id 或 (plan_id, day_number) → day_title','plan_id 在此簡稱 workout_plan_id。'),
      ('plan_exercises','PK 或 (day, exercise) 或 (day, order) → 目標組次','day、exercise、order 分別指對應 ID／exercise_order。'),
      ('workout_sessions','workout_session_id → 執行者、狀態、起訖時間','移除 plan_day_id 與可由時間推導的日期。'),
      ('workout_session_plans','workout_session_id → plan_day_id','來源日不是執行者；來源刪除只刪本關聯。'),
      ('workout_sets','PK 或 (session, exercise, set_number) → weight_kg, reps, created_at','不重存動作名、執行者或訓練量。'),
      ('body_records','PK 或 (user_id, record_date) → weight_kg, notes','每天一筆；不在 users 再保存最新體重。')],[4.4,6.4,5.9],9.5)
    para(a,'以上以已列出的業務相依與唯一約束判定：欄位為單值；非鍵欄位依賴完整候選鍵；不以非鍵欄位再決定同表另一非鍵欄位。view 可輸出推導值，但不新增儲存冗餘。保留歷史時，不須為了刪課表而刪實際訓練，亦不會因改目前體重而覆寫過去體重。')
    # Additional deterministic acceptance cases, preserved original 29 tests.
    a=find('附錄 A　正規化與 CRUD 對應')
    para(a,'6.4 邊界與資料完整性補充驗收','Heading 2',page=True)
    extra=[
      ('TC-30','FR-04 欄位邊界','A 已登入。','依序送 height_cm=99.99、100、250、250.01；再送 160.123。','依序 422、200、200、422、422；拒絕時 users 與今日體重均不變。'),
      ('TC-31','FR-12 來源原子性','可故障注入；A 無進行中。','令來源 INSERT 失敗，再查場次及來源表。','回 500；兩表新增數均 0，不留下不能續練的場次。'),
      ('TC-32','FR-11／12 併發','A 有 P，无進行中。'.replace('无','無'),'並行開始 P 的 D 與刪除 P。','開始先提交：刪除409；刪除先提交：開始404。不得留下進行中無來源。'),
      ('TC-33','NFR-03 日期','隔離資料由 fixture 指定 started_at。','存 UTC 2026-10-04 15:59:59 與 16:00:00；查 view。','日期依序 10-04（日）、10-05（一）；週起日依序 09-28、10-05，不新增日期欄位。'),
      ('TC-34','FR-10 低容量','2 天／30 分鐘，偏好 Biceps，合格候選齊全。','產課並計每日時間、主要部位與限制清單。','2 天各 3 動作；六主要部位齊全；不硬加第4動作；明列 Biceps 未獨立涵蓋。'),
      ('TC-35','FR-10 可重現','4 天／60 分鐘／Beginner／增肌／Chest、Back。','候選 seed 不變，連續產兩份。','排除 ID／時間／名稱，兩份每日動作順序、組次與時間相同；新課表 ID 不同。'),
      ('TC-36','FR-05／11 設定快照','P 建立於 Beginner、60 分鐘。','將個人改 Advanced、30 分鐘；讀 P，再編輯 P。','P 仍顯示原程度與60分鐘；修改按快照驗證。重產新課表才採新設定。'),
      ('TC-37','FR-15 狀態邊界','各有 cancelled 與 completed 場次。','對取消場次改組；對完成場次新增組；重送相同終止；改另一終態。','依序409、409、200、409；重送的 ended_at 不變，無多餘資料。'),
      ('TC-38','FR-07／18 清單與日期','A 有至少21筆動作候選。','查 page_size=20 第1/2頁；送0/101；體重改到已存在日期。','清單20/剩餘筆數；非法分頁422；日期衝突409且兩筆原值不變。'),
      ('TC-39','DB 約束','隔離DB；具直接測試寫入權限。','依序測重複Email、孤立FK、負重量、重複順位、同組號；各測完ROLLBACK。','各次SQL被約束拒絕；記錄 constraint 名稱與 SQLSTATE，不留下非法列。'),
      ('TC-40','圖片完整性','每部位至少2個Beginner啟用動作。','逐一核對圖片與動作名稱、來源授權及替代文字；模擬404。','16個初階動作皆有可用對應圖與步驟；404只顯示明確佔位圖，仍可存組。'),
      ('TC-41','FR-08 分類與引用','動作 E 已在 P 中，尚未實際訓練。','M 嘗試改 E 主要部位；再停用 E；A 開始 P。','改部位409；停用200；開始409提示替換。完成替換後可開始；不默默改動 P。'),
    ]
    table(a,['案例','功能與前置条件'.replace('条件','條件'),'操作步驟','預期結果'],[(r[0],r[1]+'\n'+r[2],r[3],r[4]) for r in extra],[1.5,4,5.5,5.7],9.5)
    para(a,'6.5 測試執行與交付證據','Heading 2',page=True)
    para(a,'每個案例均記錄版本、執行日期、環境、前置資料、實際回應、SQL 查詢截圖或輸出，以及通過／失敗／未執行。測試資料由 fixture 建立，ID 以回傳值為準；文件中的示例 ID 不假定存在。案例之間還原 fixture，避免前案刪除影響後案。')
    table(a,['驗收項目','可重現執行方式與判定'],[
      ('TC-26 效能','先暖機60秒，再30秒升至200人，維持5分鐘。讀取70%（me、歷史、Dashboard平均分配）、存組20%、產課10%；各用戶用獨立帳號與有效組號，100組後完成再開新場次。以維持期資料分端點計 P95；全部受測端點須達NFR-01，非預期錯誤率<1%。'),
      ('TC-27 安全','同一IP在15分鐘内先送20次錯誤密碼，第21次含正確密碼亦429；期間成功登入不重置失敗數。於第一筆失敗滿15分鐘後再試。另測撤銷JWT、他人ID、SQL注入字串、HTML文字跳脫、HTTPS與機密掃描。'.replace('内','內')),
      ('TC-28 易用性','記錄6個瀏覽器版本（3種各最新兩版），各測3種指定視窗。用鍵盤完成登入、產課與首組存檔；記錄按鈕操作數、44px點擊區域與錯誤提示。'),
      ('TC-29 還原','記錄備份最新成功時間T1與模擬故障時間T2，T2−T1≤24h；從啟動還原到查詢通過≤60min。還原隔離DB後核對15表、1view、筆數、約束、孤立FK為0，抽查關鍵紀錄值及下一個IDENTITY可新增。')],[3,13.7],9.5)
    para(a,'測試通過條件：TC-01～TC-41 均執行並達預期，NFR-01～15 均有對應證據；尚未執行者標示「未執行」，不得填入預估成功值。若雲端方案限制導致效能未達標，必須記錄實測值、限制與重測安排；不能僅靠文件承諾通過。')
    # Explicit operational details without inventing an extra application feature.
    a=find('5.3 易用性與相容性')
    para(a,'登入限流以整個部署共用的原子計數器執行，不得只存在單一請求或單一 PHP worker。部署交付須註明計數儲存、過期與故障策略；未能跨 worker 保持一致，TC-27 判定失敗。限流狀態是短期安全作業資料，不納入健身業務 ERD。')
    a=find('6. 驗收測試案例')
    sb=a._p.getprevious()
    while sb is not None and not sb.xpath('.//w:sectPr'): sb=sb.getprevious()
    a=Paragraph(sb,D._body)
    para(a,'動作種子資料至少 16 筆：八部位各 2 個 Beginner 動作，另可增加其他難度。每筆含步驟、注意事項及動作相符的自製或合法授權图片，資源清單記錄來源、作者及授權；不可把通用佔位圖當作該動作的教學圖片。'.replace('图片','圖片'))
    replace('例外處理：課表日不存在／非本人','例外處理：課表日不存在／非本人回 404；另一日進行中回 409 WORKOUT_ALREADY_ACTIVE；來源日含停用動作回 409 PLAN_REQUIRES_UPDATE，提示先替換，禁止自動改寫課表。')
    replace('2. 取消時由 in_progress','2. 取消時由 in_progress 改為 cancelled 並填 ended_at，保留 sets 但排除 Dashboard；重送取消回既有結果且不改 ended_at。完成與取消互不可轉換。')
    # Minimum browser preservation guarantee and clear ID contract.
    replace('1. 驗證課表日屬於本人；鎖定','1. 依 4.5 鎖順序先鎖 users 與來源 workout_plans，再確認來源日屬於本人、動作仍啟用，以及是否已有 in_progress 場次；不得在鎖定前決定刪除或開始結果。')
    replace('資料庫對應：workout_sessions、plan_days','資料庫對應：workout_sessions、workout_session_plans、plan_days、workout_plans；TC-17、TC-31、TC-32。')
    replace('資料庫對應：workout_sessions、workout_sets、plan_exercises','資料庫對應：workout_sessions、workout_session_plans、workout_sets、plan_exercises；TC-18、TC-19。')
    replace('資料庫對應：workout_plans、plan_days、plan_exercises、workout_sessions','資料庫對應：workout_plans、plan_days、plan_exercises、workout_sessions、workout_session_plans；TC-16、TC-32、TC-36。')
    replace('資料庫對應：workout_sessions、workout_sets；TC-21','資料庫對應：workout_sessions、workout_sets、workout_session_plans；TC-21、TC-22、TC-37。')
    replace('輸出／回應：GET 或 PUT →','輸出／回應：GET /api/me 或 PUT /api/me/profile → 200；回傳 profile、目前體重、訓練設定與依 priority 排序的偏好部位。')
    replace('輸入資料：workout_plan_id；可修改','輸入資料：workout_plan_id 取自網址。PATCH 接受 plan_name，以及 exercises 陣列；每個項目含 plan_exercise_id 與欲修改的 exercise_id、target_sets、target_reps、rest_seconds。未提供欄位不變，不接受其他課表的項目或重複項目 ID。')
    replace('3. 相同課表日已有進行中場次時','3. 相同來源日已有進行中場次則回原 ID；不同日回 409。') if any('3. 相同課表日已有進行中場次時' in p.text for p in D.paragraphs) else None
    a=contains('資料庫對應：workout_plans、plan_days、plan_exercises、workout_sessions')
    para(a,'本版課表編輯不新增／刪除訓練日或動作列，也不修改順序；調整整體日數須重新產課。每次 PATCH 的全部項目在單一交易中套用並驗證，任一項失敗則全部保留原值。')
    replace('本文件版本 1.3 為 PHP','本文件版本 1.3 為 PHP 與 PostgreSQL 的課程實作規格，依評分標準保留六大章節並補強正規化、完整 ERD 與驗收證據。文件描述實作目標；系統完成程度與效能以實際測試報告為準，不以規格書取代測試。')
    # Two verified primary references supporting technical additions.
    tail=D.add_paragraph('')
    para(tail,'PostgreSQL 約束：https://www.postgresql.org/docs/current/ddl-constraints.html')
    para(tail,'PostgreSQL 鎖定：https://www.postgresql.org/docs/current/explicit-locking.html')
    tail._p.getparent().remove(tail._p)
    # Same general contract, updated traceability and normalization annotations.
    for row in ts[22].rows[1:]:
        extras={'FR-04':'TC-30','FR-10':'TC-34～35','FR-11':'TC-32、36','FR-12':'TC-31～32、41','FR-15':'TC-37','FR-07':'TC-38、40','FR-08':'TC-41','FR-18':'TC-38'}
        if row.cells[0].text in extras:
            P.change(row.cells[2].paragraphs[0],row.cells[2].text+'；'+extras[row.cells[0].text])
    setcell(ts[27],2,1,'users→workout_sessions→workout_sets；users↔body_parts；plan_days↔exercises；來源關聯另表保存。')
    # Remove inherited editorial notes; insert concise comments at changed definitions.
    for el in list(D.element.xpath('.//w:commentRangeStart|.//w:commentRangeEnd|.//w:commentReference')):
        el.getparent().remove(el)
    if hasattr(D,'add_comment'):
        for target,message in [('4.1 實體關係圖','依評分要求重繪三張連通ERD，涵蓋全部15表與18條FK；R編號可與4.3交叉核對。'),('表 4.2.15　訓練來源關聯 workout_session_plans','為移除來源日與執行者間的非鍵相依，新增此關聯表；日期改由view計算，原始PHP版另行保留。'),('6.4 邊界與資料完整性補充驗收','補充欄位邊界、交易回滾、併發、跨日、低容量、圖片與DB約束測試；未宣稱實際已通過。')]:
            p=find(target)
            D.add_comment(p.runs, text=message, author='文件修訂', initials='SRS')
    # Rebuild native TOC; cache its correct page values only after rendering.
    P.TMP=TMP
    P.regenerate_toc(D)
    toc_begin=next(p for p in D.paragraphs if p._p.xpath('.//w:instrText'))
    para(toc_begin,'目錄','Title')
    find('5.3 易用性與相容性').paragraph_format.page_break_before=True
    for ri,row in enumerate(ts[22].rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing=Pt(12)
                p.paragraph_format.keep_with_next=ri==0
                for run in p.runs: L.fonts(run,9.5)
            for mar in cell._tc.xpath('./w:tcPr/w:tcMar/w:top|./w:tcPr/w:tcMar/w:bottom'):
                mar.set(qn('w:w'),'60')
    for p in D.paragraphs:
        if p.style.name=='Heading 1': p.paragraph_format.page_break_before=True
        if p.style.name.startswith('Heading'):
            p.paragraph_format.keep_with_next=True
        if p.text=='1. 專案概述': p.paragraph_format.page_break_before=False
    D.save(OUT)
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==digest
    manifest={'source':str(SOURCE),'source_sha256':digest,'output':str(OUT),'entities':15,'relationships':len(RELATIONS),'tests':41,'tables':len(D.tables),'figures':len(D.inline_shapes)}
    (TMP/'revision.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
    print(json.dumps(manifest,ensure_ascii=False))


if __name__=='__main__':
    ap=argparse.ArgumentParser(); ap.add_argument('--toc-pdf'); ap.add_argument('--check',action='store_true'); args=ap.parse_args()
    if args.toc_pdf:
        P.TMP=TMP; P.OUTPUT=OUT; P.toc_pages(args.toc_pdf,args.check)
    else: build()
