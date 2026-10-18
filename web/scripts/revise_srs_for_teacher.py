"""Revise the supplied, user-edited SRS to the six-part DBS.pdf structure.

This is a document revision only; it does not migrate or change the application.
The source is backed up once, including the student's manually edited cover.
"""
from __future__ import annotations

import argparse
import copy
import json
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output/specification/Fitness_Tracking_Management_System_SRS_中文版.docx'
BACKUP = OUT.parent / 'backups/Fitness_Tracking_Management_System_SRS_中文版_老師規範修訂前_20260930.docx'
TMP = ROOT / 'tmp/srs_teacher'
ASSETS = TMP / 'assets'
FONT = 'Arial Unicode MS'
GRAY = '44515C'
HEADINGS = []
DOC = None


def x(tag, **attrs):
    el = OxmlElement('w:' + tag)
    for k, v in attrs.items():
        el.set(qn('w:' + k), str(v))
    return el


def fonts(run, size=None, bold=None, color=None, mono=False):
    run.font.name = 'Liberation Mono' if mono else FONT
    rf = run._element.get_or_add_rPr().get_or_add_rFonts()
    for a in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        rf.set(qn('w:' + a), FONT if a == 'eastAsia' else run.font.name)
    for a in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        rf.attrib.pop(qn('w:' + a), None)
    if size is not None: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)


def paragraph(text='', style=None, keep=False):
    p = DOC.add_paragraph(text, style)
    p.paragraph_format.keep_with_next = keep
    p.paragraph_format.widow_control = True
    p._p.get_or_add_pPr().append(x('snapToGrid', val='false'))
    for r in p.runs: fonts(r)
    return p


def heading(text, level=1):
    p = paragraph(text, 'Heading ' + str(level), True)
    ident = len(HEADINGS) + 1
    name = 'SRS_' + str(ident)
    p._p.insert(1 if p._p.pPr is not None else 0, x('bookmarkStart', id=ident, name=name))
    p._p.append(x('bookmarkEnd', id=ident))
    HEADINGS.append((text, level, name))
    return p


def bullets(items):
    for t in items: paragraph(t, 'List Bullet')


def steps(items):
    for i, t in enumerate(items, 1):
        p = paragraph(f'{i}. {t}')
        p.paragraph_format.left_indent = Cm(.65)
        p.paragraph_format.first_line_indent = Cm(-.45)
        p.paragraph_format.space_after = Pt(3)


def table(headers, rows, widths=None, size=10):
    t = DOC.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    total = 16.7 if DOC.sections[-1].orientation != WD_ORIENT.LANDSCAPE else 25.4
    widths = widths or [total / len(headers)] * len(headers)
    for c, w in zip(t.columns, widths): c.width = Cm(w)
    for ridx, row in enumerate([headers] + list(rows)):
        cells = t.rows[0].cells if ridx == 0 else t.add_row().cells
        trpr = t.rows[ridx]._tr.get_or_add_trPr()
        trpr.append(x('cantSplit'))
        if ridx == 0: trpr.append(x('tblHeader'))
        for ci, (cell, text) in enumerate(zip(cells, row)):
            cell.width = Cm(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = str(text)
            pr = cell._tc.get_or_add_tcPr()
            pr.append(x('shd', fill=GRAY if ridx == 0 else ('F4F5F6' if ridx % 2 == 0 else 'FFFFFF')))
            edges = x('tcBorders')
            for e in ['top', 'left', 'bottom', 'right']:
                edges.append(x(e, val='single', sz=4, color='D9D9D9'))
            pr.append(edges)
            margins = x('tcMar')
            for edge, v in [('top', 90), ('bottom', 90), ('left', 100), ('right', 100)]:
                margins.append(x(edge, w=v, type='dxa'))
            pr.append(margins)
            for p in cell.paragraphs:
                p._p.get_or_add_pPr().append(x('snapToGrid', val='false'))
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = Pt(13.5)
                p.paragraph_format.keep_with_next = ridx == 0 or ridx == len(rows)-1 or (len(rows) <= 5 and ridx < len(rows))
                p.paragraph_format.widow_control = True
                if widths[ci] <= 1.5: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs: fonts(r, size, ridx == 0, 'FFFFFF' if ridx == 0 else '111111')
    spacer=paragraph('')
    spacer.paragraph_format.line_spacing = Pt(3)
    spacer.paragraph_format.space_after = Pt(3)
    return t


def code(text):
    lines = text.strip('\n').splitlines()
    for i, line in enumerate(lines):
        p = paragraph(line or ' ', 'Code SRS')
        p.paragraph_format.keep_with_next = i < len(lines) - 1
        p._p.get_or_add_pPr().append(x('shd', fill='F3F4F5'))
        for r in p.runs: fonts(r, 9, mono=True)
    paragraph('').paragraph_format.space_after = Pt(3)


def orientation(landscape=False):
    s = DOC.add_section(WD_SECTION.NEW_PAGE)
    s.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    s.page_width, s.page_height = (Cm(29.7), Cm(21)) if landscape else (Cm(21), Cm(29.7))
    s.left_margin = s.right_margin = Cm(2.15)
    s.top_margin = s.bottom_margin = Cm(1.9 if landscape else 2.1)
    s.header_distance = s.footer_distance = Cm(.9)
    # Only the cover has a different first-page header/footer.
    s.different_first_page_header_footer = False


def figure(path, caption, width=16.5):
    p = paragraph('', keep=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(str(path), width=Cm(width))
    p = paragraph(caption, 'Caption')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def setup():
    global DOC
    TMP.mkdir(parents=True, exist_ok=True)
    ASSETS.mkdir(parents=True, exist_ok=True)
    BACKUP.parent.mkdir(parents=True, exist_ok=True)
    if not BACKUP.exists(): shutil.copy2(OUT, BACKUP)
    DOC = Document(BACKUP)
    # Keep the user-edited author attribution, not a value from an old generator.
    author = next(p.text for p in DOC.paragraphs if '學生：' in p.text)
    for child in list(DOC.element.body):
        if child.tag != qn('w:sectPr'): DOC.element.body.remove(child)
    for st in DOC.styles:
        if st.type == 1:
            st.font.name = FONT
            st.font.color.rgb = RGBColor.from_string('111111')
            rf = st.element.get_or_add_rPr().get_or_add_rFonts()
            for a in ['ascii', 'hAnsi', 'eastAsia', 'cs']: rf.set(qn('w:' + a), FONT)
            for a in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']: rf.attrib.pop(qn('w:' + a), None)
    for n, size in [('Normal', 11), ('Title', 26), ('Heading 1', 17), ('Heading 2', 13), ('Heading 3', 11.5), ('Caption', 9)]:
        st = DOC.styles[n]
        st.font.size = Pt(size)
        st.font.bold = n in ['Title', 'Heading 1', 'Heading 2', 'Heading 3']
        st.font.color.rgb = RGBColor.from_string('000000' if n != 'Caption' else '60666D')
        st.paragraph_format.space_before = Pt(12 if n.startswith('Heading') else 0)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = Pt({'Normal':16, 'Title':34, 'Heading 1':23, 'Heading 2':19, 'Heading 3':17, 'Caption':13}[n])
        st.paragraph_format.widow_control = True
        st.paragraph_format.keep_with_next = n.startswith('Heading')
        if st.element.pPr is not None:
            for b in list(st.element.pPr):
                if b.tag in [qn('w:pBdr'), qn('w:numPr')]: st.element.pPr.remove(b)
    DOC.styles['List Bullet'].paragraph_format.left_indent = Cm(.6)
    DOC.styles['List Bullet'].paragraph_format.first_line_indent = Cm(-.3)
    if 'Code SRS' not in DOC.styles: DOC.styles.add_style('Code SRS', 1)
    cs = DOC.styles['Code SRS']
    cs.paragraph_format.line_spacing = Pt(11.5)
    cs.paragraph_format.space_after = Pt(0)
    cs.paragraph_format.space_before = Pt(0)
    cs.paragraph_format.left_indent = Cm(.25)
    s = DOC.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(2.1)
    s.left_margin = s.right_margin = Cm(2.15)
    s.header_distance = s.footer_distance = Cm(.9)
    s.different_first_page_header_footer = True
    hp = s.header.paragraphs[0]
    hp.text = 'Fitness Tracking Management System — 系統規格書'
    for r in hp.runs: fonts(r, 8, color='73777D')
    fp = s.footer.paragraphs[0]
    fp.clear(); fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fonts(fp.add_run('第 '), 9, color='73777D')
    fp._p.append(x('fldSimple', instr='PAGE'))
    fonts(fp.add_run(' 頁'), 9, color='73777D')
    for part in [s.first_page_header, s.first_page_footer]:
        for p in part.paragraphs: p.clear()
    DOC.core_properties.title = '健身紀錄管理系統 系統規格書'
    DOC.core_properties.subject = '資料庫系統課程專題'
    DOC.core_properties.author = '陳苾妏'
    DOC.core_properties.comments = ''
    paragraph('').paragraph_format.space_after = Pt(70)
    p = paragraph('Fitness Tracking\nManagement System', 'Title')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = paragraph('健身紀錄管理系統', 'Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs: fonts(r, 23, True)
    for t in ['資料庫系統規格書', 'System Requirements Specification', '資料庫系統課程專題']:
        p = paragraph(t); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
    paragraph('').paragraph_format.space_after = Pt(45)
    for t in [author, '版本 1.1　｜　2026 年 9 月 30 日']:
        p = paragraph(t); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(16)
    DOC.add_page_break()
    p = paragraph('目錄', 'Title')
    for r in p.runs: fonts(r, 19, True)
    toc_anchor = paragraph('[[TOC]]')
    DOC.add_page_break()
    return toc_anchor


def make_diagrams():
    fpath = '/Library/Fonts/Arial Unicode.ttf'
    def canvas(w, h):
        im = Image.new('RGB', (w, h), 'white')
        return im, ImageDraw.Draw(im)
    def label(d, xy, text, size=29, fill='#111111'):
        d.multiline_text(xy, text, font=ImageFont.truetype(fpath, size), fill=fill, spacing=9)
    def box(d, bounds, title, body):
        a,b,c,e = bounds
        d.rectangle(bounds, outline='#7C848B', width=2, fill='#F8F9FA')
        label(d,(a+20,b+12),title,36)
        label(d,(a+20,b+65),body,35)
    def arrow(d, pts, labeltext='', pos=None):
        d.line(pts,fill='#65717A',width=3)
        x1,y1=pts[-2]; x2,y2=pts[-1]
        if x2 == x1:
            m=1 if y2>y1 else -1
            d.polygon([(x2,y2),(x2-8,y2-15*m),(x2+8,y2-15*m)],fill='#65717A')
        else:
            m=1 if x2>x1 else -1
            d.polygon([(x2,y2),(x2-15*m,y2-8),(x2-15*m,y2+8)],fill='#65717A')
        if labeltext: label(d,pos,labeltext,24)
    im,d=canvas(1700,1090)
    box(d,(380,25,1300,170),'使用者瀏覽器','React / Vite / TypeScript / CSS　手機與電腦介面')
    box(d,(380,270,1300,450),'後端 REST API','FastAPI　JWT 與權限驗證　欄位驗證\n規則式排課　每組存檔　SQL 統計')
    box(d,(380,540,1300,710),'資料存取層','SQLAlchemy / psycopg　參數化查詢\n交易控制　FK 與 CHECK　Alembic migration')
    box(d,(380,800,1300,970),'雲端 PostgreSQL','帳號、偏好、動作、課表、訓練組數與體重紀錄')
    arrow(d,[(840,170),(840,270)],'HTTPS / JSON / Bearer JWT',(870,198))
    arrow(d,[(840,450),(840,540)],'驗證後讀寫',(880,473))
    arrow(d,[(840,710),(840,800)],'SQL / DB TLS',(880,737))
    label(d,(70,1010),'雲端主機提供 Web 與資料庫服務；首版無 Email、FCM 或 LLM 外部 API 依賴。',27)
    im.save(ASSETS/'architecture.png')
    # A full key-level ERD is split into two readable panels; users/exercises repeat as references.
    panels=[
        ('erd_account.png', [
            ('training_goals','PK training_goal_id','1 → 0..N','users','PK user_id\nFK training_goal_id'),
            ('users','PK user_id','1 → 1','user_auth','PK auth_id\nFK user_id UNIQUE'),
            ('user_auth','PK auth_id','1 → 0..N','user_sessions','PK user_session_id\nFK auth_id'),
            ('users','PK user_id','1 → 0..1','admins','PK/FK user_id'),
            ('users','PK user_id','1 → 0..N','user_body_parts','PK/FK user_id\nPK/FK body_part_id'),
            ('body_parts','PK body_part_id','1 → 0..N','user_body_parts','複合 PK (user_id, body_part_id)'),
            ('body_parts','PK body_part_id','1 → 0..N','exercises','PK exercise_id\nFK body_part_id'),
            ('users','PK user_id','1 → 0..N','body_records','PK body_record_id\nFK user_id'),
        ]),
        ('erd_workout.png', [
            ('users','PK user_id','1 → 0..N','workout_plans','PK workout_plan_id\nFK user_id / training_goal_id'),
            ('training_goals','PK training_goal_id','1 → 0..N','workout_plans','FK training_goal_id'),
            ('workout_plans','PK workout_plan_id','1 → 1..N','plan_days','PK plan_day_id\nFK workout_plan_id'),
            ('plan_days','PK plan_day_id','1 → 1..N','plan_exercises','PK plan_exercise_id\nFK plan_day_id / exercise_id'),
            ('exercises','PK exercise_id','1 → 0..N','plan_exercises','FK exercise_id'),
            ('users','PK user_id','1 → 0..N','workout_sessions','PK workout_session_id\nFK user_id / plan_day_id'),
            ('plan_days','PK plan_day_id','1 → 0..N','workout_sessions','可對應 0..1 課表日；FK 可為 NULL'),
            ('workout_sessions','PK workout_session_id','1 → 0..N','workout_sets','PK workout_set_id\nFK workout_session_id / exercise_id'),
            ('exercises','PK exercise_id','1 → 0..N','workout_sets','FK exercise_id'),
        ]),
    ]
    for name, rows in panels:
        im,d=canvas(1760, len(rows)*200+100)
        for i,(a,ak,rel,b,bk) in enumerate(rows):
            y=22+i*200
            box(d,(20,y,630,y+170),a,ak)
            box(d,(1000,y,1740,y+170),b,bk)
            arrow(d,[(630,y+78),(1000,y+78)])
            label(d,(674,y+31),rel,35)
        label(d,(24,len(rows)*200+32),'重複方塊代表同一資料表；數字表示父端與子端基數。',24)
        im.save(ASSETS/name)


# Six-column data dictionaries: name, SQL type, key, nullable, default, description/constraints.
SCHEMA = [
('users','使用者基本資料',[
('user_id','INTEGER','PK','否','IDENTITY','GENERATED BY DEFAULT AS IDENTITY；使用者識別碼。'),
('name','VARCHAR(100)','—','否','無','去頭尾空白後 1–100 字。'),
('gender','VARCHAR(20)','—','是','NULL','Female / Male / Other；未提供時 NULL。'),
('age','SMALLINT','—','是','NULL','CHECK (age BETWEEN 13 AND 100)。'),
('height_cm','NUMERIC(5,2)','—','是','NULL','CHECK (height_cm BETWEEN 100 AND 250)，單位 cm。'),
('training_experience','VARCHAR(20)','—','是','NULL','Beginner / Intermediate / Advanced。'),
('training_goal_id','INTEGER','FK','是','NULL','→ training_goals.training_goal_id；ON DELETE RESTRICT。'),
('training_days_per_week','SMALLINT','—','是','NULL','CHECK (training_days_per_week BETWEEN 2 AND 6)。'),
('training_duration_minutes','SMALLINT','—','是','NULL','CHECK (training_duration_minutes IN (30,60,90))。'),
('created_at','TIMESTAMPTZ','—','否','CURRENT_TIMESTAMP','帳號建立時間；UTC。'),
]),
('user_auth','登入認證資料',[
('auth_id','INTEGER','PK','否','IDENTITY','認證資料識別碼。'),
('user_id','INTEGER','FK','否','無','→ users.user_id；UNIQUE；ON DELETE CASCADE。'),
('email','VARCHAR(255)','—','否','無','去頭尾空白並轉小寫；UNIQUE；格式由 API 驗證。'),
('password_hash','VARCHAR(255)','—','否','無','bcrypt 雜湊值；cost = 10；不得記錄明文密碼。'),
('is_email_verified','BOOLEAN','—','否','FALSE','首版未啟用寄信驗證；不得把格式檢查當作信箱驗證。'),
('is_active','BOOLEAN','—','否','TRUE','停用後拒絕登入與所有受保護 API。'),
('created_at','TIMESTAMPTZ','—','否','CURRENT_TIMESTAMP','登入資料建立時間。'),
]),
('user_sessions','登入工作階段',[
('user_session_id','UUID','PK','否','gen_random_uuid()','作為 JWT jti；登入時每次產生不同值。'),
('auth_id','INTEGER','FK','否','無','→ user_auth.auth_id；ON DELETE CASCADE。'),
('created_at','TIMESTAMPTZ','—','否','CURRENT_TIMESTAMP','登入時間。'),
('expires_at','TIMESTAMPTZ','—','否','無','建立時間 + 2 小時；CHECK (expires_at > created_at)。'),
('revoked_at','TIMESTAMPTZ','—','是','NULL','登出或改密碼時寫入；非 NULL 即失效。'),
]),
('admins','管理員資格',[
('user_id','INTEGER','PK/FK','否','無','→ users.user_id；ON DELETE CASCADE；每位使用者最多一筆。'),
('created_at','TIMESTAMPTZ','—','否','CURRENT_TIMESTAMP','由受信任的部署／管理程序授予；前端不得自行建立。'),
]),
('training_goals','訓練目標',[
('training_goal_id','INTEGER','PK','否','IDENTITY','訓練目標識別碼。'),
('goal_name','VARCHAR(50)','—','否','無','UNIQUE；Muscle Gain / Fat Loss / Strength / General Fitness。'),
('description','TEXT','—','是','NULL','顯示用說明；最長 1,000 字。'),
]),
('body_parts','身體部位',[
('body_part_id','INTEGER','PK','否','IDENTITY','身體部位識別碼。'),
('body_part_name','VARCHAR(50)','—','否','無','UNIQUE；Chest / Back / Shoulders / Biceps / Triceps / Legs / Glutes / Core。'),
('description','TEXT','—','是','NULL','顯示用說明；最長 1,000 字。'),
]),
('user_body_parts','使用者偏好部位關聯',[
('user_id','INTEGER','PK/FK','否','無','→ users.user_id；ON DELETE CASCADE。'),
('body_part_id','INTEGER','PK/FK','否','無','→ body_parts.body_part_id；ON DELETE RESTRICT。'),
('priority','SMALLINT','—','否','1','CHECK (priority BETWEEN 1 AND 8)；數字越小越優先。'),
]),
('exercises','運動動作',[
('exercise_id','INTEGER','PK','否','IDENTITY','運動動作識別碼。'),
('exercise_name','VARCHAR(120)','—','否','無','UNIQUE；去空白後 1–120 字。'),
('body_part_id','INTEGER','FK','否','無','→ body_parts.body_part_id；ON DELETE RESTRICT；只記主要部位。'),
('difficulty_level','VARCHAR(20)','—','否','無','Beginner / Intermediate / Advanced。'),
('equipment','VARCHAR(100)','—','是','NULL','器材名稱；徒手填 Bodyweight。'),
('description','TEXT','—','否','無','1–4,000 字；包含動作步驟與注意事項。'),
('image_url','TEXT','—','是','NULL','HTTPS URL，最長 2,048 字；無圖時顯示預設示意。'),
('is_active','BOOLEAN','—','否','TRUE','FALSE 不再作為新課表候選；歷史仍可查詢。'),
]),
('workout_plans','訓練課表',[
('workout_plan_id','INTEGER','PK','否','IDENTITY','課表識別碼。'),
('user_id','INTEGER','FK','否','無','→ users.user_id；ON DELETE CASCADE。'),
('plan_name','VARCHAR(120)','—','否','無','1–120 字；預設名稱由 API 建立。'),
('training_goal_id','INTEGER','FK','否','無','→ training_goals.training_goal_id；ON DELETE RESTRICT；代表產生當時的目標。'),
('created_at','TIMESTAMPTZ','—','否','CURRENT_TIMESTAMP','建立時間；既有課表不隨個人目標變更而改寫。'),
]),
('plan_days','課表訓練日',[
('plan_day_id','INTEGER','PK','否','IDENTITY','課表日識別碼。'),
('workout_plan_id','INTEGER','FK','否','無','→ workout_plans.workout_plan_id；ON DELETE CASCADE。'),
('day_number','SMALLINT','—','否','無','CHECK (day_number BETWEEN 1 AND 6)。'),
('day_title','VARCHAR(120)','—','否','無','1–120 字；例如「第 1 天：胸、背、腿」。'),
]),
('plan_exercises','課表中的動作',[
('plan_exercise_id','INTEGER','PK','否','IDENTITY','課表動作識別碼。'),
('plan_day_id','INTEGER','FK','否','無','→ plan_days.plan_day_id；ON DELETE CASCADE。'),
('exercise_id','INTEGER','FK','否','無','→ exercises.exercise_id；ON DELETE RESTRICT。'),
('exercise_order','SMALLINT','—','否','無','CHECK (exercise_order BETWEEN 1 AND 6)。'),
('target_sets','SMALLINT','—','否','無','CHECK (target_sets BETWEEN 1 AND 10)。'),
('target_reps','SMALLINT','—','否','無','CHECK (target_reps BETWEEN 1 AND 50)；首版用單一目標次數。'),
('rest_seconds','SMALLINT','—','否','無','CHECK (rest_seconds BETWEEN 0 AND 300)。'),
]),
('workout_sessions','實際訓練場次',[
('workout_session_id','INTEGER','PK','否','IDENTITY','實際訓練場次識別碼；不同於登入工作階段。'),
('user_id','INTEGER','FK','否','無','→ users.user_id；ON DELETE CASCADE。'),
('plan_day_id','INTEGER','FK','是','NULL','→ plan_days.plan_day_id；ON DELETE SET NULL；課表刪除時保留歷史。'),
('session_date','DATE','—','否','無','API 以 Asia/Taipei 的開始日期寫入。'),
('status','VARCHAR(20)','—','否',"'in_progress'",'CHECK：in_progress / completed / cancelled。'),
('started_at','TIMESTAMPTZ','—','否','CURRENT_TIMESTAMP','開始時間；UTC。'),
('ended_at','TIMESTAMPTZ','—','是','NULL','完成或取消時間；終止狀態必填且不得早於 started_at。'),
('notes','TEXT','—','是','NULL','0–1,000 字。'),
]),
('workout_sets','每組實際訓練紀錄',[
('workout_set_id','INTEGER','PK','否','IDENTITY','每組訓練識別碼。'),
('workout_session_id','INTEGER','FK','否','無','→ workout_sessions.workout_session_id；ON DELETE CASCADE。'),
('exercise_id','INTEGER','FK','否','無','→ exercises.exercise_id；ON DELETE RESTRICT。'),
('set_number','SMALLINT','—','否','無','CHECK (set_number BETWEEN 1 AND 100)。'),
('weight_kg','NUMERIC(6,2)','—','否','無','CHECK (weight_kg BETWEEN 0 AND 999.99)；外加重量，0 表示徒手。'),
('reps','SMALLINT','—','否','無','CHECK (reps BETWEEN 1 AND 999)。'),
('created_at','TIMESTAMPTZ','—','否','CURRENT_TIMESTAMP','首次存檔時間。'),
]),
('body_records','體重歷史紀錄',[
('body_record_id','INTEGER','PK','否','IDENTITY','身體紀錄識別碼。'),
('user_id','INTEGER','FK','否','無','→ users.user_id；ON DELETE CASCADE。'),
('record_date','DATE','—','否','無','API 驗證不得晚於 Asia/Taipei 今日；每日一筆。'),
('weight_kg','NUMERIC(5,2)','—','否','無','CHECK (weight_kg BETWEEN 20 AND 500)。'),
('notes','TEXT','—','是','NULL','0–1,000 字。'),
]),
]


def fr(fid, title, trigger, inputs, logic, success, errors, db):
    heading(fid + ' ' + title, 3)
    for label, text in [('觸發條件',trigger),('輸入資料',inputs)]:
        p=paragraph(label+'：'+text,keep=True)
        p.paragraph_format.space_after=Pt(4)
    paragraph('處理邏輯：',keep=True)
    steps(logic)
    paragraph('輸出／回應：'+success,keep=True)
    paragraph('例外處理：'+errors,keep=True)
    paragraph('資料庫對應：'+db)


def requirements():
    heading('3. 功能性需求')
    paragraph('本章以 FR 編號識別需求。所有以 /api 開頭的端點均為本版預定介面契約；驗收須依第 6 章執行，文件中的目標值不代表已通過實測。')
    heading('3.1 共通規則',2)
    bullets([
        '受保護 API 由 Authorization: Bearer <JWT> 取得 user_id，禁止以請求中的 user_id 作為授權依據。查詢子表時必須 JOIN 回擁有者。',
        '查無資料或資料不屬於本人均回傳 404；缺少、過期、撤銷的 JWT 回傳 401；一般會員呼叫管理端點回傳 403。',
        '欄位型別、長度、範圍或日期錯誤回傳 422；資源重複或狀態衝突回傳 409；成功查詢／更新回傳 200，新增 201，刪除 204。',
        '錯誤 JSON 固定含 code、message；422 另含 field_errors。清單預設 page=1、page_size=20，page_size 範圍 1–100。',
        '多表新增與關聯替換使用同一資料庫交易；COMMIT 成功後才回覆成功。失敗時 ROLLBACK，介面保留未送出的輸入並提供重試。',
        '時間戳以 UTC 的 TIMESTAMPTZ 儲存；日期與每週統計以 Asia/Taipei 為準，每週由星期一開始。',
    ])
    heading('3.2 會員與權限管理',2)
    fr('FR-01','使用者註冊','點擊「註冊」。','name 1–100 字、email 最長 255 字、password 10–72 UTF-8 bytes，至少含英文字母與數字。',[
        'Email 去除頭尾空白並轉小寫，再驗證格式及唯一性；不接受用戶提交管理員角色。',
        '密碼以 bcrypt cost=10 雜湊；同一交易建立 users、user_auth 與 user_sessions。註冊後未建立 admins 紀錄，角色為 USER。',
        'JWT 使用 HS256，包含 sub=user_id、jti=user_session_id、role、iat、exp；有效時間 7,200 秒。交易成功後回傳。',
    ],'POST /api/auth/register → 201；回傳 user_id、access_token、token_type=Bearer、expires_in=7200。','重複 Email 回傳 409 EMAIL_EXISTS；欄位不合法回傳 422；交易失敗不得留下孤立 users。','users、user_auth、user_sessions；TC-01、TC-02、TC-03。')
    paragraph('Email 驗證範圍：首版驗證 Email 格式與唯一性，但不寄送驗證信，也不宣稱已確認信箱持有者。is_email_verified 預設 FALSE，不阻擋登入。寄信驗證與忘記密碼流程需要 Email service、驗證 token、有效期限及寄送重試，列為後續版本；本版不建立這些服務。')
    fr('FR-02','登入與登出','點擊「登入」或「登出」。','登入輸入 email、password；登出附帶 Bearer JWT。',[
        '依正規化 Email 查詢 user_auth，以 bcrypt 驗證密碼並檢查 is_active。成功登入時建立一筆 user_sessions，過期時間為建立時間加 2 小時。',
        '每次受保護請求均驗證 JWT 簽章、exp、jti、user_sessions.revoked_at 及帳號狀態。role 由 admins 資格推導，管理權限以資料庫即時結果為準。',
        '登出時將對應工作階段的 revoked_at 設為目前時間，前端移除 Token；相同 Token 不得再次讀取受保護資料。Token 只放在前端記憶體，重新載入頁面後需重新登入。',
    ],'POST /api/auth/login → 200 與 JWT；POST /api/auth/logout → 204。','錯誤帳密統一回 401 INVALID_CREDENTIALS；停用帳號、過期或已登出 Token 回 401。','user_auth、user_sessions、admins；TC-04、TC-05。')
    fr('FR-03','變更密碼與停用自己的帳號','在帳號設定確認變更密碼或停用。','current_password；變更密碼另傳 new_password（規則同 FR-01）。',[
        '驗證本人目前密碼；修改密碼時寫入新 password_hash，並撤銷該帳號所有登入工作階段。',
        '停用時將 user_auth.is_active 設為 FALSE 並撤銷所有工作階段；保留課表及歷史資料，但禁止存取。此操作不等於物理刪除。',
        '管理員不得透過此介面變更他人密碼或停用他人；首版帳號停用不提供前端復原。',
    ],'PATCH /api/auth/password 或 POST /api/auth/deactivate → 204；前端回到登入頁。','目前密碼錯誤回 401；新密碼不合法回 422。','user_auth、user_sessions；TC-06、TC-07。')
    heading('3.3 使用者基本資料管理',2)
    fr('FR-04','建立查看與修改基本資料','首次登入填寫資料，或在個人資料頁儲存。','name；gender 選填；age 13–100；height_cm 100–250；weight_kg 20–500；training_experience 為三個固定等級。',[
        'GET /api/me 僅載入本人；尚未填寫的欄位回傳 NULL。設定完成前可以登入，但不能產生課表。',
        'PUT /api/me/profile 驗證所有必填值，更新 users；體重寫入今日的 body_records，當日已有一筆則更新該筆，與 users 更新放入同一交易。',
        '目前體重由 body_records 中最新 record_date 的資料取得，users 不另存 weight_kg，以避免兩份體重不一致。',
    ],'GET 或 PUT → 200；回傳 profile 與目前體重。','缺少必填資料、範圍或等級錯誤回 422；不得修改其他使用者。','users、body_records；TC-08。')
    heading('3.4 訓練目標與偏好管理',2)
    fr('FR-05','設定訓練目標天數與時間','在訓練設定點擊儲存。','training_goal_id；training_days_per_week=2–6；training_duration_minutes ∈ {30,60,90}。',[
        '確認 training_goal_id 存在；四種目標為增肌、減脂、力量與一般健身。',
        '在 users 更新選項；已產生課表保留原 training_goal_id，重新產生才套用新設定。',
    ],'PUT /api/me/training-settings → 200。','不存在的目標或無效天數／時間回 422。','users、training_goals；TC-09。')
    fr('FR-06','設定多個想加強部位','勾選部位並儲存。','body_part_ids：1–8 個互不重複的整數；陣列順序為優先順序。',[
        '先驗證所有 ID 存在且不重複；在一個交易刪除本人舊關聯，再新增新的 user_body_parts。priority 按陣列順序設為 1 至 N。',
        '任一 ID 失效時回滾整筆交易；複合主鍵防止重複部位，不使用逗號字串欄位。',
    ],'PUT /api/me/body-parts → 200；回傳已保存部位。','空陣列、重複 ID、無效 ID 或超過 8 個回 422；原資料保留。','user_body_parts、body_parts；TC-10。')
    heading('3.5 運動動作與參考資料管理',2)
    fr('FR-07','瀏覽搜尋與動作說明','開啟動作資料庫或點擊課表中的動作。','keyword 0–100 字；可選 body_part_id、difficulty_level；分頁參數。',[
        '以參數化查詢 JOIN exercises 與 body_parts；keyword 比對動作與部位名稱，預設只顯示 is_active=TRUE，依 exercise_name、exercise_id 排序。',
        '回傳主要部位、難度、器材、動作說明與 image_url；圖片載入失敗時顯示本機預設示意，文字仍可閱讀。',
    ],'GET /api/exercises 或 GET /api/exercises/{id} → 200；無搜尋結果回空陣列。','非法篩選／分頁回 422；不存在的動作回 404。','exercises、body_parts；TC-11。')
    fr('FR-08','管理運動動作','管理員新增編輯或移除動作。','exercise_name、body_part_id、difficulty_level、equipment、description、image_url、is_active；限制見 4.2。',[
        '確認登入者於 admins 有資格；驗證 FK、名稱唯一性、等級與 HTTPS 圖片 URL。後端只保存 URL，不下載使用者提供的檔案。',
        '若 exercise 未被課表或訓練組引用，DELETE 可物理刪除；已被引用時回 409，管理員改用 PATCH is_active=FALSE。歷史依 exercise_id 仍可查詢。',
        '若已有 workout_sets 引用，不得修改主要 body_part_id，以免歷史部位統計被改寫；需要重新分類時新增動作版本並停用原動作。',
    ],'POST /api/admin/exercises → 201；PATCH → 200；DELETE → 204。','非管理員回 403；重複名稱、受引用刪除或修改部位回 409；輸入不合法回 422。','admins、exercises、body_parts、plan_exercises、workout_sets；TC-12。')
    fr('FR-09','維護訓練目標與身體部位說明','管理員編輯參考資料說明。','training_goal_id 或 body_part_id；description 最長 1,000 字。',[
        '初始 migration/seed 建立四個目標與八個部位；首版名稱與集合固定，以便排課規則可重現。',
        '管理員可讀取與修改 description；前端不提供新增、改名或刪除固定分類。此類擴充須連同規則與 migration 一起調整。',
    ],'GET /api/reference → 200；PATCH /api/admin/training-goals/{id} 或 /api/admin/body-parts/{id} → 200。','非管理員回 403；ID 不存在回 404；不支援的欄位回 422。','training_goals、body_parts、admins；TC-13。')
    heading('3.6 訓練課表產生與管理',2)
    fr('FR-10','產生建議課表','完成設定後點擊「產生我的訓練課表」。','無需傳 user_id；選用 plan_name 1–120 字。',[
        '以登入者讀取基本資料、目標、偏好、天數與時間；任一必要設定缺漏時拒絕產生。',
        '依附錄 C 的固定規則查詢啟用且難度不高於使用者等級的 exercises；配置每日動作順序、組數、次數與休息。',
        '後端驗證天數、候選 ID、難度、每週主要部位覆蓋及時間預算後，以單一交易新增 workout_plans、plan_days、plan_exercises。每次成功操作產生新課表，不覆蓋舊課表。',
    ],'POST /api/plans/generate → 201，回傳 workout_plan_id、每日動作、部位、圖片、組次、休息與預估時間。','設定不完整回 422 PROFILE_INCOMPLETE；候選不足或約束無法滿足回 409 PLAN_CONSTRAINT_UNSATISFIED；不得保存半份課表。','users、user_body_parts、exercises、workout_plans、plan_days、plan_exercises；TC-14、TC-15。')
    fr('FR-11','查看修改與刪除課表','進入「我的課表」，編輯或確認刪除。','workout_plan_id；可修改 plan_name 與各 plan_exercise 的 exercise_id、target_sets、target_reps、rest_seconds。',[
        '讀取時 JOIN plan_days、plan_exercises、exercises 與 body_parts，依 day_number、exercise_order 排序。',
        '更新時驗證本人所有權、難度、啟用狀態與組次範圍；同日 exercise 不重複，估計時間不得超過本人設定。',
        '該課表存在進行中訓練時，禁止修改與刪除；其餘時段刪除會 CASCADE 刪除 days／plan_exercises，歷史 workout_sessions.plan_day_id 設為 NULL，保留 sets。',
    ],'GET /api/plans、GET/PATCH /api/plans/{id} → 200；DELETE /api/plans/{id} → 204。','非本人或不存在回 404；進行中引用回 409；無效動作、重複或超時回 422。','workout_plans、plan_days、plan_exercises、workout_sessions；TC-16。')
    heading('3.7 開始運動與每組紀錄',2)
    fr('FR-12','開始或繼續運動','在當日課表點擊「開始運動」。','plan_day_id。',[
        '驗證課表日屬於本人；鎖定使用者資料，查詢該使用者是否已有 in_progress 場次。',
        '無進行中場次時 INSERT workout_sessions，寫入 user_id、plan_day_id、台北日期、started_at，status=in_progress。',
        '相同課表日已有進行中場次時回傳該筆與已存 sets；不同課表日已有場次時要求先完成或取消。部分唯一索引防止並行點擊建立兩筆。',
    ],'POST /api/workout-sessions → 新增 201，續練 200；回傳 workout_session_id、今日動作及已存組數。GET /api/workout-sessions/current 可取得進行中場次。','課表日不存在／非本人回 404；另一日進行中回 409 WORKOUT_ALREADY_ACTIVE。','workout_sessions、plan_days、workout_plans；TC-17。')
    fr('FR-13','每完成一組立即儲存','完成一組後按「儲存此組」。','workout_session_id、exercise_id、set_number 1–100、weight_kg 0–999.99、reps 1–999。',[
        '鎖定本人 workout_sessions，確認 status=in_progress；exercise_id 必須存在於該課表日，並驗證數值範圍。',
        '以 (workout_session_id, exercise_id, set_number) 唯一鍵保存一組；相同資料重送回傳既有結果，不多建一組。同鍵不同值須使用 FR-14 修改。',
        '每組獨立 COMMIT 後才顯示「已儲存」；斷線或失敗顯示「未儲存」並保留欄位。重新登入後可載入已存資料。',
    ],'POST /api/workout-sessions/{id}/sets → 首次 201、相同資料重送 200；回傳 workout_set_id、set_volume=weight_kg×reps。','場次完成／取消、同鍵不同值回 409；動作或數值無效回 422；非本人回 404。','workout_sessions、workout_sets、plan_exercises；TC-18、TC-19。')
    fr('FR-14','修正或刪除每組紀錄','在訓練頁或歷史詳情編輯錯誤組數。','workout_set_id；PATCH 可修改 weight_kg、reps。',[
        'JOIN workout_sessions 確認擁有者並鎖定場次；允許修改 in_progress 或 completed 的數值，不改 exercise_id 或 set_number。',
        '取消場次不可改組數；DELETE 刪除指定一組。已完成場次至少保留一組，刪除最後一組時改用刪除整場訓練。',
        'Dashboard 由實際資料重新彙總，下一次查詢必須反映更新／刪除結果。',
    ],'PATCH /api/workout-sets/{id} → 200；DELETE → 204。','非本人回 404；取消場次或刪除完成場次最後一組回 409；數值無效回 422。','workout_sets、workout_sessions；TC-20。')
    fr('FR-15','完成取消或刪除訓練','點擊「完成訓練」「取消本次訓練」或確認刪除歷史。','workout_session_id；notes 選填、最長 1,000 字。',[
        '完成時鎖定本人場次並確認至少一組；status 由 in_progress 改為 completed，ended_at 設為目前時間。連續點擊完成回傳既有結果。',
        '取消時由 in_progress 改為 cancelled，保留已存 sets，但排除於 Dashboard；完成與取消互不可轉換。',
        '只允許刪除 completed／cancelled 場次；同一交易刪除場次及其 sets。進行中場次需先取消。',
    ],'POST /api/workout-sessions/{id}/complete 或 /cancel → 200；DELETE /api/workout-sessions/{id} → 204。','空場次完成、非法狀態轉換或刪除進行中場次回 409；非本人回 404。','workout_sessions、workout_sets；TC-21、TC-22。')
    heading('3.8 歷史紀錄與統計分析',2)
    fr('FR-16','查看訓練與動作歷史','進入訓練歷史或查看某動作歷史。','start_date、end_date、exercise_id、status 選填；分頁參數。',[
        'Workout History 預設列出 completed 與 cancelled，進行中紀錄放於 current；只查本人並依 session_date、workout_session_id 倒序。',
        '詳情 JOIN workout_sets、exercises、body_parts 顯示每組重量、次數、組號與量；Exercise History 加上 exercise_id 篩選。課表刪除後仍能顯示實際動作。',
    ],'GET /api/workout-sessions 或 /api/exercises/{id}/history → 200；含項目與分頁資訊，無資料回空陣列。','日期起訖錯誤／超過 366 天回 422；非本人明細回 404。','workout_sessions、workout_sets、exercises、body_parts；TC-23。')
    fr('FR-17','訓練量與 Dashboard','開啟統計儀表板或變更日期篩選。','start_date、end_date；預設含今日的最近 28 天，最多 366 天。',[
        '只計入本人 status=completed 的場次；每組訓練量為外加重量 kg × 次數。徒手紀錄的訓練量為 0，不推估體重負荷。',
        '總量使用 SUM(weight_kg*reps)，每週次數使用 COUNT(DISTINCT workout_session_id)，以台北日期及週一分組。',
        '最常部位及動作以完成組數 COUNT(*) 排名，平手依對應 ID 升序；體重歷史讀取 body_records，平均值用 AVG(weight_kg)。',
        '無訓練資料回總量 0、次數 0、排行榜空陣列；無體重紀錄回 NULL 與空趨勢。',
    ],'GET /api/dashboard → 200；回傳總量、每週次數、部位與動作排行、體重趨勢。','日期或範圍不合法回 422；尚未完成的組數不得影響正式統計。','workout_sessions、workout_sets、exercises、body_parts、body_records；TC-24。')
    heading('3.9 身體紀錄管理',2)
    fr('FR-18','新增查看修改與刪除體重','進入體重歷史並新增、修正或刪除。','record_date≤台北今日；weight_kg 20–500；notes≤1,000 字；body_record_id 用於修改／刪除。',[
        '每位使用者每天一筆；新增若日期已存在回 409，提示使用 PATCH 修改。',
        '查詢依 record_date 排序；修改與刪除均驗證本人。刪除最新紀錄後，目前體重回退到上一筆；若已無紀錄則回 NULL。',
    ],'POST /api/body-records → 201；GET /api/body-records、PATCH /api/body-records/{id} → 200；DELETE → 204。','重複日期回 409；未來日期或不合法體重回 422；非本人回 404。','body_records、users；TC-25。')


def database():
    heading('4. 資料模型與結構設計')
    paragraph('本版包含 14 張資料表。users 保存會員資料，user_auth 保存登入憑證，user_sessions 保存可撤銷登入狀態，admins 表示管理員資格。workout_sessions 與 user_sessions 分別代表運動與登入，兩者不得混用。')
    heading('4.1 實體關係圖',2)
    figure(ASSETS/'erd_account.png','圖 2　實體關係圖：帳號、偏好與動作資料')
    figure(ASSETS/'erd_workout.png','圖 3　實體關係圖：課表與訓練紀錄')
    paragraph('users 與 body_parts 透過 user_body_parts 形成 M:N；plan_days 與 exercises 透過 plan_exercises 形成 M:N。users 對 user_auth 的業務規則為 1:1，FK UNIQUE 保證最多一筆，註冊交易保證至少一筆。每份課表及每個訓練日的最低子項數由產生課表交易驗證。')
    orientation(True)
    heading('4.2 資料表詳細規格',2)
    paragraph('以下採老師範本的六欄格式。IDENTITY 為 PostgreSQL 自動流水號；「無」表示沒有資料庫 DEFAULT，由 API 必須提供；NULL 僅供選填或尚未完成設定。欄位範圍與列舉均建立 CHECK，必填欄位均建立 NOT NULL。')
    paragraph('設計修正：目前體重改由 body_records 的最新紀錄讀取；plan_exercises.target_reps 改為整數目標次數；補齊原稿已提及的 admins 與 user_sessions。這些欄位與關係以本章為本版規格，實作時須透過 migration 套用。')
    for i,(name,title,rows) in enumerate(SCHEMA,1):
        heading(f'表 4.2.{i}　{title} {name}',3)
        display_rows=[list(row) for row in rows]
        for row in display_rows:
            if row[4]=='CURRENT_TIMESTAMP': row[4]='CURRENT_\nTIMESTAMP'
        table(['欄位名稱','資料型態','PK/FK','允許\nNULL','預設值','說明與約束'],display_rows,[4.8,3.2,1.65,1.4,4.0,10.35],10)
    orientation(False)
    heading('4.3 關聯唯一性與刪除規則',2)
    table(['約束／關係','規格與用途'],[
        ('user_body_parts 複合主鍵','PRIMARY KEY (user_id, body_part_id)；同一使用者同一部位只存在一筆。'),
        ('課表日／動作唯一性','UNIQUE (workout_plan_id, day_number)；UNIQUE (plan_day_id, exercise_order)；UNIQUE (plan_day_id, exercise_id)。'),
        ('訓練組與體重唯一性','UNIQUE (workout_session_id, exercise_id, set_number)；UNIQUE (user_id, record_date)。'),
        ('同時只有一場進行中訓練',"workout_sessions(user_id) 的部分唯一索引，WHERE status='in_progress'。"),
        ('訓練狀態與時間',"CHECK：in_progress 時 ended_at IS NULL；completed/cancelled 時 ended_at IS NOT NULL AND ended_at >= started_at。"),
        ('使用者關聯 CASCADE','物理清除 users 時，auth、sessions、admins、偏好、課表、場次、體重沿 FK 刪除；首版前端僅提供停用帳號。'),
        ('課表關聯 CASCADE / SET NULL','刪除課表連帶刪除 days、plan_exercises；歷史場次保留，plan_day_id 設為 NULL。'),
        ('動作與參考資料 RESTRICT','動作被課表／sets 引用時拒絕物理刪除；目標與部位被引用時拒絕刪除。'),
        ('擁有者與跨表規則','API 必須 JOIN 驗證課表日／訓練組屬於登入者；一般 FK 無法單獨保證跨表所有權。'),
    ],[5.3,11.4])
    heading('4.4 交易索引與資料管理',2)
    bullets([
        '註冊、設定偏好、儲存完整課表與刪除場次皆為原子交易；注入錯誤時驗證 ROLLBACK 後各表筆數不變。',
        '新增／修改訓練組與完成／取消訓練須先以 SELECT … FOR UPDATE 鎖定同一 workout_sessions，避免完成後又插入組數。',
        'PK 與 UNIQUE 由 PostgreSQL 自動建立索引，勿重複建立 email 與 user_body_parts(user_id) 索引。EXPLAIN (ANALYZE, BUFFERS) 在測試資料庫驗證查詢計畫。',
        'Alembic 管理 schema 版本；seed 在重跑時不得新增重複目標、部位或動作。應用程式資料庫帳號僅取得必要 DML 權限，migration 使用部署專用帳號。',
    ])
    table(['索引','欄位與用途'],[
        ('idx_exercises_body_difficulty','exercises(body_part_id, difficulty_level)，WHERE is_active=TRUE；查詢候選動作。'),
        ('idx_sessions_user_date','workout_sessions(user_id, session_date DESC)；使用者日期篩選。'),
        ('uq_active_workout_per_user',"UNIQUE workout_sessions(user_id) WHERE status='in_progress'；防重複開始。"),
        ('uq_workout_sets_session_exercise_no','唯一鍵左側包含 workout_session_id，可讀取整場組數，不另建相同前綴索引。'),
        ('idx_sets_exercise','workout_sets(exercise_id)；配合 ownership JOIN 查動作歷史及引用。'),
        ('uq_body_records_user_date','UNIQUE (user_id, record_date)，支援體重趨勢及最新紀錄。'),
        ('idx_user_sessions_auth','user_sessions(auth_id, expires_at)；查詢登入階段及整批撤銷。'),
    ],[6.7,10])


TESTS = [
('TC-01','FR-01 註冊成功','Email new@example.test 尚未存在。','1. 填姓名 Test 與未使用 Email。\n2. 輸入符合規則的測試密碼並註冊。','201；users/user_auth/user_sessions 各 +1；admins 不新增；password_hash 不等於輸入值；JWT exp−iat=7200。'),
('TC-02','FR-01 重複註冊','已有 new@example.test。','1. 使用 NEW@example.test 再註冊。\n2. 比較前後三張表筆數。','409 EMAIL_EXISTS；畫面顯示 Email 已使用；所有表均無新增。'),
('TC-03','FR-01 交易與欄位','可注入 DB 失敗的測試環境。','1. 送出 3 字密碼。\n2. 合法註冊但令 user_auth INSERT 失敗。','步驟 1 為 422；步驟 2 為錯誤回應並回滾；無孤立 users。'),
('TC-04','FR-02 登入／保護','已有啟用帳號 A。','1. 錯誤密碼登入。\n2. 正確密碼登入。\n3. 不帶 Token 呼叫 GET /api/me。','依序 401、200、401；僅成功登入增加 1 個 user_sessions。'),
('TC-05','FR-02 登出與期限','已登入並保存原 JWT。','1. 登出再用原 Token 查 /api/me。\n2. 用過期 2 小時的 Token 查詢。','均 401；登出後 revoked_at 非 NULL；不回傳個人資料。'),
('TC-06','FR-03 變更密碼','帳號 A 有兩個有效 Token。','1. 輸入目前與新密碼。\n2. 使用兩個舊 Token。\n3. 以舊／新密碼登入。','204；舊 Token 全部 401；舊密碼失敗、新密碼成功；hash 更新。'),
('TC-07','FR-03 停用帳號','帳號 A 已登入且有課表。','1. 提交正確目前密碼並停用。\n2. 嘗試重新登入。','204 後 is_active=FALSE、登入 401；課表與歷史仍保留。'),
('TC-08','FR-04 個人資料','新帳號 A；今日無體重。','1. 填 age=21、height=160、weight=50、Beginner。\n2. 重載 GET /api/me。','200；users 更新、body_records 今日 +1；返回體重 50；users 無重複體重欄位。'),
('TC-09','FR-05 訓練設定','A 已完成基本資料。','1. 設增肌、4 天、60 分鐘。\n2. 改送 7 天或 45 分鐘。','步驟 1 保存成功；步驟 2 為 422，原 4 天／60 分鐘不變。'),
('TC-10','FR-06 多選部位','A 原有一個偏好部位。','1. 保存 Chest、Back。\n2. 再送含不存在 ID 的陣列。','成功時 junction 有 2 筆且 priority=1/2；失敗為 422，兩筆舊關聯保留。'),
('TC-11','FR-07 搜尋圖片','存在啟用胸部動作及無效圖片 URL。','1. 篩 Chest 與 Beginner。\n2. 點開動作。\n3. 搜尋不存在關鍵字。','只回符合條件動作；文字可讀、失效圖片顯示預設圖；不存在結果為 200 空陣列。'),
('TC-12','FR-08 動作管理','管理員 M、會員 A；動作 E 已被 sets 引用。','1. A 嘗試新增。\n2. M 新增後刪除未引用動作。\n3. M 刪 E，再停用 E。','依序 403、201/204、409/200；E 停用後仍可讀歷史、不能成新課表候選。'),
('TC-13','FR-09 參考管理','管理員 M 與會員 A。','1. M 更新 Core 說明。\n2. A 做相同修改。\n3. M 嘗試改分類名稱。','200 且 DB 說明更新；A 為 403；名稱變更 422，固定八個部位不變。'),
('TC-14','FR-10 課表產生','Beginner、增肌、Chest/Back、4 天／60 分鐘；每部位有兩個初階候選。','1. 產生課表。\n2. JOIN 檢查日與動作。','201；1 plan、4 days；每日至少 3 動作且估計 ≤60 分鐘；八部位週覆蓋、無高於初階動作。'),
('TC-15','FR-10 候選／回滾','測試環境暫無任何啟用 Core 候選。','1. 產生課表。\n2. 比較 plans/days/exercises 筆數。','409 PLAN_CONSTRAINT_UNSATISFIED；三張課表表均不新增；回傳缺少部位。'),
('TC-16','FR-11 課表 CRUD','A 有課表 P 與已完成場次 S。','1. 改 P 名稱與目標次數。\n2. 刪 P。\n3. 查 S 詳情。','PATCH 200；DELETE 204；P/day/plan_exercises 消失；S.plan_day_id=NULL，sets 與歷史仍存在。'),
('TC-17','FR-12 開始與續練','A 有有效課表日 D，沒有進行中場次。','1. 同時送兩次開始 D。\n2. 再次載入 current。\n3. 開始另一日。','只一筆 in_progress；重複請求回同 ID；另一日回 409；current 可載入已存 sets。'),
('TC-18','FR-13 逐組存檔','A 有進行中場次 S 與胸推 E。','1. 送 E 第 1 組 60kg×12。\n2. 以相同內容重送。\n3. 重新登入後讀 S。','首次 201、重送 200；sets 只 +1；set_volume=720；重新登入仍看到該組。'),
('TC-19','FR-13 非法資料','S 已存在第 1 組。','1. 送 weight=-1 或 reps=0。\n2. 同組號改送 65kg×10。','無效數值 422；同鍵不同值 409；sets 原值與筆數不變。'),
('TC-20','FR-14 修改刪組','完成場次有 60×12、65×10、65×8 三組。','1. 將第 3 組改為 65×9。\n2. 刪除第 2 組。','PATCH 200；DELETE 204；剩 2 組；總量 720+585=1305；非本人操作 404。'),
('TC-21','FR-15 完成與競態','A 場次含至少一組。','1. 完成訓練再重送完成。\n2. 同時測試新增組與完成。','completed 且 ended_at 有值；重送不新增；鎖序使新增先提交或被 409 拒絕，完成後不再新增。'),
('TC-22','FR-15 取消刪場','A 有空的進行中場次。','1. 完成空場次。\n2. 取消後刪除。','完成回 409；取消 200；刪除 204 且子 sets 一併刪除；不列入 Dashboard。'),
('TC-23','FR-16 歷史與權限','A、B 各有完成場次。','1. A 查看自己的日期／動作歷史。\n2. A 用 B 的場次 ID 查詳情。','A 只見本人正確 JOIN 結果；B 明細為 404；移除 Token 回 401。'),
('TC-24','FR-17 精確統計','同週 completed S1：60×12、65×10、65×8；S2：20×10；另有取消場次 100×10。','1. 查看該週 Dashboard。\n2. 查無資料日期區間。','總量 2090、每週次數 2；取消場次不計；胸推三組若最多則排名第一；空區間回 0／空陣列。'),
('TC-25','FR-18 體重 CRUD','A 今日無體重紀錄。','1. 新增 50kg、同日再新增。\n2. 改成 51kg、查詢、刪除。\n3. 送未來日期。','依序 201、409、200、200、204、422；趨勢與目前體重隨有效紀錄更新。'),
]


def nonfunctional():
    heading('5. 非功能性需求')
    heading('5.1 效能需求',2)
    paragraph('下列為驗收門檻，並非已完成的量測結果。測試使用預先喚醒的雲端服務、同區資料庫，以及 1,000 位會員、10,000 場訓練與 100,000 組紀錄；測試報告需記錄實際 CPU、記憶體、DB 方案、資料量與測試時間。')
    table(['編號','量化門檻','驗證方式'],[
        ('NFR-01','一般讀取 API P95 ≤500 ms；每組儲存 P95 ≤1,000 ms；產生課表 P95 ≤3,000 ms。','API 用戶端量測完整 HTTP 回應，排除測試前 60 秒暖機；依端點分別計算。'),
        ('NFR-02','200 個並行虛擬使用者連續操作 5 分鐘；預期成功請求的非預期失敗率 <1%。','30 秒升至 200 人；每人請求後停 1 秒；讀取／存組／產課表比例 70/20/10，不得崩潰或違反唯一鍵。'),
        ('NFR-03','收到儲存成功回應後，重新載入可讀到相同紀錄；不接受假成功。','另一次 HTTP 查詢與 DB SELECT 交叉確認；故障注入不得留下部分資料。'),
    ],[1.8,7.8,7.1])
    paragraph('部署冷啟動時間另列測試結果；正式效能門檻於暖機後量測。免費方案若無法達標，必須在驗收報告揭露實測值，並以學校資源或符合需求的雲端規格重測，不能以未測試結果宣稱通過。')
    heading('5.2 安全性需求',2)
    bullets([
        'NFR-04：公開站點強制 HTTPS；應用程式連線雲端資料庫使用 TLS。HTTP 入口重新導向 HTTPS。',
        'NFR-05：密碼只存 bcrypt hash；JWT 使用環境變數中的 SECRET_KEY，效期為 2 小時；登出或改密碼後舊工作階段立即被拒絕。',
        'NFR-06：SQL 使用參數綁定，排序欄位使用允許清單；每個個人資源檢查所有權，管理端點確認 admins 資格。SQL 注入輸入不能擴大結果集。',
        'NFR-07：DATABASE_URL、SECRET_KEY、API key 與部署密碼不得進 Git；.env.example 僅列空值或非機密範例。日誌不得輸出明文密碼、JWT 或完整連線字串。',
        'NFR-08：同 IP 於 15 分鐘內最多 20 次失敗登入，第 21 次回 429；成功登入不解除該 IP 的限制計數。',
    ])
    heading('5.3 易用性與相容性',2)
    bullets([
        'NFR-09：驗收當日的 Chrome、Safari、Edge 各最新兩個主要版本均須完成註冊、產生課表與逐組存檔；報告記錄實際版本。',
        'NFR-10：375×812、768×1024 與 1440×900 三種視窗下，不出現整頁水平捲動；主要按鈕最小可點擊範圍為 44×44 CSS px。',
        'NFR-11：已具備訓練設定的使用者由課表頁至首組輸入最多 2 次主要按鈕操作；成功、失敗與未存檔狀態必須以文字顯示。',
        'NFR-12：欄位具有可辨識標籤、錯誤提示與鍵盤焦點順序；圖片提供替代文字，圖片故障不阻擋運動紀錄。',
    ])
    heading('5.4 資料可靠性與可維護性',2)
    bullets([
        'NFR-13：所有 schema 變更具 Alembic migration；新空白資料庫可由 migration 與 seed 重建，14 張表與約束均存在。',
        'NFR-14：每日建立一份備份並保留最近 7 份；驗收時至少完成一次還原到獨立測試資料庫，核對表數、各表筆數與 FK。RPO ≤24 小時，RTO 目標 ≤60 分鐘。',
        'NFR-15：README 包含安裝、環境變數、啟動、migration、seed、測試、部署與 Live Demo；伺服器重新啟動不影響雲端資料。',
    ])


def acceptance():
    orientation(True)
    heading('6. 驗收測試案例')
    heading('6.1 功能驗收',2)
    paragraph('以隔離的測試資料庫執行，A、B 為一般會員，M 為管理員。每個案例記錄 HTTP 狀態、畫面結果及 SQL 查詢證據；破壞性與故障注入操作不得對正式使用者資料執行。表格沿用老師要求的案例編號、測試功能、前置條件、操作步驟與預期結果五欄。')
    table(['案例編號','測試功能','前置條件','操作步驟','預期結果'],TESTS,[1.9,3.25,5.1,6.6,8.55],10)
    orientation(False)
    heading('6.2 非功能驗收與追溯',2)
    table(['案例','對應需求','操作與通過標準'],[
        ('TC-26','NFR-01～03','依 5.1 負載執行並附原始報告；所有端點 P95 與錯誤率達門檻，200 人測試未破壞資料完整性。'),
        ('TC-27','NFR-04～08','檢查 HTTPS/TLS；以注入字串查詢不得洩漏他人資料；掃描 Git 機密；第 21 次失敗登入為 429；已撤銷 Token 為 401。'),
        ('TC-28','NFR-09～12','各指定瀏覽器與三種尺寸執行註冊至逐組儲存；無全頁水平捲動，按鈕尺寸、鍵盤與替代文字符合規格。'),
        ('TC-29','NFR-13～15','在空資料庫執行 migration/seed，再用備份還原；表數、筆數、FK 符合來源；記錄還原時間與 README 重現結果。'),
    ],[1.9,3.2,11.6])
    heading('6.3 功能資料表與測試對照',2)
    mappings=[
        ('FR-01','users / user_auth / user_sessions','TC-01～03'),('FR-02','user_auth / user_sessions / admins','TC-04～05'),
        ('FR-03','user_auth / user_sessions','TC-06～07'),('FR-04','users / body_records','TC-08'),
        ('FR-05','users / training_goals','TC-09'),('FR-06','user_body_parts / body_parts','TC-10'),
        ('FR-07','exercises / body_parts','TC-11'),('FR-08','admins / exercises / plan_exercises / workout_sets','TC-12'),
        ('FR-09','admins / training_goals / body_parts','TC-13'),('FR-10','users / user_body_parts / exercises / workout_plans / plan_days / plan_exercises','TC-14～15'),
        ('FR-11','workout_plans / plan_days / plan_exercises / workout_sessions','TC-16'),('FR-12','workout_sessions / plan_days / workout_plans','TC-17'),
        ('FR-13','workout_sessions / workout_sets / plan_exercises','TC-18～19'),('FR-14','workout_sets / workout_sessions','TC-20'),
        ('FR-15','workout_sessions / workout_sets','TC-21～22'),('FR-16','workout_sessions / workout_sets / exercises / body_parts','TC-23'),
        ('FR-17','workout_sessions / workout_sets / exercises / body_parts / body_records','TC-24'),('FR-18','body_records / users','TC-25'),
    ]
    table(['功能編號','主要資料表','驗收案例'],mappings,[2.1,10.8,3.8])


def appendices():
    heading('附錄 A　正規化與 CRUD 對應')
    paragraph('1NF：所有欄位保留單值，使用者的多個偏好部位拆成 user_body_parts。課表每個動作、實際訓練每一組皆為獨立列，不儲存逗號清單。')
    paragraph('2NF：user_body_parts 的候選鍵為 (user_id, body_part_id)，priority 取決於使用者對該部位的排序，依賴整個複合鍵；其他資料表的非鍵欄位均依賴各自的候選鍵。')
    paragraph('3NF：users 只存 training_goal_id，不重複存目標名稱；exercises 只存 body_part_id，不重複存部位名稱；workout_sets 不重複保存使用者、日期、場次狀態或動作名稱，這些值經由 JOIN 取得。')
    paragraph('目前體重以 body_records 最新值推導，訓練量以 weight_kg×reps 計算，不儲存可重算的總量。workout_plans.training_goal_id 描述產生時的課表目標，不是 users 目前目標的副本，兩者可不同。角色由 admins 存在與否推導，不在 user_auth 重複保存 role。')
    paragraph('users 與 user_auth 分表是責任與權限隔離，單靠拆表不代表已符合 3NF；本設計以實際函數相依及約束為依據。workout_sessions.user_id 表示實際執行者，即使來源課表刪除仍須保存；系統允許同一課表日跨日期執行，因此 plan_day_id 並非場次的候選鍵。')
    table(['資料表／模組','Create','Read','Update','Delete'],[
        ('users / auth','註冊','本人資料','個資、密碼','停用帳號'),('user_sessions','登入','驗證 jti','撤銷','到期清理'),
        ('user_body_parts','選擇部位','偏好清單','重排順位','移除偏好'),('exercises','管理員新增','搜尋／明細','修改或停用','未引用可刪'),
        ('workout_plans / days / plan_exercises','產生課表','課表明細','名稱／組次','刪課表'),('workout_sessions','開始運動','目前／歷史','完成／取消','刪除終止場次'),
        ('workout_sets','逐組儲存','動作歷史','修正量／次','刪錯誤組'),('body_records','新增體重','時間趨勢','修正數值','刪錯誤紀錄'),
        ('目標／部位 / admins','部署 seed','選項／授權','說明／受信任管理','非一般用戶功能'),
    ],[4.6,2.5,3.1,3.4,3.1],9.7)
    heading('附錄 B　SQL 查詢需求')
    paragraph('以下為以 SQLAlchemy text() 執行的參數化 SQL，:user_id 由已驗證登入狀態取得；在 DBeaver 執行時以參數視窗綁定，不將外部輸入拼接入字串。')
    heading('B.1 新增與更新個人資料',2)
    code('''INSERT INTO users (name)
VALUES (:name)
RETURNING user_id;

UPDATE users
SET age = :age,
    height_cm = :height_cm,
    training_experience = :experience
WHERE user_id = :user_id;

INSERT INTO body_records (user_id, record_date, weight_kg)
VALUES (:user_id, :record_date, :weight_kg)
ON CONFLICT (user_id, record_date)
DO UPDATE SET weight_kg = EXCLUDED.weight_kg;''')
    paragraph('註冊 INSERT 必須與 user_auth、user_sessions 共用交易；個人資料 UPDATE 與當日體重 upsert 共用另一個交易。不得只執行示範中的單表 INSERT 就宣稱註冊完成。')
    heading('B.2 JOIN 查詢訓練歷史',2)
    code('''SELECT
    ws.session_date, e.exercise_name, bp.body_part_name,
    s.set_number, s.weight_kg, s.reps,
    s.weight_kg * s.reps AS set_volume
FROM workout_sessions AS ws
JOIN workout_sets AS s
    ON s.workout_session_id = ws.workout_session_id
JOIN exercises AS e ON e.exercise_id = s.exercise_id
JOIN body_parts AS bp ON bp.body_part_id = e.body_part_id
WHERE ws.user_id = :user_id
    AND ws.status IN ('completed', 'cancelled')
    AND ws.session_date BETWEEN :start_date AND :end_date
ORDER BY ws.session_date DESC, ws.workout_session_id DESC,
    e.exercise_id, s.set_number;''')
    heading('B.3 每週訓練量與次數',2)
    code('''SELECT
    DATE_TRUNC('week', ws.session_date::timestamp)::date
        AS training_week,
    SUM(s.weight_kg * s.reps) AS weekly_volume,
    COUNT(DISTINCT ws.workout_session_id) AS workout_count
FROM workout_sessions AS ws
JOIN workout_sets AS s
    ON s.workout_session_id = ws.workout_session_id
WHERE ws.user_id = :user_id AND ws.status = 'completed'
    AND ws.session_date BETWEEN :start_date AND :end_date
GROUP BY training_week
ORDER BY training_week DESC;''')
    heading('B.4 最常訓練部位與動作',2)
    code('''SELECT
    bp.body_part_id, bp.body_part_name,
    COUNT(*) AS total_sets,
    SUM(s.weight_kg * s.reps) AS total_volume
FROM workout_sets AS s
JOIN workout_sessions AS ws
    ON ws.workout_session_id = s.workout_session_id
JOIN exercises AS e ON e.exercise_id = s.exercise_id
JOIN body_parts AS bp ON bp.body_part_id = e.body_part_id
WHERE ws.user_id = :user_id AND ws.status = 'completed'
    AND ws.session_date BETWEEN :start_date AND :end_date
GROUP BY bp.body_part_id, bp.body_part_name
ORDER BY total_sets DESC, bp.body_part_id ASC;''')
    paragraph('最常動作使用相同來源與篩選，將 SELECT／GROUP BY 改為 e.exercise_id、e.exercise_name，仍以 COUNT(*) 排序；同分依 exercise_id 升序。')
    heading('B.5 搜尋與體重彙總',2)
    code('''SELECT e.exercise_id, e.exercise_name, bp.body_part_name,
    e.difficulty_level, e.equipment, e.image_url
FROM exercises AS e
JOIN body_parts AS bp ON bp.body_part_id = e.body_part_id
WHERE e.is_active = TRUE
    AND (e.exercise_name ILIKE :pattern
         OR bp.body_part_name ILIKE :pattern)
ORDER BY e.exercise_name, e.exercise_id
LIMIT :page_size OFFSET :offset;

SELECT AVG(weight_kg) AS average_weight,
    MIN(weight_kg) AS minimum_weight,
    MAX(weight_kg) AS maximum_weight
FROM body_records
WHERE user_id = :user_id
    AND record_date BETWEEN :start_date AND :end_date;''')
    heading('B.6 刪除與防重複約束',2)
    code('''DELETE FROM body_records
WHERE body_record_id = :body_record_id
    AND user_id = :user_id
RETURNING body_record_id;

CREATE UNIQUE INDEX uq_active_workout_per_user
ON workout_sessions (user_id)
WHERE status = 'in_progress';

ALTER TABLE workout_sets
ADD CONSTRAINT uq_workout_sets_session_exercise_no
UNIQUE (workout_session_id, exercise_id, set_number);''')
    heading('附錄 C　課表推薦演算法')
    paragraph('首版使用資料庫與規則式演算法，不呼叫 LLM。規則輸出是可重現的課表模板，重量由使用者逐組輸入；系統不依目標自動指定高重量。LLM 可列為未來擴充，但須限制候選 exercise_id 並通過相同後端驗證。')
    heading('C.1 候選動作與每日配置',2)
    bullets([
        '難度順序為 Beginner < Intermediate < Advanced；候選必須 is_active=TRUE 且 difficulty≤使用者等級。同日不得重複 exercise_id。',
        '每次 30／60／90 分鐘分別最多安排 3／5／6 個動作，所有方案每日至少 3 個動作。固定熱身 5 分鐘；每下估計 4 秒、動作轉換 60 秒。',
        '單日預估秒數 = 300 + Σ[sets×reps×4 + (sets−1)×rest_seconds] + (動作數−1)×60；必須 ≤duration_minutes×60。',
        '先放入下表的每日三個必備部位，再以 (週出現次數÷偏好權重、priority、body_part_id) 升序挑選新增部位；偏好權重為 2，其餘為 1；同日不重複部位。',
        '同部位候選依 (本週使用次數、與使用者等級的差距、exercise_id) 升序選取。同樣的設定與資料集產生相同結果。',
    ])
    table(['每週天數','每日三個必備部位；依 D1 至 Dn 排列'],[
        ('2 天','D1：Chest、Back、Legs；D2：Shoulders、Glutes、Core。必須在額外欄位補 Biceps、Triceps。'),
        ('3 天','D1：Chest、Shoulders、Triceps；D2：Back、Biceps、Core；D3：Legs、Glutes、Core。'),
        ('4 天','D1：Chest、Back、Legs；D2：Shoulders、Glutes、Core；D3：Chest、Back、Biceps；D4：Legs、Glutes、Triceps。'),
        ('5 天','D1：Chest、Shoulders、Triceps；D2：Back、Biceps、Core；D3：Legs、Glutes、Core；D4：Chest、Back、Shoulders；D5：Legs、Glutes、Biceps。'),
        ('6 天','D1/D4：Chest、Shoulders、Triceps；D2/D5：Back、Biceps、Core；D3/D6：Legs、Glutes、Core。'),
    ],[2,14.7])
    paragraph('覆蓋規則：一週至少包含 Chest、Back、Shoulders、Legs、Glutes、Core；Biceps、Triceps 於每週容量足以容納八部位時亦必須各至少一次。2 天×30 分鐘只有 6 個動作欄位，明確採六個主要部位，手臂不保證獨立動作；不能同時承諾八部位且超過容量。任何未覆蓋的偏好部位需在輸出列為限制說明。')
    paragraph('額外欄位填入順序：先補足可容納的未覆蓋部位，再依偏好加權排序；若新增後超時，跳過該欄位並驗證最低配置。若必備配置亦超時，回 409 並提示改成 60／90 分鐘或調整目標，不保存不合規課表。')
    heading('C.2 組數次數與休息規則',2)
    table(['目標','原設計區間','首版固定輸出'],[
        ('Muscle Gain','3–4 組；8–12 下','3 組 × 10 下；休息 60 秒'),
        ('Fat Loss','2–4 組；12–15 下','2 組 × 12 下；休息 45 秒'),
        ('Strength','3–5 組；3–6 下','3 組 × 5 下；休息 120 秒'),
        ('General Fitness','2–3 組；10–15 下','2 組 × 12 下；休息 60 秒'),
    ],[4.1,5.9,6.7])
    paragraph('以上數值是專題軟體的預設參數；目的在定義可測試的輸出。使用者可以在規格範圍內修改組次，後端再驗證難度與時間。')
    heading('C.3 處理流程',2)
    code('''Input: authenticated user_id
1. Read profile, goal, preferred parts, days and duration.
2. Validate required settings and enabled candidates.
3. Build required body-part slots from the weekly table.
4. Fill missing coverage, then weighted preference slots.
5. For each day:
    a. Select unique exercise IDs from eligible candidates.
    b. Assign fixed sets, reps and rest for the selected goal.
    c. Check estimated time and required coverage.
6. Validate all days, IDs, difficulty and numeric limits.
7. Begin one transaction:
    a. INSERT one workout_plans row.
    b. INSERT all plan_days rows.
    c. INSERT all plan_exercises rows.
    d. COMMIT; on any error ROLLBACK.
8. Return the complete plan and coverage limitations.''')
    heading('附錄 D　部署與設定計畫')
    table(['項目','規格與交付內容'],[
        ('GitHub','沿用現有 repository；提交程式、migration、seed、測試、README、.env.example；排除 .env、備份與其他敏感資訊。'),
        ('公開網站','Render 或學校提供的雲端主機部署前端與 FastAPI；首頁可公開開啟，個人資料操作需註冊登入，無需 GitHub 帳號。'),
        ('雲端資料庫','雲端 PostgreSQL；以 DATABASE_URL 連線，禁止使用 localhost 資料庫作為公開網站唯一持久化來源。'),
        ('部署程序','建立雲端 DB → 注入環境變數 → migration → seed → 建置前端 → 啟動 API → HTTPS 健康檢查 → 執行核心驗收。'),
        ('帳號與角色','管理員資格由受信任部署程序授予；不在公開網頁提供升級管理員按鈕，也不把展示帳密寫入 repository。'),
        ('備份與還原','每日備份保留 7 份；備份存於受控位置；在獨立測試 DB 執行還原及筆數／FK 核對。'),
        ('Live Demo','正式驗收前在 README 填入通過測試的公開網址與測試日期；本規格書不宣稱既有部署已符合新規格。'),
        ('README','包含 Project Description、Database Design、ER Diagram、Installation、Environment Variables、Local Development、Deployment、Live Demo。'),
    ],[3.5,13.2])
    code('''# .env.example - placeholders only
DATABASE_URL=
SECRET_KEY=
JWT_ALGORITHM=HS256
ACCESS_TOKEN_EXPIRE_MINUTES=120
CORS_ORIGINS=

# Optional future integration; unused in version 1.1
LLM_API_KEY=''')
    paragraph('上述為新規格預定設定名稱；真正金鑰由雲端平台管理，.env 不提交 Git。首版沒有啟用 LLM，LLM_API_KEY 留空不影響排課。')
    heading('附錄 E　資料庫課程重點與結論')
    table(['課程重點','本專題對應'],[
        ('PK / FK / 1:1','users、user_auth 與 admins；4.1–4.3 可核對鍵與基數。'),
        ('1:N / M:N','users→workout_sessions→workout_sets；users↔body_parts；plan_days↔exercises。'),
        ('3NF','附錄 A 說明原子值、完整鍵依賴、傳遞相依與避免可推導重複值。'),
        ('CRUD / SQL','FR-04、08、11、13–18；附錄 A CRUD 對照與附錄 B 參數化 SQL。'),
        ('JOIN / GROUP BY / Aggregate','訓練歷史、SUM 訓練量、COUNT DISTINCT 次數、AVG/MIN/MAX 體重。'),
        ('交易與併發','註冊／課表跨表交易、逐組 COMMIT、場次鎖與唯一鍵；TC-03、15、17–21。'),
        ('管理與維護','migration、seed、最小 DB 權限、索引、備份與還原；TC-29。'),
    ],[4.9,11.8])
    paragraph('本系統以關聯式資料庫支援帳號、訓練設定、課表、每組紀錄及統計。使用者可從「今天做什麼」進入實際訓練，已完成組數逐筆保存，並透過可驗證的 SQL 查詢追蹤結果。實作以 migration、約束、交易與核心 API 為先，再完成介面與雲端驗收。')
    paragraph('文件依「資料庫系統規格書撰寫指南」的六大章節整理；參考來源為課程提供的 DBS.pdf。本文件版本 1.1 定義待實作與驗收的目標規格，完成情形以實際測試報告為準。')


def build_toc(anchor):
    # Word-compatible native field, with cached visible entries and bookmarks.
    begin = paragraph('')
    begin._p.append(x('r')); begin._p[-1].append(x('fldChar', fldCharType='begin'))
    r=x('r'); ins=x('instrText'); ins.set(qn('xml:space'),'preserve'); ins.text=' TOC \\o "1-2" \\h \\z \\u '
    r.append(ins); begin._p.append(r)
    r=x('r'); r.append(x('fldChar',fldCharType='separate')); begin._p.append(r)
    anchor._p.addprevious(begin._p)
    for text,level,name in HEADINGS:
        if level>2: continue
        sty='TOC '+str(level)
        if sty not in DOC.styles: DOC.styles.add_style(sty,1)
        p=paragraph('',sty)
        if text=='5. 非功能性需求': p.paragraph_format.page_break_before=True
        p.paragraph_format.left_indent=Cm(0 if level==1 else .55)
        p.paragraph_format.space_after=Pt(3)
        p.paragraph_format.line_spacing=Pt(13)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(16.6),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
        link=x('hyperlink',anchor=name)
        r=x('r'); pr=x('rPr'); pr.append(x('rFonts',ascii=FONT,hAnsi=FONT,eastAsia=FONT)); pr.append(x('sz',val=20)); r.append(pr)
        t=x('t'); t.text=text; r.append(t); link.append(r); p._p.append(link)
        p.add_run('\t')
        run=p.add_run('1'); run._r.set('data-page-bookmark',name)
        # Custom XML attributes are removed after caching; map by paragraph text.
        for r in p.runs: fonts(r,10,bold=level==1)
        anchor._p.addprevious(p._p)
    end=paragraph(''); r=x('r'); r.append(x('fldChar',fldCharType='end')); end._p.append(r)
    anchor._p.addprevious(end._p)
    anchor._p.getparent().remove(anchor._p)
    # Do not force an update dialog on opening; native TOC remains updatable.
    settings=DOC.settings.element
    for item in list(settings):
        if item.tag==qn('w:updateFields'): settings.remove(item)
    settings.append(x('updateFields',val='false'))


def main():
    anchor=setup()
    make_diagrams()
    heading('1. 專案概述')
    heading('1.1 開發背景與痛點',2)
    paragraph('本專題的開發動機來自我自己的健身經驗。我目前有健身習慣，但每次到健身房時，經常不知道當天應該做哪些動作、安排哪些部位，以及如何依自己的目標與程度選擇課表，因此需要先上網搜尋或詢問 AI。這些資訊分散在不同工具中，查詢與訓練紀錄無法接續，也不方便比較每次的重量、次數與進步情況。')
    paragraph('本系統整合會員資料、健身目標、偏好部位、運動動作資料庫、課表規劃與訓練紀錄。使用者查看當日課表後可直接開始運動，每完成一組立即存入資料庫，並從歷史紀錄與 Dashboard 觀察訓練量及體重變化。資料庫設計是專題核心，重點包含 PK、FK、正規化、交易、查詢與管理。')
    heading('1.2 系統目標',2)
    table(['目標','可衡量的完成條件'],[
        ('OBJ-01 課表規劃','完成基本資料與訓練設定後，可一次產生 2–6 天課表；每日 3–6 個動作、組次及休息清楚可見；暖機後 API P95≤3 秒。'),
        ('OBJ-02 持續紀錄','每組獨立保存；每組 API P95≤1 秒；相同內容重送不新增重複組；重新登入後可讀取所有已成功存檔資料。'),
        ('OBJ-03 可查詢可管理','完成 14 張關聯資料表及約束；以雲端 PostgreSQL 保存資料，18 項 FR 均可追溯到驗收案例；統計測試必須得到指定數值。'),
    ],[4.1,12.6])
    heading('1.3 目標使用者與權限',2)
    table(['角色','可執行操作','限制'],[
        ('訪客','查看公開首頁、註冊與登入。','不得讀取個人資料、課表、歷史或管理介面。'),
        ('一般會員 USER','管理本人資料、目標、課表、訓練組數、體重與統計。','僅限本人所有權；不得建立管理員資格或管理共用動作。'),
        ('管理員 ADMIN','具一般會員能力；另可管理動作及修改固定參考資料說明。','管理員資格由 admins 決定；本版無權瀏覽其他會員的私人訓練紀錄。'),
    ],[3.3,7.3,6.1])
    paragraph('網站使用流程：公開首頁 → 註冊／登入 → 基本資料 → 目標與偏好 → 產生課表 → 查看今日動作與圖片 → 開始運動 → 逐組儲存 → 完成訓練 → 查看歷史與統計。網站可公開存取，登入的是本系統帳號，不需要 GitHub 帳號。')
    paragraph('首版範圍包含註冊登入、可撤銷登入狀態、個人資料、規則式排課、圖片參考、逐組紀錄與統計；Email 寄信驗證、忘記密碼、支付、推播及 LLM 排課列為後續擴充。')
    heading('2. 系統架構與技術棧')
    heading('2.1 系統架構圖',2)
    paragraph('前端負責操作與顯示，後端負責身份、權限、欄位與排課驗證，PostgreSQL 保存持久資料。產生課表與統計在後端執行，瀏覽器不得直接取得資料庫憑證。')
    figure(ASSETS/'architecture.png','圖 1　系統架構與資料互動')
    heading('2.2 技術棧選型清單',2)
    table(['層級','選用技術／工具','選用理由與說明'],[
        ('前端','React / Vite / TypeScript / CSS','沿用現有前端基礎，提供響應式介面、表單、課表與訓練操作。'),
        ('後端','Python / FastAPI','沿用現有 API 架構，定義型別驗證、授權及交易處理。'),
        ('資料庫','PostgreSQL','支援關聯鍵、CHECK、交易、JOIN、GROUP BY 及彙總查詢。'),
        ('ORM 與驅動','SQLAlchemy / psycopg','與現有依賴一致；管理模型、連線與參數化 SQL。'),
        ('版本管理','Alembic / Git / GitHub','追蹤 schema 及程式版本，提供可重現的 migration。'),
        ('圖表','Recharts','依 API 統計結果繪製訓練與體重趨勢；不在前端另存正式資料。'),
        ('身份與密碼','JWT / bcrypt','本版要求新增／整合的認證能力；Token 2 小時，hash cost=10。'),
        ('部署與機密','Render 或學校雲端 / 環境變數','提供公開 HTTPS 與雲端 PostgreSQL；沿用原稿部署選項。'),
        ('資料管理與測試','DBeaver / pytest / HTTP 測試工具','查表、檢視 ERD、驗證 SQL、API 與交易結果。'),
    ],[3,5.3,8.4])
    paragraph('技術清單以現有 React、FastAPI、SQLAlchemy、Alembic 基礎及本版待實作功能為準。外部 Email、推播及 AI API 不屬於首版必要依賴，架構圖不將其列為已啟用服務。')
    requirements()
    database()
    nonfunctional()
    acceptance()
    appendices()
    build_toc(anchor)
    # Manual heading numbers follow the supplied university template.
    for p in DOC.paragraphs:
        if p.style.name == 'Heading 1': p.paragraph_format.page_break_before=True
    # Avoid blank page after an explicit section/page break before a chapter.
    for p in DOC.paragraphs:
        if p.text=='1. 專案概述': p.paragraph_format.page_break_before=False
    for r in DOC.element.xpath('.//w:r'):
        r.attrib.pop('data-page-bookmark',None)
    for sect in DOC.sections:
        for node in list(sect._sectPr):
            if node.tag==qn('w:docGrid'): sect._sectPr.remove(node)
    DOC.save(OUT)
    (TMP/'headings.json').write_text(json.dumps(HEADINGS,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'docx':str(OUT),'backup':str(BACKUP),'tables':len(DOC.tables),'headings':len(HEADINGS),'entities':len(SCHEMA),'functional_tests':len(TESTS)},ensure_ascii=False))


if __name__ == '__main__':
    main()
