#!/usr/bin/env python3
"""Build the editable v1.1 SRS from the reviewed Chinese Word source."""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT.parent / "參考資料" / "Fitness_Tracking_Management_System_SRS_中文版_教師建議修訂版.docx"
OUTPUT_DIR = ROOT / "output" / "specification"
OUTPUT = OUTPUT_DIR / "Fitness_Tracking_Management_System_SRS_中文版_正式修訂版_v1.1.docx"
DIAGRAM_DIR = ROOT / "docs" / "diagrams"

SLATE = "334155"
MID = "64748B"
PALE = "F1F5F9"
BORDER = "CBD5E1"
INK = RGBColor(31, 41, 55)
CJK_FONT = "Arial Unicode MS"


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "w:" + edge
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_keep(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def shade_cell(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_text(paragraph, text: str) -> None:
    paragraph.clear()
    paragraph.add_run(text)


def replace_exact(doc: Document, old: str, new: str) -> None:
    matches = [p for p in doc.paragraphs if p.text == old]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph for replacement, found {len(matches)}: {old[:80]}")
    set_text(matches[0], new)


def replace_prefix(doc: Document, prefix: str, new: str) -> None:
    matches = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph starting with {prefix!r}, found {len(matches)}")
    set_text(matches[0], new)


def reset_table(table, rows: list[list[str]]) -> None:
    while table.rows:
        table._tbl.remove(table.rows[0]._tr)
    for values in rows:
        row = table.add_row()
        set_row_keep(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.text = str(value)
            set_cell_margins(cell)
    set_repeat_table_header(table.rows[0])


def style_table(table) -> None:
    table.autofit = True
    for row_index, row in enumerate(table.rows):
        set_row_keep(row)
        for cell in row.cells:
            set_cell_margins(cell)
            shade_cell(cell, SLATE if row_index == 0 else ("FFFFFF" if row_index % 2 else "F8FAFC"))
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.05
                paragraph.paragraph_format.keep_together = True
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CJK_FONT)
                    run.font.size = Pt(8.6)
                    run.font.color.rgb = RGBColor(255, 255, 255) if row_index == 0 else INK
                    if row_index == 0:
                        run.bold = True


def insert_heading_and_table_after(doc: Document, anchor_table, heading: str, rows: list[list[str]]):
    paragraph = doc.add_paragraph(style="Heading 3")
    set_text(paragraph, heading)
    table = doc.add_table(rows=1, cols=len(rows[0]))
    reset_table(table, rows)
    anchor_table._tbl.addnext(paragraph._p)
    paragraph._p.addnext(table._tbl)
    return table


def replace_picture_before_caption(doc: Document, caption: str, image_path: Path, alt_text: str) -> None:
    target = next((p for p in doc.paragraphs if p.text.startswith(caption)), None)
    if target is None:
        raise RuntimeError(f"Caption not found: {caption}")
    previous = target._p.getprevious()
    while previous is not None and previous.tag == qn("w:p") and not previous.xpath(".//w:drawing"):
        if "".join(previous.itertext()).strip():
            break
        remove = previous
        previous = previous.getprevious()
        remove.getparent().remove(remove)
    if previous is not None and previous.tag == qn("w:p") and previous.xpath(".//w:drawing"):
        previous.getparent().remove(previous)
    picture_paragraph = doc.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.paragraph_format.space_before = Pt(6)
    picture_paragraph.paragraph_format.space_after = Pt(5)
    run = picture_paragraph.add_run()
    run.add_picture(str(image_path), width=Cm(16.0))
    for node in run._r.xpath(".//wp:docPr"):
        node.set("descr", alt_text)
    target._p.addprevious(picture_paragraph._p)


def format_document(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.line_spacing = 1.28
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True

    for style_name, size, color in (
        ("Heading 1", 18, SLATE),
        ("Heading 2", 14, "475569"),
        ("Heading 3", 11.5, "475569"),
    ):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12 if style_name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.widow_control = True
        if paragraph.style.name.startswith("Heading") or paragraph.style.name == "Caption":
            paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            run.font.name = "Arial"
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CJK_FONT)

    code_starts = (
        "BEGIN;",
        "SET TRANSACTION",
        "SELECT user_id",
        "SELECT workout_session_id",
        "WHERE user_id",
        "AND status",
        "-- ",
        "COMMIT;",
        "CREATE UNIQUE INDEX",
        "ON workout_sessions",
        "SELECT",
        "ws.workout_session_id",
        "ev.exercise_version_id",
        "ev.exercise_name",
        "SUM(",
        "FROM workout_sessions",
        "JOIN workout_",
        "ON ",
        "WHERE ws.user_id",
        "AND ws.status",
        "GROUP BY",
        "ORDER BY",
        "Input:",
    )
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(code_starts):
            p_pr = paragraph._p.get_or_add_pPr()
            shd = p_pr.find(qn("w:shd"))
            if shd is None:
                shd = OxmlElement("w:shd")
                p_pr.append(shd)
            shd.set(qn("w:fill"), PALE)
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.left_indent = Cm(0.45)
            paragraph.paragraph_format.right_indent = Cm(0.25)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            for run in paragraph.runs:
                run.font.name = "Consolas"
                run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CJK_FONT)
                run.font.size = Pt(8.5)

    for table in doc.tables:
        style_table(table)

    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.0)
        header = section.header
        if header.paragraphs:
            p = header.paragraphs[0]
        else:
            p = header.add_paragraph()
        set_text(p, "Fitness Tracking Management System — SRS v1.1")
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in p.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(148, 163, 184)


def update_toc_from_pdf(pdf_path: Path) -> None:
    doc = Document(OUTPUT)
    reader = PdfReader(str(pdf_path))
    page_map: dict[str, int] = {}

    def walk(items) -> None:
        for item in items:
            if isinstance(item, list):
                walk(item)
            elif isinstance(item, dict) and "/Title" in item:
                page_map[str(item["/Title"])] = reader.get_destination_page_number(item) + 1

    walk(reader.outline)
    missing: list[str] = []
    updated = 0
    for paragraph in doc.paragraphs:
        if paragraph.style.name.lower() not in ("toc 1", "toc 2", "目錄 13", "目錄 23"):
            continue
        texts = paragraph._p.xpath(".//w:t")
        if len(texts) < 2:
            continue
        # Word often splits a Chinese TOC title into several w:t nodes.  The
        # paragraph text before the tab is the reliable complete key, while
        # the final w:t node stores the cached page number.
        title = paragraph.text.rsplit("\t", 1)[0]
        if title not in page_map:
            missing.append(title)
            continue
        texts[-1].text = str(page_map[title])
        updated += 1
    if missing:
        raise RuntimeError(f"TOC headings missing from PDF outline: {missing}")
    doc.save(OUTPUT)
    print(json.dumps({"toc_entries_updated": updated, "pdf_pages": len(reader.pages)}, ensure_ascii=False))


def build() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document(SOURCE)

    # Cover and version.
    replace_exact(doc, "資料庫系統規格書", "軟體需求規格書（SRS）")
    cover_student = next(p for p in doc.paragraphs if p.text.startswith("學生："))
    set_text(cover_student, "版本：1.1（2026-10-06）\n學生：S11359001 陳苾妡　　指導老師：高啟洲")

    replacements = {
        "前端網頁與後端 API 使用相同的網站來源。受保護 API 以 Authorization: Bearer <Access JWT> 驗證身分；只有 /api/auth/refresh 與 /api/auth/logout 會讀取同源的 HttpOnly Refresh Cookie。系統不開放任意外部網站跨來源讀取 API，Refresh 請求另驗證 Origin。所有新增、修改及刪除操作均由後端重新檢查資料格式、身分與權限，不依賴前端驗證結果。":
            "前端與後端 API 使用同一網站來源。登入成功後，伺服器將效期 7,200 秒的 Access JWT 放入 Secure、HttpOnly、SameSite=Strict Cookie；瀏覽器 JavaScript 無法讀取 Token。寫入要求同時驗證 Origin／Referer 與 double-submit CSRF Token。系統不開放任意 CORS，所有新增、修改與刪除均由 PHP 重新驗證格式、所有權與權限。",
        "受保護 API 由 Authorization: Bearer <JWT> 取得 user_id，禁止以請求中的 user_id 作為授權依據。查詢子表時必須 JOIN 回擁有者。":
            "受保護 API 從 HttpOnly Cookie 取得 JWT，驗證簽章、exp、sub 與 jti 後，再 JOIN user_sessions、user_auth 與 users 取得當前使用者。禁止以請求中的 user_id 作為授權依據；查詢子表必須 JOIN 回擁有者。",
        "3. 使用 firebase/php-JWT 簽發固定 HS256 Access JWT：sub 為 user_id 的十進位字串、jti=user_session_id，另含 iat、exp，效期 7,200 秒且不含 role。另產生 256-bit 不透明 Refresh Token；瀏覽器只以 Secure、HttpOnly、SameSite=Strict、Path=/ 且不設定 Domain 的 __Host-refresh_token Cookie 保存原值，資料庫只存 SHA-256。提交成功後才回傳 Access Token 並設定 Cookie。":
            "3. 後端以固定 HS256 簽發 Access JWT：payload 僅含 sub、jti、iat、exp，效期 7,200 秒，不含 role。jti 對應 user_sessions.user_session_id。正式環境將 JWT 放入 __Host-fittrack_access Cookie（Secure、HttpOnly、SameSite=Strict、Path=/、無 Domain），並另設可由前端讀取的 CSRF Cookie。",
        "輸出／回應：POST /api/auth/register → 201；JSON 回傳 user_id、access_token、token_type=Bearer、expires_in=7200，回應同時設定 __Host-refresh_token Cookie。Refresh Token 不出現在 JSON、URL、前端儲存區或日誌。":
            "輸出／回應：POST /api/auth/register → 201；JSON 回傳 user_id、role 與 expires_in=7200，Access Token 僅透過 HttpOnly Cookie 傳送，不出現在 JSON、URL、localStorage、sessionStorage 或日誌。",
        "輸入資料：登入輸入 email、password；登出附帶 Bearer JWT。":
            "輸入資料：登入輸入 email、password；登出、改密碼與停用帳號由同源 Cookie 提供身分，寫入要求必須通過 CSRF 驗證。",
        "1. PHP 依正規化 Email 查詢 user_auth，以 password_verify() 比對密碼並檢查 is_active。成功登入時建立 user_sessions：Access JWT 效期 2 小時，Refresh Token 絕對效期 7 天。資料庫只存 Refresh Token 雜湊。":
            "1. PHP 依正規化 Email 查詢 user_auth，以 password_verify() 比對密碼並檢查 is_active。成功後建立 user_sessions，效期與 Access JWT 均為 2 小時。本版不導入 Refresh Token；逾期後重新登入。",
        "2. 每次受保護請求驗證簽章、固定 HS256、exp、iat 與 jti，再 JOIN user_sessions、user_auth 確認 sub 對應本人、期限有效、未撤銷且帳號啟用。只有管理端點另查 admins；一般會員端點不查管理員表。頁面重整後以 POST /api/auth/refresh、同源 Origin 檢查及 HttpOnly Cookie 輪替 Refresh Token，取得新的 Access Token；舊 Refresh Token 立即失效。管理寫入與資格撤銷的併發順序見 4.7。":
            "2. 每次受保護請求驗證簽章、固定 HS256、exp、iat 與 jti，再 JOIN user_sessions、user_auth、users 確認 sub 對應本人、Session 未到期、未撤銷且帳號啟用。管理端點每次再 LEFT JOIN／查詢 admins，而不信任 JWT 內的角色。頁面重整時呼叫 GET /api/auth/session，Cookie 仍有效即恢復使用者資料，不延長原到期時間。",
        "3. POST /api/auth/refresh 以 Cookie 原值雜湊查找工作階段，驗證 Origin、未撤銷與 refresh_expires_at 後，在同一交易輪替 Refresh Token 雜湊、更新 expires_at 與 last_refreshed_at，再發新 Access JWT。舊 Refresh Token 重用時撤銷整個工作階段。登出時寫入 revoked_at 並以過期 Set-Cookie 清除 Cookie。":
            "3. 登出時將當前 user_sessions.revoked_at 設為現在，並以過期 Set-Cookie 清除 Access 與 CSRF Cookie。修改密碼或停用帳號時，同一交易撤銷該帳號所有未撤銷 Session。",
        "輸出／回應：POST /api/auth/login → 200 與 Access JWT，並設定 Refresh Cookie；POST /api/auth/refresh → 200 與新 Access JWT、輪替 Cookie；POST /api/auth/logout → 204 並清除 Cookie。":
            "輸出／回應：POST /api/auth/login → 200 並設定 Access／CSRF Cookie；GET /api/auth/session → 200 與當前 user_id、name、email、role；POST /api/auth/logout → 204 並清除 Cookie。",
        "例外處理：錯誤帳密統一回 401 INVALID_CREDENTIALS；停用帳號、過期、撤銷或重用 Refresh Token 回 401。登入限流依 NFR-08 回 429；限流儲存不可用時回 503，不得繞過。":
            "例外處理：錯誤帳密統一回 401 INVALID_CREDENTIALS；停用帳號、過期或撤銷 Session 回 401。登入限流依 NFR-08 回 429 與 Retry-After；限流儲存不可用時回 503，不得繞過。",
        "FR-11 查看複製修訂與封存課表": "FR-11 查看、編輯與封存課表",
        "觸發條件：進入「我的課表」、複製修訂或確認封存。介面不將封存標示為永久刪除。":
            "觸發條件：進入「我的課表」、儲存課表調整或確認封存。介面不將封存標示為永久刪除。",
        "輸入資料：來源 workout_plan_id；POST /api/plans/{id}/revisions 接受 plan_name 與 exercises 陣列，各項含來源 plan_exercise_id 及欲替換的 exercise_version_id、target_sets、target_reps、rest_seconds。不得混入其他課表項目。":
            "輸入資料：workout_plan_id；PATCH 可傳 plan_name 與 exercises 陣列，每項含 plan_exercise_id、exercise_id、target_sets、target_reps、rest_seconds 與選填 target_weight_kg。不得混入其他使用者或其他課表項目。",
        "1. 讀取時 JOIN plan_days、plan_exercises 與其固定 exercise_versions，依 day_number、exercise_order 排序；不改連最新版本。":
            "1. 讀取時 JOIN plan_days、plan_exercises 與固定 exercise_versions，依 day_number、exercise_order 排序。課表有 in_progress 場次時，API 與 DB Trigger 均禁止修改或封存來源課表。",
        "2. 修訂時完整複製來源週課表與其規劃輸入，再套用修改、檢查難度、主要部位覆蓋與時間；新 workout_plan_id 的 copied_from_plan_id 指向來源。原始課表不變。本次是人工修訂，algorithm_version 記錄 manual-revision-v1。":
            "2. 沒有進行中場次時，PATCH 在單一交易內確認所有權、動作仍啟用且數值合法，再更新課表模板。已經開始的運動場次使用 workout_session_exercises 快照，所以不會被之後的模板編輯回寫。",
        "3. 已發布課表的名稱、規劃輸入、訓練日及動作不可原地修改。封存只填 archived_at，保留來源關聯；封存後禁止新開始，既有進行中場次仍可繼續。進行中若發現目標重量或組數不合適，使用者仍可照實記錄實際 weight_kg 與 reps，但不得改寫來源處方；要更改後續處方時，須先取消本場再複製修訂，或完成本場後修訂下一次課表。介面必須提示此差異。舊課表與動作版本不得物理刪除。":
            "3. 進行中場次只允許調整本場 workout_session_exercises 的 target_sets、target_reps、rest_seconds 與 target_weight_kg，不回寫 plan_exercises。DELETE /api/plans/{id} 是軟封存 archived_at，不刪除來源關聯或歷史。",
        "輸出／回應：GET /api/plans、GET /api/plans/{id} → 200；POST /api/plans/{id}/revisions → 201 與新 ID；POST /api/plans/{id}/archive → 204，可重送。原 PATCH/DELETE 課表端點不再提供，回 405。":
            "輸出／回應：GET /api/plans、GET /api/plans/{id} → 200；PATCH /api/plans/{id} → 200；DELETE /api/plans/{id} → 204（軟封存）。PATCH /api/workout-sessions/{sessionId}/exercises/{sessionExerciseId} → 200（僅調整本場）。",
        "例外處理：非本人或不存在回 404；無效動作版本、同日相同動作身份、超時回 422。封存不刪歷史；改用停用動作的新修訂不得發布。":
            "例外處理：非本人或不存在回 404；進行中場次鎖定來源課表時回 409 PLAN_IN_USE；無效動作、組次、休息或重量回 422。",
        "複製修訂不增加／刪除訓練日與動作列，也不改順序；調整整體日數須依目前會員設定重新產課。新課表發布前允許交易內編輯，COMMIT 時不得留下未發布或不完整課表。":
            "本版的 copied_from_plan_id 保留為日後「複製為新版課表」的來源欄位；現行介面的 PATCH 是編輯沒有進行中場次的模板。歷史依場次快照還原，不以目前模板反推。",
        "2. 同一交易 INSERT workout_sessions（user_id、started_at、status=in_progress）及 workout_session_plans（workout_session_id、plan_day_id）。完整歷史處方由來源日連回不可變週課表及固定動作版本取得，不再逐組複製名稱。":
            "2. 同一交易新增 workout_sessions、workout_session_plans，並把來源日的每個 plan_exercises 複製為 workout_session_exercises，保留 plan_exercise_id、exercise_version_id、順序與目標組次／休息／重量。場次快照建立後，來源課表之後的修改不影響該場歷史。",
        "3. uq_active_workout_per_user 防止同一會員提交兩筆 in_progress。遇此索引的 23505 必須先 ROLLBACK，再用新交易查既有場次；同日回 200、不同日回 409。其他唯一性錯誤不得誤判為續練。來源與完成組數的 DB 約束見 4.5。":
            "3. PostgreSQL 在 READ COMMITTED 下先以 SELECT ... FOR UPDATE 鎖定 users 資料列，使同一使用者的兩個開始請求依序執行。部分唯一索引 uq_workout_sessions_user_active 是最後防線。同日雙擊的實測結果為一個 201、一個 200，兩者回傳同一 workout_session_id；不同日則回 409。",
        "輸入資料：workout_session_id 取自網址；JSON 傳 plan_exercise_id、set_number 1–100、weight_kg 0–999.99、reps 1–999。不得另傳 user_id、exercise_id、名稱或圖片。首版不接受課表外臨時動作；若日後支援，須新增獨立的場次動作來源模型，不得以放寬 FK 取代。":
            "輸入資料：workout_session_id 取自網址；JSON 傳 workout_session_exercise_id、set_number 1–100、weight_kg 0–999.99、reps 1–999。不得另傳 user_id、exercise_id、名稱或圖片。首版不開放課表外臨時加做；日後若開放，必須新增可說明來源的場次動作模型，不能放寬 FK。",
        "1. 依固定順序鎖 users 與本人 workout_sessions，確認 in_progress；驗證 plan_exercise_id 所屬 plan_day_id 等於場次來源日。後端先檢查，資料庫觸發器再驗證，防止跨課表日組數。":
            "1. 後端以 workout_session_exercise_id JOIN 回 workout_sessions 驗證本人且 status=in_progress。workout_sets 的 FK 只能指向已建立的場次快照；workout_session_exercises 的 Trigger 另驗證 plan_exercise_id 必須屬於本場來源日，後端與 DB 雙重防止非法動作。",
        "2. 以 (workout_session_id, plan_exercise_id, set_number) 唯一鍵保存；同鍵同值重送回既有結果，不多建一組。同鍵不同值回 409，需使用 FR-14 修改。":
            "2. 以 UNIQUE(workout_session_exercise_id,set_number) 去除重送。同鍵同值重送回 200 與既有結果，同鍵不同值回 409 SET_ALREADY_EXISTS，必須使用 FR-14 編輯。",
        "1. 經 workout_sessions 驗證本人並先鎖父場次；允許修正 in_progress 或 completed 的 weight_kg、reps，不改所屬場次、plan_exercise_id 或 set_number。":
            "1. 經 workout_sets → workout_session_exercises → workout_sessions 驗證本人。允許修正 in_progress 或 completed 的 weight_kg、reps，不允許把組數移到其他場次動作，也不改 set_number。",
        "2. 明細 JOIN workout_sets → plan_exercises → exercise_versions 顯示當時名稱、部位、器材、步驟、圖片、目標組次及實際組數；另由來源日回查整份週課表及規劃輸入。Exercise History 以 exercise_versions.exercise_id 篩選，跨版本視為同一動作。":
            "2. 明細以 workout_sets → workout_session_exercises → exercise_versions 顯示當場固定的動作名稱、器材、步驟、圖片、目標處方與實際組數；再由 exercise_versions.exercise_id 連回穩定動作身份與主部位。Exercise History 以穩定 exercise_id 篩選，不同版本視為同一動作。",
        "訓練頁按不可變來源課表順序顯示當時動作、部位、器材、版本圖片與說明、目標組次、休息與輸入欄位。今天指使用者選擇的課表日，不依星期猜測；計時僅為操作提示。歷史完整性範圍與不保存項目見 4.6。":
            "訓練頁依 workout_session_exercises.exercise_order 顯示當場快照的動作、部位、器材、版本圖片與說明、目標組次、休息與輸入欄位。今天指使用者選擇的課表日，不依星期猜測。使用者可調整本場目標，但不回寫來源課表。",
        "介面必須區分「尚未儲存」「儲存中」「已儲存」「儲存失敗」。只有成功回應才標示已儲存；失敗保留輸入並允許重試。F5 後 Access Token 雖從記憶體消失，前端先呼叫 /api/auth/refresh，成功即恢復登入，不要求重新輸入帳密；Refresh 失敗才顯示登入頁。已提交的每組資料不受影響，尚未送出的欄位仍只存在目前頁面，離開前須提示。":
            "介面必須區分「尚未儲存」「儲存中」「已儲存」「儲存失敗」。只有成功回應才標示已儲存；失敗保留輸入並允許重試。F5 後瀏覽器會自動夾帶 HttpOnly Access Cookie，前端呼叫 GET /api/auth/session 恢復使用者資料；Cookie 到期、被撤銷或帳號停用才回到登入頁。已提交每組資料不受影響，未送出欄位只存在當前頁面。",
        "1. 在管理寫入交易內確認 admins 資格；image_url 只作 API 輸入，必須指向本站已核准的不可覆寫 PNG 路徑。部署程序驗證授權、保存圖片實體並計算 SHA-256，後端只存 image_sha256，不再存可由雜湊推導的 URL；不得下載任意用戶 URL。":
            "1. 管理端點在每次 Request 都由資料庫即時查詢 admins，不採信 JWT 內的角色。image_url 只保存 HTTPS 圖片網址；後端不主動下載使用者提供的遠端內容，並驗證長度與格式。動作描述修改會新增 exercise_versions，不覆寫舊版本。",
        "圖片 URL 並非圖片快照。發布版本前，部署流程須把合法素材存於不可覆寫路徑（例如 /assets/exercises/<sha256>.png），並在版本只保存內容雜湊；URL 由固定資源前綴、雜湊與 .png 副檔名推導；備份包括圖片實體。外部網址被覆寫或失效不能影響舊版本。沒有圖片時明確記 NULL，顯示非此動作示範的佔位圖，不宣稱還原不存在的圖片。":
            "exercise_versions 保存每個版本當時的 image_url 與文字說明，使舊課表與訓練歷史仍能指出當時引用的圖片網址；本系統不保存遠端圖片二進位內容，也不宣稱網址永久可用。若網址失效，介面顯示動作名稱、說明與明確的缺圖替代內容。需要永久影像保存時，應在未來加入具版本與備份政策的物件儲存服務。",
        "1. Workout History 預設列出本人 completed 與 cancelled。workout_session_details 以 INNER JOIN 取得非空 source_plan_day_id，從 started_at 換算台北 session_date，再依日期與 ID 倒序。":
            "1. Workout History 預設列出本人的 completed 與 cancelled 場次。查詢直接 JOIN workout_sessions、workout_session_plans、plan_days 與 workout_plans，依 started_at 與 workout_session_id 倒序；不存在未實作的唯讀 View。",
        "本版以 18 張領域資料表保存目前狀態、不可變處方與實際訓練，另以 login_rate_limits 保存跨 PHP Worker 共用的登入限流狀態，共 19 張應用資料表。exercise_versions 保存動作版本；planning_profiles 與 planning_profile_body_parts 保存產課時輸入。歷史經固定外鍵取得，不使用目前會員設定或最新動作版本代替當時資料。1 個唯讀 view workout_session_details 不另存數值。":
            "本版共有 20 張應用資料表與 25 條外鍵。exercise_versions 保存不可變的動作版本；planning_profiles 與 planning_profile_body_parts 保存產課當時的目標、年齡、身高、體重、程度、天數、時間與偏好順序；workout_session_exercises 保存開始運動時的動作版本與當場處方。login_failures 是跨 PHP Worker 共用的滑動視窗事件表。",
        "ERD 為鍵層級模型，PK／FK／UQ 分別表示主鍵、外鍵與唯一鍵。端點數字代表可參與筆數；圖示以基本鍵約束的基數為準，發布完整性及每場恰一來源由 4.5 的提交前約束補強。各圖重複節點是同一資料表，不是資料複製。R01 至 R23 與 4.3 節逐列 1:1 對應；login_rate_limits 沒有 FK，因此不編列 R 編號。":
            "ERD 以 PK／FK／UQ 表示主鍵、外鍵與唯一鍵。R01–R25 與 4.3 節逐列 1:1 對應；login_failures 不指向個人帳號而沒有 FK，因此不編 R 號。圖中重複出現的資料表代表同一實體，不是資料複製。",
        "圖 3 的課表自關聯 copied_from_plan_id 只記修訂來源；原課表不會被覆寫。圖 5 的 workout_sets 不再保存 exercise_id 或名稱，版本由 plan_exercise_id 唯一決定。場次來源日與組數所屬課表日必須一致，由 DB 觸發器驗證。":
            "圖 5 的 workout_session_exercises 是場次快照：既保留來源 plan_exercise_id，也固定 exercise_version_id 及當場目標。workout_sets 只指向場次動作，不重複保存 exercise_id、名稱或課表日。DB Trigger 驗證快照來源日與場次一致。",
        "workout_session_details 為 INNER JOIN workout_sessions 與 workout_session_plans 的唯讀 view，提供來源日與台北 session_date。沒有來源的損毀資料不可被默默解讀為正常歷史；一致性巡檢須另外以 LEFT JOIN 找出並報錯。":
            "歷史明細不另外建立重複資料的 View；PHP 使用參數化 SQL 直接 JOIN workout_sessions、workout_session_plans、workout_session_exercises、exercise_versions 與 workout_sets。若場次來源或快照缺失，FK、NOT NULL 與 Trigger 先阻止寫入，巡檢仍可用 LEFT JOIN 找出異常。",
        "除標示可空者外一律 NOT NULL。主鍵 INTEGER 使用 GENERATED BY DEFAULT AS IDENTITY；時間戳為 TIMESTAMPTZ，created_at 預設 CURRENT_TIMESTAMP。身高欄位使用 NUMERIC(5,2) 並限制 90–250 cm；所有 kg 欄位（體重快照、體重紀錄及訓練重量）統一使用 NUMERIC(6,2)，再以各欄 CHECK 限制合理範圍。未註明 DEFAULT 的欄位須由後端提供。FK 刪除政策見 4.3；CHAR(64) 雜湊須 CHECK 為 64 個小寫十六進位字元。":
            "除標示可空者外一律 NOT NULL。主鍵 INTEGER 使用 GENERATED BY DEFAULT AS IDENTITY；時間戳為 TIMESTAMPTZ，created_at 預設 CURRENT_TIMESTAMP。身高使用 NUMERIC(5,2) 並限制 100–250 cm；體重快照、體重紀錄及訓練重量均使用 NUMERIC(6,2)。資料庫的 CHECK 負責結構範圍，PHP 另拒絕超過兩位小數的輸入而不默默四捨五入。",
        "共 23 條 FK。發布交易須讓每份課表有且僅有一份 planning_profiles、1–8 個偏好、2–6 個訓練日及每日 3–6 動作。每場訓練在 COMMIT 時必有一筆來源；completed 必至少一組。圖中的 0 下限不是允許正常業務提交空內容，而是區分單純 FK 與跨表約束能力。":
            "共 25 條 FK。workout_session_exercises 透過場次、來源課表項目與動作版本三組 FK 保留當場快照；workout_sets 再指向該快照。completed 場次至少一組為跨表不變量，由 DEFERRABLE CONSTRAINT TRIGGER 於 COMMIT 前檢查。",
        "本版所有業務寫入明確採 READ COMMITTED。PHP 在 beginTransaction() 後先執行 SET TRANSACTION ISOLATION LEVEL READ COMMITTED，驗收以 SHOW transaction_isolation 確認，不依賴伺服器預設。每一 SQL 指令取新的已提交快照；SELECT FOR UPDATE 僅鎖既有資料列，不鎖「查不到的空集合」，也不等同 SERIALIZABLE。[1]":
            "PostgreSQL 連線建立後明確執行 SET SESSION CHARACTERISTICS AS TRANSACTION ISOLATION LEVEL READ COMMITTED，業務交易以此隔離層級執行；驗收以 SHOW transaction_isolation 確認為 read committed。每一 SQL 指令取新的已提交快照；SELECT FOR UPDATE 僅鎖既有資料列，不鎖「查不到的空集合」，也不等同 SERIALIZABLE。[1]",
        "只在短交易內做 DB 讀寫；圖片下載、HTTP 呼叫、等待使用者操作與昂貴雜湊均不放在持鎖區間。SET LOCAL lock_timeout='3s'、statement_timeout='5s' 作為本版上限，部署驗收確認與效能門檻相容。":
            "只在短交易內進行 DB 讀寫；圖片下載、外部 HTTP 呼叫、等待使用者操作與 password_hash() 不放在長時間持鎖區間。死結、鎖等待與連線中斷均必須回滾交易，不得回傳假成功。",
        "本版的「完整歷史」指能還原當次採用的整週課表、來源訓練日、動作版本、目標處方、產課輸入及現存實際組數，不表示備份整個網站或追蹤所有操作。開始訓練只綁定已發布的不可變課表；不同場次可共享同一份不可變處方，不必每次再複製整份。":
            "本版的「完整歷史」指能還原當次的產課輸入、來源課表日、動作版本、場次順序、場次目標處方與現存實際組數。開始運動時把 plan_exercises 複製為 workout_session_exercises，因此後續編輯或封存課表模板，不會改變已建立場次的處方。",
        "Access JWT 只攜帶 sub、jti、iat、exp，不保存 role。Refresh Token 是隨機不透明值，只由 HttpOnly Cookie 傳送，資料庫只存雜湊；兩者共用 user_session_id。JWT 簽章只能證明內容未被改動，不能保證管理員資格至今未變。本版查 user_sessions 以支援登出、密碼變更、Refresh 輪替與重用撤銷，因此不宣稱完全無狀態。":
            "Access JWT 僅攜帶 sub、jti、iat、exp，不存 role。JWT 簽章只能證明內容未被竄改，不能證明管理員資格仍然有效。因此每次受保護請求都查 user_sessions 與 user_auth，管理端點另即時查 admins。JWT 的用途是帶入簽名過的使用者與 Session 識別值，不是當作永久權限快取。",
        "每次受保護請求驗證 Token 與登入狀態；一般資源再驗所有權，管理端點再查 admins。GET /api/me 的 role 只是即時查詢結果供顯示。即使前端仍顯示管理按鈕，後端也以當次授權為準；DB 不可用或查詢失敗時不放行。[4]":
            "每次受保護請求驗證 Token 與登入狀態；一般資源再驗所有權，管理端點再查 admins。GET /api/me 與 GET /api/auth/session 的 role 是當下查詢結果，僅供介面顯示。admins 列被刪除後，下一個管理 API 即回 403；DB 不可用或查詢失敗時不放行。[4]",
        "管理寫入先以 FOR SHARE 鎖 user_auth、對應 user_sessions 及 admins 資格列，再修改資源並提交。撤銷管理員資格的維運程序依相同順序鎖 user_auth 後刪 admins；兩者互斥。若撤銷先提交，之後管理寫入回 403；若管理寫入先取得資格鎖，該筆可先提交，撤銷等待其結束。本版不承諾撤銷能取消已通過授權並持鎖的交易。讀取端點在授權查詢時作判斷。":
            "管理員資格在每一個 API Request 的驗證階段由 admins 即時查詢。若資格在該 Request 通過驗證後才被撤銷，已進入處理中的單一 Request 可能完成；撤銷不追溯取消已通過的交易，但下一個管理 Request 必須回 403。文件不宣稱使用尚未實作的資格列鎖定流程。",
        "既有 PK／UNIQUE 已有索引，不重複建立相同前綴。以 EXPLAIN (ANALYZE, BUFFERS) 驗證讀取計畫及資料量，不能由「已建索引」推論必然變快。migration 與備份帳號和應用帳號分離；Phinx 版本與種子需可重跑。":
            "既有 PK／UNIQUE 已有索引，不重複建立相同前綴。以 EXPLAIN (ANALYZE, BUFFERS) 驗證讀取計畫及資料量，不能由「已建索引」推論必然變快。Migration 與備份帳號和應用帳號分離；版本化 SQL、schema_migrations 與 Seeder 必須可在隔離測試庫重跑。",
        "登入失敗計數放在 PostgreSQL 的 login_rate_limits，使多個 PHP Worker 或 Container 共用同一份狀態。PHP 以伺服器端 RATE_LIMIT_SECRET 對正規化 IP 做 HMAC-SHA-256，不保存原始 IP。單一交易使用 INSERT ... ON CONFLICT DO UPDATE 原子更新視窗與 failed_count；成功登入刪除對應列。第 20 次失敗提交後設定 blocked_until，第 21 次請求在驗證密碼前回 429；封鎖到期後重設新視窗。":
            "登入限流使用 PostgreSQL login_failures 事件表。PHP 以 RATE_LIMIT_SECRET 對來源 IP 做 HMAC-SHA-256，不儲存原始 IP。每次請求在交易內對 ip_hash 取得 pg_advisory_xact_lock，先刪除 15 分鐘視窗外的事件，再 COUNT 目前視窗。前 20 次錯誤帳密回 401 並新增事件；第 21 次在驗密前回 429 與 Retry-After。成功登入會刪除該 ip_hash 的失敗事件。",
        "Render Cron Job 每日執行 php bin/cleanup-expired.php。腳本以 1,000 筆為一批，使用短交易清除 refresh_expires_at 或 revoked_at 已超過 7 天的 user_sessions，並清除 updated_at 超過 1 天且 blocked_until 已到期的 login_rate_limits。Access Token 的 expires_at 已到期但 Refresh Token 尚有效時不可刪除該登入工作階段。每批提交後再處理下一批，避免長交易與大量鎖定。":
            "Render Cron Service 以 UTC 19:00 排程，等於台北時間每日 03:00 執行 php bin/cleanup-expired.php。腳本取得全局 advisory lock，以 500 筆為一批、FOR UPDATE SKIP LOCKED 刪除 expires_at <= CURRENT_TIMESTAMP 的 user_sessions，並刪除超過 15 分鐘的 login_failures。每批獨立提交，避免長交易。Render 上排程服務需 Starter 或當時支援 Cron 的方案；若未啟用，應由維運者以外部排程執行同一命令。",
        "NFR-05：密碼以 PHP bcrypt cost=10 儲存。Access JWT 固定 HS256、效期 2 小時且不含 role，只存於前端記憶體；Refresh Token 為 256-bit 不透明值，效期 7 天，只放在 Secure、HttpOnly、SameSite=Strict Cookie，資料庫只存雜湊並在每次刷新時輪替。每次受保護請求查登入工作階段，管理端點另查 admins。撤銷生效與重用偵測見 4.7、TC-05、TC-18。":
            "NFR-05：密碼以 PHP bcrypt cost=10 儲存。Access JWT 固定 HS256、效期 2 小時且不含 role；正式環境只放在 Secure、HttpOnly、SameSite=Strict Cookie，不寫入 localStorage。寫入要求驗證同源與 CSRF Token。每次受保護請求查 user_sessions 與 user_auth，管理端點另即時查 admins。",
        "NFR-08：同 IP 於 15 分鐘內最多 20 次失敗登入，第 21 次回 429；成功登入不解除該 IP 的限制計數。":
            "NFR-08：同一來源 IP 在 15 分鐘滑動視窗內可有 20 次失敗驗密；第 21 次要求回 429 與 Retry-After。成功登入會清除該 ip_hash 的失敗事件。",
        "NFR-08：登入限流以 PostgreSQL 的 login_rate_limits 作為整個部署共用的原子計數器，不存於單一 PHP Worker。來源 IP 先以環境變數 RATE_LIMIT_SECRET 做 HMAC-SHA-256；15 分鐘內允許 20 次失敗，第 21 次回 429。計數以 INSERT ... ON CONFLICT 的原子 UPSERT 更新，成功登入刪除該列；DB 不可用時登入採 fail closed 並回 503。依 4.9 與 TC-27 驗收。":
            "NFR-08：登入限流以 PostgreSQL login_failures 與 pg_advisory_xact_lock 在多個 PHP Worker 間實作共用滑動視窗。來源 IP 先用 RATE_LIMIT_SECRET 做 HMAC-SHA-256；DB 不可用時登入 fail closed 回 503。X-Forwarded-For 只在 REMOTE_ADDR 屬 TRUSTED_PROXY_IPS 時採信，否則使用直接連線 IP。",
        "NFR-13：Phinx migration 必須重建 19 張應用資料表、23 條 FK、索引、觸發器、唯讀 view 及 seed；另有工具表 phinxlog。composer.JSON 與 composer.lock 固定 PHP 套件版本。":
            "NFR-13：database/postgres.sql 必須在空 PostgreSQL 建立 20 張應用資料表、25 條 FK、索引與觸發器。database/migrations/20261006_teacher_review_v1_1.sql 必須將舊版 15 表 Schema 原地升級，先做資料條件檢查、在單一交易內回填快照，失敗即全部回滾。",
        "NFR-15：README 包含 PHP、Composer 與 PDO_PGSQL 安裝、環境變數、啟動、migration、seed、PHPUnit 測試與部署。公開目錄限 public/，伺服器重啟不影響雲端資料。":
            "NFR-15：README 包含 PHP、PDO_PGSQL、環境變數、啟動、Migration、Seed、SQLite 煙霧測試、PostgreSQL 完整性／併發測試、Schema 一致性檢查、清理排程與部署。公開目錄限 public/，伺服器重啟不影響雲端 PostgreSQL。",
    }
    for old, new in replacements.items():
        replace_exact(doc, old, new)

    bodyweight_future = (
        "5. 未來擴充的徒手等效負荷 = 場次開始時體重快照 × 具版本的動作係數 × reps；"
        "估算值必須與實際外加重量分開顯示，本版不使用未定義係數。"
    )
    if not any(p.text == bodyweight_future for p in doc.paragraphs):
        anchor = next(p for p in doc.paragraphs if p.text.startswith("4. 無訓練資料回 weighted_volume=0"))
        future_paragraph = doc.add_paragraph(bodyweight_future, style=anchor.style)
        anchor._p.addnext(future_paragraph._p)

    # Replace obsolete fixed-window SQL block with the actual sliding-window pseudocode.
    start = next(i for i, p in enumerate(doc.paragraphs) if p.text.startswith("INSERT INTO login_rate_limits"))
    end = next(i for i, p in enumerate(doc.paragraphs) if p.text.startswith("Render Cron Service"))
    block = doc.paragraphs[start:end]
    set_text(block[0], "BEGIN;")
    lines = [
        "SELECT pg_advisory_xact_lock(hashtextextended(:ip_hash, 0));",
        "DELETE FROM login_failures",
        "WHERE ip_hash = :ip_hash",
        "  AND failed_at <= CURRENT_TIMESTAMP - INTERVAL '15 minutes';",
        "SELECT COUNT(*), MIN(failed_at)",
        "FROM login_failures WHERE ip_hash = :ip_hash;",
        "-- COUNT >= 20：COMMIT 後回 429 + Retry-After，不驗密。",
        "-- 密碼錯誤：INSERT 一筆 failed_at，COMMIT 後回 401。",
        "-- 登入成功：DELETE 該 ip_hash 的失敗事件。",
        "COMMIT;",
    ]
    for paragraph, line in zip(block[1:], lines):
        set_text(paragraph, line)
    for paragraph in block[len(lines) + 1 :]:
        paragraph._element.getparent().remove(paragraph._element)

    # Replace the obsolete Refresh Token cleanup example with the executable job SQL.
    cleanup_start = next(i for i, p in enumerate(doc.paragraphs) if p.text == "WITH doomed AS (")
    cleanup_end = next(i for i, p in enumerate(doc.paragraphs) if p.text.startswith("排程記錄開始時間"))
    cleanup_block = doc.paragraphs[cleanup_start:cleanup_end]
    cleanup_lines = [
        "WITH doomed AS (",
        "    SELECT user_session_id",
        "    FROM user_sessions",
        "    WHERE expires_at <= CURRENT_TIMESTAMP",
        "    ORDER BY expires_at",
        "    LIMIT 500",
        "    FOR UPDATE SKIP LOCKED",
        ")",
        "DELETE FROM user_sessions AS us",
        "USING doomed",
        "WHERE us.user_session_id = doomed.user_session_id;",
        "DELETE FROM login_failures",
        "WHERE failed_at <= CURRENT_TIMESTAMP - INTERVAL '15 minutes';",
    ]
    for paragraph, line in zip(cleanup_block, cleanup_lines):
        set_text(paragraph, line)
    for paragraph in cleanup_block[len(cleanup_lines) :]:
        paragraph._element.getparent().remove(paragraph._element)
    cleanup_note = next(p for p in doc.paragraphs if p.text.startswith("排程記錄開始時間"))
    set_text(
        cleanup_note,
        "排程輸出開始後的刪除筆數、執行毫秒與時區，但不記錄 Cookie、Token、雜湊或原始 IP。"
        "排程失敗不得影響使用者既有資料；監控連續失敗並由維運人員重跑。TC-27 驗證跨 Worker 限流，"
        "TC-31 驗證分批清理只刪除已到期 Session 與視窗外事件。",
    )

    # Data dictionaries.
    tables = list(doc.tables)
    reset_table(tables[0], [
        ["目標", "可衡量的完成條件"],
        ["課表規劃", "完成基本資料與訓練設定後，可一次產生 2–6 天課表；每日 3–6 個動作、組次及休息清楚可見；暖機後 API P95≤3 秒。"],
        ["持續紀錄", "每組獨立保存；每組 API P95≤1 秒；相同內容重送不新增重複組；重新載入後以有效 Cookie 恢復已存資料。"],
        ["可查詢可管理", "完成 20 張應用資料表、25 條外鍵與跨表約束；正式環境以雲端 PostgreSQL 保存資料。"],
    ])
    reset_table(tables[2], [
        ["層級", "選用技術／工具", "選用理由與說明"],
        ["前端", "HTML / CSS / JavaScript", "同源介面、Fetch API 與 RWD；不需額外前端編譯服務。"],
        ["後端", "PHP 8.4 / Apache", "路由、欄位驗證、Cookie／JWT Session 驗證、排課與 JSON API。"],
        ["資料庫", "PostgreSQL 16", "20 張應用表、25 條 FK，支援 Trigger、部分唯一索引、JOIN、GROUP BY 與 Aggregate；登入限流事件存於 login_failures，配合 advisory lock 在多個 PHP Worker 間共享，不依賴 Redis 或單一程序記憶體。"],
        ["資料存取", "PDO / PDO_PGSQL", "手寫參數化 SQL、交易與列鎖；關閉模擬預備語句。"],
        ["Schema 版本", "Versioned SQL Migration / Git", "bin/migrate.php 執行版本化 SQL；schema_migrations 記錄套用版本。"],
        ["身分與密碼", "HS256 JWT / PHP password_hash", "bcrypt cost=10；JWT 以 HttpOnly Cookie 傳送，jti 連回 user_sessions。"],
        ["部署", "Docker / Render / Environment Variables", "Web Service 連線 Cloud PostgreSQL；敏感資料不進 Git。"],
        ["排程", "Render Cron / cleanup-expired.php", "每日 03:00（Asia/Taipei）分批清理過期 Session 與限流事件。"],
        ["測試", "SQLite smoke / PostgreSQL integration", "功能回歸、Schema 一致性、Trigger、Migration 與兩個平行開始請求。"],
    ])
    reset_table(tables[6], [
        ["欄位", "型態與鍵", "語意、預設值及約束"],
        ["user_session_id", "UUID；PK", "JWT jti；DEFAULT gen_random_uuid()。"],
        ["auth_id", "INTEGER；FK", "→ user_auth.auth_id；ON DELETE CASCADE。"],
        ["created_at", "TIMESTAMPTZ", "DEFAULT CURRENT_TIMESTAMP。"],
        ["expires_at", "TIMESTAMPTZ", "必須晚於 created_at；登入後 7,200 秒。"],
        ["revoked_at", "TIMESTAMPTZ；可空", "登出、改密碼或停用時寫入；非 NULL 即失效。"],
    ])
    reset_table(tables[8], [
        ["欄位", "型態與鍵", "語意、預設值及約束"],
        ["login_failure_id", "BIGINT；PK", "每次失敗驗密事件的 IDENTITY 識別值。"],
        ["ip_hash", "CHAR(64)", "來源 IP 經 RATE_LIMIT_SECRET 做 HMAC-SHA-256；只允許 64 個小寫十六進位字元。"],
        ["failed_at", "TIMESTAMPTZ", "DEFAULT CURRENT_TIMESTAMP；15 分鐘滑動視窗的事件時間。"],
    ])
    heading_login = next(p for p in doc.paragraphs if "4.2.5" in p.text and "login_rate_limits" in p.text)
    set_text(heading_login, "表 4.2.5 login_failures")
    reset_table(tables[4], [
        ["欄位", "型態與鍵", "語意、預設值及約束"],
        ["user_id", "INTEGER；PK", "GENERATED BY DEFAULT AS IDENTITY；使用者識別碼。"],
        ["name", "VARCHAR(100)", "去頭尾空白後 1–100 字。"],
        ["gender", "VARCHAR(20)；可空", "Female / Male / Other；未提供時 NULL。"],
        ["age", "SMALLINT；可空", "13–100；完成 Profile 前可空。"],
        ["height_cm", "NUMERIC(5,2)；可空", "100.00–250.00 cm；完成 Profile 前可空。"],
        ["training_experience", "VARCHAR(20)；可空", "Beginner / Intermediate / Advanced。"],
        ["training_goal_id", "INTEGER；FK；可空", "→ training_goals；RESTRICT。"],
        ["training_days_per_week", "SMALLINT；可空", "2–6。"],
        ["training_duration_minutes", "SMALLINT；可空", "30 / 60 / 90。"],
        ["profile_completed_at", "TIMESTAMPTZ；可空", "完成基本資料並同交易建立初始體重後寫入。"],
        ["created_at", "TIMESTAMPTZ", "DEFAULT CURRENT_TIMESTAMP。"],
    ])
    reset_table(tables[12], [
        ["欄位", "型態與鍵", "語意、預設值及約束"],
        ["body_record_id", "INTEGER；PK", "身體紀錄識別碼。"],
        ["user_id", "INTEGER；FK", "→ users；CASCADE。"],
        ["record_date", "DATE", "不得晚於 CURRENT_DATE；UNIQUE(user_id,record_date)。"],
        ["weight_kg", "NUMERIC(6,2)", "20.00–500.00 kg。"],
        ["notes", "TEXT；可空", "最多 1,000 字。"],
    ])
    reset_table(tables[14], [
        ["欄位", "型態與鍵", "語意、預設值及約束"],
        ["exercise_version_id", "INTEGER；PK", "完整描述版本識別值。"],
        ["exercise_id", "INTEGER；FK", "→ exercises；RESTRICT。"],
        ["version_no", "INTEGER", ">=1；UNIQUE(exercise_id,version_no)。"],
        ["exercise_name", "VARCHAR(120)", "去空白後 1–120 字；名稱不是穩定身份。"],
        ["difficulty_level", "VARCHAR(20)", "Beginner / Intermediate / Advanced。"],
        ["equipment", "VARCHAR(100)；可空", "器材名稱；徒手可填 Bodyweight。"],
        ["description", "TEXT", "1–4,000 字；包含步驟與注意事項。"],
        ["image_url", "TEXT；可空", "參考圖片的 HTTPS URL；最長 2,048 字。"],
        ["created_at", "TIMESTAMPTZ", "DEFAULT CURRENT_TIMESTAMP；版本建立後禁止 UPDATE／DELETE。"],
    ])
    reset_table(tables[15], [
        ["欄位", "型態與鍵", "語意、預設值及約束"],
        ["workout_plan_id", "INTEGER；PK", "課表識別值。"],
        ["user_id", "INTEGER；FK", "→ users；CASCADE。"],
        ["copied_from_plan_id", "INTEGER；FK；可空", "→ workout_plans；RESTRICT；不得指向自己。"],
        ["plan_name", "VARCHAR(120)", "去空白後 1–120 字。"],
        ["published_at", "TIMESTAMPTZ", "NOT NULL；DEFAULT CURRENT_TIMESTAMP。"],
        ["archived_at", "TIMESTAMPTZ；可空", "軟封存時間；不可早於 created_at。"],
        ["created_at", "TIMESTAMPTZ", "DEFAULT CURRENT_TIMESTAMP。"],
    ])
    reset_table(tables[16], [
        ["欄位", "型態與鍵", "語意、預設值及約束"],
        ["workout_plan_id", "INTEGER；PK/FK", "→ workout_plans；ON DELETE RESTRICT。"],
        ["training_goal_id", "INTEGER；FK", "→ training_goals；RESTRICT。"],
        ["age_snapshot", "SMALLINT", "產課當時年齡；13–100。"],
        ["height_cm_snapshot", "NUMERIC(5,2)", "產課當時身高；100.00–250.00 cm。"],
        ["weight_kg_snapshot", "NUMERIC(6,2)", "產課當時最新體重；20.00–500.00 kg。"],
        ["weight_record_date", "DATE", "快照體重所屬日期。"],
        ["experience_snapshot", "VARCHAR(20)", "Beginner / Intermediate / Advanced。"],
        ["training_days_per_week", "SMALLINT", "2–6。"],
        ["duration_minutes", "SMALLINT", "30 / 60 / 90。"],
        ["algorithm_version", "VARCHAR(50)", "rules-v1、llm-validated-v1 或 legacy-backfill-v1。"],
        ["captured_at", "TIMESTAMPTZ", "DEFAULT CURRENT_TIMESTAMP；規劃輸入捕捉時間。"],
    ])
    reset_table(tables[19], [
        ["欄位", "型態與鍵", "語意、預設值及約束"],
        ["plan_exercise_id", "INTEGER；PK", "課表動作識別值。"],
        ["plan_day_id", "INTEGER；FK", "→ plan_days；RESTRICT。"],
        ["exercise_version_id", "INTEGER；FK", "→ exercise_versions；RESTRICT；固定描述版本。"],
        ["exercise_order", "SMALLINT", "1–10；UNIQUE(plan_day_id,exercise_order)。"],
        ["target_sets", "SMALLINT", "1–10。"],
        ["target_reps", "SMALLINT", "1–50。"],
        ["rest_seconds", "SMALLINT", "0–600 秒。"],
        ["target_weight_kg", "NUMERIC(6,2)；可空", "0.00–999.99 kg；目標重量。"],
    ])
    reset_table(tables[20], [
        ["欄位", "型態與鍵", "語意、預設值及約束"],
        ["workout_session_id", "INTEGER；PK", "實際訓練場次識別值。"],
        ["user_id", "INTEGER；FK", "→ users；CASCADE。"],
        ["status", "VARCHAR(20)", "in_progress / completed / cancelled；DEFAULT in_progress。"],
        ["started_at", "TIMESTAMPTZ", "DEFAULT CURRENT_TIMESTAMP。"],
        ["ended_at", "TIMESTAMPTZ；可空", "終止狀態必填且不得早於 started_at。"],
        ["notes", "TEXT；可空", "最多 1,000 字。"],
    ])
    reset_table(tables[22], [
        ["欄位", "型態與鍵", "語意、預設值及約束"],
        ["workout_session_exercise_id", "INTEGER；PK", "場次動作快照識別值。"],
        ["workout_session_id", "INTEGER；FK", "→ workout_sessions；CASCADE。"],
        ["plan_exercise_id", "INTEGER；FK", "→ plan_exercises；RESTRICT；Trigger 驗證屬於本場來源日。"],
        ["exercise_version_id", "INTEGER；FK", "→ exercise_versions；RESTRICT；必須等於來源 plan_exercise 的版本。"],
        ["exercise_order", "SMALLINT", "1–10；UNIQUE(workout_session_id,exercise_order)。"],
        ["target_sets", "SMALLINT", "1–10；可只調整本場。"],
        ["target_reps", "SMALLINT", "1–50；可只調整本場。"],
        ["rest_seconds", "SMALLINT", "0–600；可只調整本場。"],
        ["target_weight_kg", "NUMERIC(6,2)；可空", "0.00–999.99；目標重量，不等於實際紀錄。"],
        ["created_at", "TIMESTAMPTZ", "DEFAULT CURRENT_TIMESTAMP。"],
    ])
    heading_sets = next(p for p in doc.paragraphs if "4.2.19" in p.text and "workout_sets" in p.text)
    set_text(heading_sets, "表 4.2.19 workout_session_exercises")
    new_sets = insert_heading_and_table_after(doc, tables[22], "表 4.2.20 workout_sets", [
        ["欄位", "型態與鍵", "語意、預設值及約束"],
        ["workout_set_id", "INTEGER；PK", "每組實際訓練識別值。"],
        ["workout_session_exercise_id", "INTEGER；FK", "→ workout_session_exercises；CASCADE。"],
        ["set_number", "SMALLINT", "1–100；UNIQUE(workout_session_exercise_id,set_number)。"],
        ["weight_kg", "NUMERIC(6,2)", "0.00–999.99；外加重量，0 表示徒手。"],
        ["reps", "SMALLINT", "1–999。"],
        ["created_at", "TIMESTAMPTZ", "DEFAULT CURRENT_TIMESTAMP。"],
    ])

    relationship_rows = [
        ["編號", "父表 → 子表", "子表外鍵", "父端 : 子端／刪除"],
        ["R01", "users → user_auth", "user_id", "1 : 0..1；CASCADE"],
        ["R02", "user_auth → user_sessions", "auth_id", "1 : 0..N；CASCADE"],
        ["R03", "users → admins", "user_id", "1 : 0..1；CASCADE"],
        ["R04", "training_goals → users", "training_goal_id", "1 : 0..N；RESTRICT"],
        ["R05", "users → user_body_parts", "user_id", "1 : 0..N；CASCADE"],
        ["R06", "body_parts → user_body_parts", "body_part_id", "1 : 0..N；RESTRICT"],
        ["R07", "users → body_records", "user_id", "1 : 0..N；CASCADE"],
        ["R08", "body_parts → exercises", "body_part_id", "1 : 0..N；RESTRICT"],
        ["R09", "exercises → exercise_versions", "exercise_id", "1 : 0..N；RESTRICT"],
        ["R10", "users → workout_plans", "user_id", "1 : 0..N；CASCADE"],
        ["R11", "workout_plans → workout_plans", "copied_from_plan_id", "0..1 : 0..N；RESTRICT"],
        ["R12", "workout_plans → planning_profiles", "workout_plan_id", "1 : 0..1；RESTRICT"],
        ["R13", "training_goals → planning_profiles", "training_goal_id", "1 : 0..N；RESTRICT"],
        ["R14", "planning_profiles → planning_profile_body_parts", "workout_plan_id", "1 : 0..N；RESTRICT"],
        ["R15", "body_parts → planning_profile_body_parts", "body_part_id", "1 : 0..N；RESTRICT"],
        ["R16", "workout_plans → plan_days", "workout_plan_id", "1 : 0..N；RESTRICT"],
        ["R17", "plan_days → plan_exercises", "plan_day_id", "1 : 0..N；RESTRICT"],
        ["R18", "exercise_versions → plan_exercises", "exercise_version_id", "1 : 0..N；RESTRICT"],
        ["R19", "users → workout_sessions", "user_id", "1 : 0..N；CASCADE"],
        ["R20", "workout_sessions → workout_session_plans", "workout_session_id", "1 : 0..1；CASCADE"],
        ["R21", "plan_days → workout_session_plans", "plan_day_id", "1 : 0..N；RESTRICT"],
        ["R22", "workout_sessions → workout_session_exercises", "workout_session_id", "1 : 0..N；CASCADE"],
        ["R23", "plan_exercises → workout_session_exercises", "plan_exercise_id", "1 : 0..N；RESTRICT"],
        ["R24", "exercise_versions → workout_session_exercises", "exercise_version_id", "1 : 0..N；RESTRICT"],
        ["R25", "workout_session_exercises → workout_sets", "workout_session_exercise_id", "1 : 0..N；CASCADE"],
    ]
    reset_table(tables[23], relationship_rows)
    reset_table(tables[24], [
        ["唯一性或檢查", "規則"],
        ["帳號", "UNIQUE(user_auth.user_id)、UNIQUE(email)。"],
        ["偏好", "PK(user_id,body_part_id)、UNIQUE(user_id,priority)；產課快照另有相同的 plan 範圍鍵。"],
        ["動作版本", "UNIQUE(exercise_id,version_no)；Trigger 禁止 UPDATE／DELETE。"],
        ["課表", "UNIQUE(workout_plan_id,day_number)、UNIQUE(plan_day_id,exercise_order)；Trigger 防止同日重複穩定 exercise_id。"],
        ["場次", "UNIQUE(user_id) WHERE status='in_progress'；status 與 ended_at 組合 CHECK。"],
        ["場次動作", "UNIQUE(session_id,plan_exercise_id)、UNIQUE(session_id,exercise_order)；Trigger 驗證來源日與版本。"],
        ["每組", "UNIQUE(workout_session_exercise_id,set_number)；完成場次至少一組由延遲 Constraint Trigger 保護。"],
        ["體重", "UNIQUE(user_id,record_date)；Profile 完成後不得刪除最後一筆。"],
    ])
    reset_table(tables[27], [
        ["觸發時機", "保證內容"],
        ["exercise_versions BEFORE UPDATE/DELETE", "舊版本不可覆寫或刪除；修改動作必須 INSERT 新版本。"],
        ["plan_exercises BEFORE INSERT/UPDATE", "同一課表日不得放入同一穩定 exercise_id 的不同版本。"],
        ["workout_session_exercises BEFORE INSERT/UPDATE", "快照 plan_exercise_id 必須屬於場次來源日，exercise_version_id 必須與來源處方一致。"],
        ["body_records BEFORE DELETE", "以 advisory transaction lock 序列化同一使用者刪除；完整 Profile 至少保留一筆體重。"],
        ["workout_sessions / workout_sets / workout_session_exercises CONSTRAINT TRIGGER", "於 COMMIT 前確認 completed 場次存在至少一組；刪除最後一組或刪除其場次動作快照都會使交易失敗。"],
        ["workout_plans / plan_days / plan_exercises BEFORE UPDATE/DELETE", "來源課表有 in_progress 場次時禁止修改或刪除；本場調整必須寫入場次快照。"],
    ])
    reset_table(tables[28], [
        ["歷史項目", "保存位置", "後續修改的歷史行為"],
        ["產課目標、年齡、身高、體重與日期、程度、天數、時間", "planning_profiles", "存當時快照；不跟 users 或 body_records 同步。"],
        ["偏好部位與順序", "planning_profile_body_parts", "每個部位一列，不用 VARCHAR 清單；不跟目前偏好同步。"],
        ["動作名稱、難度、器材、說明與圖片 URL", "exercise_versions", "版本不可覆寫；穩定 exercise_id 可跨版本彙總。"],
        ["開始當下的動作順序、版本與目標處方", "workout_session_exercises", "場次建立後獨立於來源模板；本場可調整目標。"],
        ["場次時間、狀態與實際重量／次數", "workout_sessions、workout_sets", "實際紀錄與處方是不同事實；修正錯記後統計重算。"],
    ])
    reset_table(tables[29], [
        ["關係與候選鍵", "主要函數相依", "3NF 判定與刻意重複"],
        ["exercise_versions；PK version_id，AK(exercise_id,version_no)", "版本鍵 → 名稱、難度、器材、說明、圖片", "穩定主部位放 exercises；歷史描述不依賴目前版本。"],
        ["planning_profiles；PK workout_plan_id", "課表 ID → 當次所有排課輸入", "此為不同時點的業務事實，與 users 目前值不同是預期行為。"],
        ["planning_profile_body_parts；複合 PK", "(plan_id,body_part_id) → priority", "消除多值字串；非鍵不依賴部分主鍵。"],
        ["plan_exercises；PK plan_exercise_id", "處方 ID → 動作版本與目標", "不重複名稱、器材或部位。"],
        ["workout_session_exercises；PK session_exercise_id", "場次動作 ID → 來源、版本、順序與當場目標", "與 plan_exercises 數值重複是刻意的時間快照；建立後獨立變動。"],
        ["workout_sets；PK set_id，AK(session_exercise_id,set_number)", "組的鍵 → weight_kg、reps、created_at", "volume 由查詢計算；不保存可推導的動作名稱或場次 user_id。"],
    ])
    reset_table(tables[30], [
        ["索引", "用途與邊界"],
        ["uq_workout_sessions_user_active", "UNIQUE(user_id) WHERE status='in_progress'；同一使用者最後防重複。"],
        ["idx_user_sessions_expires", "user_sessions(expires_at)；排程分批清理。"],
        ["idx_login_failures_ip_time", "login_failures(ip_hash,failed_at DESC)；15 分鐘滑動視窗。"],
        ["idx_body_records_user_date", "body_records(user_id,record_date DESC,body_record_id DESC)；取目前體重。"],
        ["idx_exercises_body_active", "exercises(body_part_id,is_active)；候選動作篩選。"],
        ["idx_exercise_versions_latest", "exercise_versions(exercise_id,version_no DESC)；目錄取最新版。"],
        ["idx_workout_plans_user_created", "workout_plans(user_id,created_at DESC)；課表清單。"],
        ["idx_plan_days_plan / idx_plan_exercises_day", "課表依天數與動作順序讀取。"],
        ["idx_workout_sessions_user_started", "workout_sessions(user_id,started_at DESC)；訓練歷史。"],
        ["idx_session_exercises_session", "workout_session_exercises(workout_session_id,exercise_order)；場次詳情。"],
        ["idx_workout_sets_session_exercise", "workout_sets(workout_session_exercise_id,set_number)；逐組讀取與統計。"],
    ])

    # Acceptance cases aligned with executable v1.1 behavior.
    acceptance = tables[32]
    case_rows = {row.cells[0].text: row for row in acceptance.rows[1:]}
    case_updates = {
        "TC-01": ["FR-01 註冊", "新 Email", "註冊，查 users/auth/sessions 與 Set-Cookie。", "201；三表各 +1；JWT 無 role、exp−iat=7200；Access Cookie 為 Secure／HttpOnly／SameSite=Strict；無 admins。"],
        "TC-05": ["FR-02 登入、重整與登出", "有效帳號", "登入→F5→GET /auth/session→登出→用舊 Cookie 呼叫。", "重整後不重輸帳密；登出撤銷 Session 並清 Cookie；舊 Cookie 回 401。"],
        "TC-07": ["FR-04 基本資料", "已有會員", "存身高 100.00 及 250.00；送 99.99；同日改體重；檢查 profile_completed_at。", "邊界可存，非法 422；同日體重僅 1 筆；完成時間與初始體重同交易提交，users 無體重副本。"],
        "TC-11": ["FR-11 課表與本場目標", "課表 P 有進行中場次", "修改／封存 P；調整場次目標；取消後修改 P。", "來源模板操作回 409；本場目標 200 且不回寫 P；場次結束後可修改，舊歷史不變。"],
        "TC-12": ["FR-12 同日併發", "使用者無 active；兩個獨立 PostgreSQL 連線", "兩個 Request 同時開始同日。", "僅 1 場 in_progress；一個 201、一個 200；兩者 workout_session_id 相同。"],
        "TC-14": ["FR-13 組數重送與來源保護", "in_progress 且有場次快照", "同 session_exercise_id／set_number 同值送 2 次，再送不同值；直接 SQL 插入跨日快照。", "201／200／409；只 1 組。跨日快照被 DB Trigger 拒絕。"],
        "TC-16": ["FR-15 DB 非空約束", "空場次與已有一組的 completed 場次；非 superuser 測試角色", "直接 SQL 完成空場次、刪除 completed 最後一組、刪除其場次動作快照。", "三種繞過均回 23514 且交易回滾；無空 completed 場次留存。"],
        "TC-18": ["管理員撤權", "有效 JWT 與 admins 列", "刪除 admins 列後以原 Cookie 呼叫管理 API。", "立即 403；JWT 無 role；一般會員 API 仍依所有權執行。"],
        "TC-19": ["FR-16 完整歷史", "P1/V1 有已完成場次", "開始後修改課表模板、會員設定與動作產生 V2。", "歷史仍顯示場次快照的 V1 名稱、處方、順序，並查得當時排課輸入。"],
        "TC-20": ["FR-17 統計", "完成組 60×12、65×10、65×8 及徒手 0×15", "查 Dashboard 外加重量訓練量、總次數與徒手次數。", "external_load_volume=1890 kg·次；total_reps=45；bodyweight_reps=15；取消場次排除。"],
        "TC-27": ["安全與限流", "兩個 PHP Worker；相同來源 IP", "15 分鐘內連續 21 次失敗；換 IP；成功登入；模擬 DB 不可用。", "前 20 次 401，第 21 次 429 + Retry-After；不同 IP 獨立；成功後清除；DB 失敗回 503。"],
        "TC-28": ["FR-08 圖片版本", "既有歷史引用 V1 圖片 URL", "管理員修改 image_url 建立 V2；查詢舊場次；模擬舊圖片網址失效。", "V2 新增；舊場次仍引用 V1 的 image_url 與文字說明；缺圖時顯示文字替代內容，不宣稱系統備份圖片二進位檔。"],
        "TC-30": ["Schema／Migration 一致性", "空測試庫與含代表資料的舊版庫", "建立新 Schema；套用 v1.1 Migration；執行 schema_consistency.py。", "20 張應用表、25 條 FK；規劃快照與舊訓練歷史回填；Manifest 比對 PASS。"],
        "TC-31": ["Session 與限流清理", "同時有效、過期 Session 與超過 15 分鐘的失敗事件", "執行 cleanup-expired.php；查詢剩餘資料與 JSON 日誌。", "過期資料刪除，有效 Session 保留；每批≤500；日誌無 Token、雜湊或原始 IP。"],
    }
    for case_id, values in case_updates.items():
        row = case_rows[case_id]
        for index, value in enumerate([case_id, *values]):
            row.cells[index].text = value

    # SQL query example follows the snapshot chain.
    sql_replacements = {
        "JOIN workout_sets AS s": "JOIN workout_session_exercises AS wse",
        "ON s.workout_session_id = ws.workout_session_id": "ON wse.workout_session_id = ws.workout_session_id",
        "JOIN plan_exercises AS pe": "JOIN workout_sets AS s",
        "ON pe.plan_exercise_id = s.plan_exercise_id": "ON s.workout_session_exercise_id = wse.workout_session_exercise_id",
        "ON ev.exercise_version_id = pe.exercise_version_id": "ON ev.exercise_version_id = wse.exercise_version_id",
    }
    for old, new in sql_replacements.items():
        matches = [p for p in doc.paragraphs if p.text.strip() == old]
        if len(matches) != 1:
            raise RuntimeError(f"Expected one SQL line {old!r}, found {len(matches)}")
        set_text(matches[0], new)

    # Normalize the final index name in prose and SQL code blocks.
    for paragraph in doc.paragraphs:
        if "uq_active_workout_per_user" in paragraph.text:
            set_text(paragraph, paragraph.text.replace(
                "uq_active_workout_per_user", "uq_workout_sessions_user_active"
            ))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if "uq_active_workout_per_user" in paragraph.text:
                        set_text(paragraph, paragraph.text.replace(
                            "uq_active_workout_per_user", "uq_workout_sessions_user_active"
                        ))

    # Replace diagrams and captions.
    diagram_specs = [
        ("圖 1", "architecture_v1_1.png", "圖 1　系統架構與資料互動", "同源 PHP 架構、Cloud PostgreSQL 與排程清理"),
        ("圖 2", "erd_accounts_v1_1.png", "圖 2　帳號、目前偏好與登入保護 ERD", "R01 至 R07；login_failures 無 FK"),
        ("圖 3", "erd_plans_v1_1.png", "圖 3　課表與規劃輸入 ERD", "R10 至 R18，含產課輸入快照"),
        ("圖 4", "erd_exercises_v1_1.png", "圖 4　動作身份與不可變版本 ERD", "R08 至 R09，穩定身份與描述版本分離"),
        ("圖 5", "erd_workouts_v1_1.png", "圖 5　運動場次、來源快照與實際組數 ERD", "R19 至 R25，workout_sets 經場次動作快照連回版本"),
    ]
    for prefix, filename, caption, alt in diagram_specs:
        replace_picture_before_caption(doc, prefix, DIAGRAM_DIR / filename, alt)
        caption_para = next(p for p in doc.paragraphs if p.text.startswith(prefix))
        set_text(caption_para, caption)

    # Append concise revision appendix, while the detailed traceability remains a separate deliverable.
    appendix = doc.add_paragraph(style="Heading 1")
    set_text(appendix, "附錄 A 教師審查建議修訂摘要")
    appendix.paragraph_format.page_break_before = True
    intro = doc.add_paragraph(
        "本版已將教師審查的 12 項問題分別落實於規格書、ERD、PostgreSQL DDL、Migration、PHP API、排程與測試。完整的舊位置、新位置、修正方式與驗證證據見隨附 TEACHER_REVIEW_TRACEABILITY.md。"
    )
    intro.paragraph_format.keep_with_next = True
    appendix_table = doc.add_table(rows=1, cols=3)
    reset_table(appendix_table, [
        ["分類", "修正重點", "驗證方式"],
        ["歷史與關聯", "場次動作快照、不可變動作版本、完整產課輸入快照", "SQLite 功能回歸 + PostgreSQL FK／Trigger 反向測試"],
        ["業務邏輯", "最後一筆體重保護、本場目標調整、徒手次數獨立統計", "API 與直接 SQL 雙層測試"],
        ["安全與維運", "HttpOnly Cookie + CSRF、即時 admins 授權、滑動視窗限流、每日清理", "Cookie／授權測試、21 次限流邊界、Cron 實際執行"],
        ["Schema 與文件", "20 表／25 FK 一致、R01–R25 對照、正式 ERD、名詞大小寫統一", "Schema Manifest 自動比對、DOCX／PDF 全頁視覺檢查"],
    ])

    # Ask Word to refresh fields; cached TOC pages are updated after the first PDF render.
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    format_document(doc)
    style_table(new_sets)
    style_table(appendix_table)
    doc.core_properties.title = "Fitness Tracking Management System — SRS v1.1"
    doc.core_properties.subject = "資料庫系統課程專題規格書"
    doc.core_properties.comments = "依教師審查建議同步修訂文件、ERD、Schema、Migration、排程與測試。"
    doc.save(OUTPUT)
    print(json.dumps({"output": str(OUTPUT), "tables": len(doc.tables), "paragraphs": len(doc.paragraphs)}, ensure_ascii=False))


if __name__ == "__main__":
    if len(sys.argv) == 3 and sys.argv[1] == "--update-toc":
        update_toc_from_pdf(Path(sys.argv[2]))
    elif len(sys.argv) == 1:
        build()
    else:
        raise SystemExit("Usage: finalize_teacher_review_v1_1.py [--update-toc rendered.pdf]")
